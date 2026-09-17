# -*- coding: utf-8 -*-
"""Z_AUX_make_tablas_xlsx.py — Genera Tablas_Completas_Resultados.xlsx con los
resultados AGREGADOS A NIVEL REGIONAL (sin división por país), desagregados por
escenario y horizonte temporal, para los 9 indicadores del reporte:

  T1 Capacidad instalada de generación por fuente [GW]      (años 2030/2040/2050)
  T2 Generación anual por fuente [TWh]                      (años 2030/2040/2050)
  T3 Capacidad instalada de transmisión por tipo [GW]       (años 2030/2040/2050)
  T4 Kilómetros de línea instalados por tipo [km]           (periodos 2026-2050)
  T5 Capacidad instalada de almacenamiento [GW]             (años 2030/2040/2050)
  T6 Inversión anual promedio por tipo [MUSD/año]           (periodos 2026-2050)
  T7 Costo total del sistema, promedio anual [MUSD/año]     (periodos 2026-2050)
  T8 Costo unitario del sistema [USD/MWh]                   (periodos 2026-2050)
  T9 Consumo de combustibles fósiles [MBEP]                 (años 2030/2040/2050)

Es la versión REGIONAL y en Excel del antiguo Z_AUX_make_tablas_docx.py (tablas
por país del Entregable 2), actualizada a las fórmulas VIGENTES del proyecto:
  - Escenarios: los 14 de la corrida 2026-09-16 (solve_scenarios de
    Config_MOMF_T1_AB.yaml), en el MISMO orden y con los MISMOS nombres que el
    dashboard (SCENARIO_ALIAS de dashboard_config): familia OPT (BAC base, BSR
    sin repotenciación, BFA/BFB/BRA/BRB sensibilidades de costo), familia ETT
    (ISR base, INV con repotenciación, IFA/IFB/IRA/IRB), OPC (PLAN) y VSR
    (ETT-GP). BAU/OPT/VGB crudos (sin tope) se ignoran aunque estén en el CSV.
  - Transmisión (T3): fórmula nominal del chart_04 / fig_capacidad_transmision:
    Existentes = TCA−ANC (PLAN); Repo No Planif. = TCA−AMCI (RPO);
    Total = ΣTCA de los 3 grupos (SIN los factores /0.8 y /1.8 antiguos).
  - km (T4): load_capacity_and_distances() con fallback BAC->BAU / OPC->OPT
    (el xlsx de distancias no trae los códigos "con tope").
  - Costo unitario (T8): CAPEX anualizado AL VUELO desde CapitalInvestment
    (CRF con DISCOUNT_RATE/ASSET_LIFETIME de Z_AUX_capital_annualization_script);
    NO se usa CapitalInvestmentAnnualized del CSV (vacía en el snapshot).
  - Combustibles fósiles (T9): actividad MIN* fósil (chart_08a) convertida de
    PJ a millones de bep (conversión IEA de fig_combustibles_fosiles_2050.py);
    reemplaza a la tabla de emisiones del docx original.

El CSV (~1,4 GB) se reduce en UNA pasada por chunks a valores por
(Scenario, YEAR, TECHNOLOGY); la reducción se cachea en un pickle en
outputs/Figures/ (junto al xlsx), invalidado por mtime del CSV. NO correr en
paralelo con los fig_*.py (riesgo de OOM con el CSV grande).

Layout (spec 2026-09-16 §3): vive en scripts/figures/; solo scripts/ entra en
sys.path y todo se importa como paquete (figures.common.dashboard_config,
pipeline.Z_AUX_capital_annualization_script, common.relac_paths). Salida:
outputs/Figures/Tablas_Completas_Resultados.xlsx.

Uso (desde cualquier cwd):
    PYTHONUTF8=1 python scripts/figures/Z_AUX_make_tablas_xlsx.py
"""

from __future__ import annotations

import os
import sys

import pandas as pd

# Solo scripts/ entra en sys.path; todo se importa como paquete (spec 2026-09-16 §3).
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from common import relac_paths as P  # noqa: E402
from figures.common.dashboard_config import (  # noqa: E402
    COUNTRY_ISO3_BY_NAME,
    CSV_PATH,
    SCENARIO_ALIAS,
    SOURCE_FAMILY_NAMES,
    classify_line_group,
    classify_line_group_raw,
    classify_min_fossil_group,
    classify_tech_generation,
    classify_tech_type,
    load_capacity_and_distances,
)
from pipeline.Z_AUX_capital_annualization_script import (  # noqa: E402
    ASSET_LIFETIME,
    DISCOUNT_RATE,
    calculate_crf,
)

# Salidas bajo outputs/Figures/ (relac_paths.FIGURES), como el resto de figuras.
OUT_XLSX = str(P.FIGURES / "Tablas_Completas_Resultados.xlsx")
CACHE_PKL = str(P.FIGURES / ".tablas_xlsx_reduced.pkl")

# Los 14 escenarios de la corrida 2026-09-16, en el orden del dashboard
# (dashboard_config._SCEN_14): familia OPT, familia ETT, PLAN, ETT-GP.
SCENARIOS = [
    "BAC", "BFA", "BFB", "BRA", "BRB", "BSR",
    "ISR", "IFA", "IFB", "INV", "IRA", "IRB",
    "OPC", "VSR",
]
# Nombres de display idénticos a los del dashboard (build_dashboard.py usa
# SCENARIO_ALIAS.get(sc, sc)); sin etiquetas cortas propias.
ALIAS = {sc: SCENARIO_ALIAS.get(sc, sc) for sc in SCENARIOS}

YEARS = [2030, 2040, 2050]
YEAR_COLS = [str(y) for y in YEARS]
PERIODS = {
    "2026-2030": list(range(2026, 2031)),
    "2031-2035": list(range(2031, 2036)),
    "2036-2040": list(range(2036, 2041)),
    "2041-2045": list(range(2041, 2046)),
    "2046-2050": list(range(2046, 2051)),
}
PERIOD_COLS = list(PERIODS)
ACC_YEARS = list(range(2026, 2051))

FAMILY_ORDER = list(SOURCE_FAMILY_NAMES.values())
RENEWABLE_FAMS = {SOURCE_FAMILY_NAMES[c] for c in
                  ("PWRBIO", "PWRCSP", "PWRGEO", "PWRHYD",
                   "PWRSPV", "PWRWAS", "PWRWOF", "PWRWON")}
TRANS_CATS = ["Líneas Existentes", "Líneas Nuevas Planificadas",
              "Líneas Repotenciadas Planificadas", "Líneas Nuevas No Planificadas",
              "Líneas Repotenciadas No Planificadas"]
KM_CATS = ["Líneas Planificadas", "Líneas Nuevas No Planificadas",
           "Líneas Repotenciadas No Planificadas"]
INV_CATS = ["Generación", "Transmisión", "Almacenamiento"]
COST_CATS = ["CAPEX", "O&M", "Combustible", "Energía no Suministrada"]
FOSSIL_CATS = ["Carbón", "Gas natural", "Petróleo/derivados"]
STO_TYPES = {"PWRSDS": "Baterías corta duración", "PWRLDS": "Baterías larga duración",
             "PWRBDS": "Baterías (BDS)"}  # mismos prefijos que fig_almacenamiento_2050
VOLL_MUSD_PJ = 1500.0 / 3.6   # VOLL $1.500/MWh en MUSD por PJ (chart_14/15)
PJ_PER_TWH = 0.277778         # PJ -> TWh (misma constante que chart_11)

# IEA: 1 tep = 41,868 GJ; 1 tep ≈ 7,33 bep (barriles equivalentes de petróleo).
# Misma conversión que fig_combustibles_fosiles_2050.py.
GJ_PER_TEP = 41.868
BEP_PER_TEP = 7.33
GJ_PER_BEP = GJ_PER_TEP / BEP_PER_TEP  # ≈ 5,7119 GJ/bep
# PJ (10^15 J) / (GJ/bep) (10^9 J/bep) = 10^6 bep = 1 millón de bep -> el
# factor de escala es exactamente 1/GJ_PER_BEP, sin necesidad de otro término.
MBEP_PER_PJ = 1.0 / GJ_PER_BEP

# ================================================================
# 1) Reducción del CSV por (Scenario, YEAR, TECHNOLOGY)
# ----------------------------------------------------------------
# Parámetros de capacidad/costo: el valor por tech-año aparece en UNA sola de
# las muchas filas (timeslices), el resto NaN/repetido -> max. La producción sí
# se reparte por timeslice -> sum (mismos gotchas que el dashboard).
# ================================================================
MAX_COLS = ["TotalCapacityAnnual", "AccumulatedNewCapacity",
            "AccumulatedTotalAnnualMinCapacityInvestment", "CapitalInvestment",
            "OperatingCost", "NewCapacity", "TotalTechnologyAnnualActivity"]
SUM_COLS = ["ProductionByTechnology"]


def load_reduced() -> tuple[pd.DataFrame, pd.DataFrame]:
    mt = os.path.getmtime(CSV_PATH)
    if os.path.exists(CACHE_PKL):
        blob = pd.read_pickle(CACHE_PKL)
        if blob.get("mtime") == mt and blob.get("scen") == sorted(SCENARIOS):
            print("reducción cargada de cache", flush=True)
            return blob["per_max"], blob["per_sum"]

    usecols = ["Scenario", "YEAR", "TECHNOLOGY"] + MAX_COLS + SUM_COLS
    max_parts, sum_parts, nrows = [], [], 0
    for i, ch in enumerate(pd.read_csv(CSV_PATH, usecols=usecols,
                                       chunksize=2_000_000)):
        nrows += len(ch)
        ch = ch[ch["Scenario"].isin(SCENARIOS)]
        ch = ch.dropna(subset=["YEAR", "TECHNOLOGY"])
        ch["YEAR"] = pd.to_numeric(ch["YEAR"], errors="coerce")
        ch = ch.dropna(subset=["YEAR"])
        ch = ch[(ch["YEAR"] >= 2023) & (ch["YEAR"] <= 2050)]
        if ch.empty:
            continue
        ch["YEAR"] = ch["YEAR"].astype(int)
        for c in MAX_COLS + SUM_COLS:
            ch[c] = pd.to_numeric(ch[c], errors="coerce")
        key = ["Scenario", "YEAR", "TECHNOLOGY"]
        max_parts.append(ch.groupby(key)[MAX_COLS].max().reset_index())
        sum_parts.append(ch.groupby(key)[SUM_COLS].sum(min_count=1).reset_index())
        print(f"chunk {i + 1}: acumuladas {nrows:,} filas", flush=True)

    key = ["Scenario", "YEAR", "TECHNOLOGY"]
    per_max = (pd.concat(max_parts, ignore_index=True)
               .groupby(key)[MAX_COLS].max().reset_index())
    per_sum = (pd.concat(sum_parts, ignore_index=True)
               .groupby(key)[SUM_COLS].sum(min_count=1).reset_index())
    pd.to_pickle({"mtime": mt, "scen": sorted(SCENARIOS),
                  "per_max": per_max, "per_sum": per_sum}, CACHE_PKL)
    print(f"reducido: max={len(per_max):,}  sum={len(per_sum):,}", flush=True)
    return per_max, per_sum


per_max, per_sum = load_reduced()
missing = [sc for sc in SCENARIOS if sc not in set(per_max["Scenario"].unique())]
if missing:
    raise SystemExit(f"ERROR: escenarios sin datos en el CSV: {missing}")

# ================================================================
# 2) Métricas regionales — data[sc][fila][columna] = valor
# ================================================================
def nested() -> dict:
    return {sc: {} for sc in SCENARIOS}


def put(tbl, sc, row, col, val):
    tbl[sc].setdefault(row, {})[col] = float(val)


# ---- T1 capacidad de generación [GW] / T2 generación [TWh] ----
cap = per_max.dropna(subset=["TotalCapacityAnnual"]).copy()
cap["grp"] = cap["TECHNOLOGY"].map(classify_tech_generation)
cap = cap[cap["grp"].isin(["Renovable", "No Renovable"])]
cap["fam"] = cap["TECHNOLOGY"].str[:6].map(SOURCE_FAMILY_NAMES)

gen = per_sum.dropna(subset=["ProductionByTechnology"]).copy()
gen["grp"] = gen["TECHNOLOGY"].map(classify_tech_generation)
gen = gen[gen["grp"].isin(["Renovable", "No Renovable"])]
gen["fam"] = gen["TECHNOLOGY"].str[:6].map(SOURCE_FAMILY_NAMES)
gen["twh"] = gen["ProductionByTechnology"] / 3.6

cap_tbl, gen_tbl = nested(), nested()
for df, tbl, vcol in ((cap, cap_tbl, "TotalCapacityAnnual"), (gen, gen_tbl, "twh")):
    d = df[df["YEAR"].isin(YEARS) & df["fam"].notna()]
    for (sc, fam, y), v in d.groupby(["Scenario", "fam", "YEAR"])[vcol].sum().items():
        put(tbl, sc, fam, str(y), v)

# ---- T3 capacidad de transmisión [GW] (fórmula vigente del chart_04) ----
tr = per_max.copy()
tr["LG"] = tr["TECHNOLOGY"].map(classify_line_group)
tr = tr[tr["LG"].notna()]
cols4 = ["AccumulatedNewCapacity", "TotalCapacityAnnual",
         "AccumulatedTotalAnnualMinCapacityInvestment"]
g4 = tr[tr["YEAR"].isin(YEARS)].groupby(["Scenario", "YEAR", "LG"])[cols4].sum()
tr_tbl = nested()
for sc in SCENARIOS:
    for y in YEARS:
        def val(group, col):
            k = (sc, y, group)
            return float(g4.at[k, col]) if k in g4.index else 0.0

        acc_new_plan = val("PLAN", "AccumulatedNewCapacity")
        tca_plan = val("PLAN", "TotalCapacityAnnual")
        tca_nli = val("NLI", "TotalCapacityAnnual")
        tca_rpo = val("RPO", "TotalCapacityAnnual")
        acc_min_rpo = val("RPO", "AccumulatedTotalAnnualMinCapacityInvestment")
        vals = [tca_plan - acc_new_plan, acc_new_plan, acc_min_rpo,
                tca_nli, tca_rpo - acc_min_rpo]
        for cat, v in zip(TRANS_CATS, vals):
            put(tr_tbl, sc, cat, str(y), v)

# ---- T4 kilómetros de línea por tipo [km] (chart_10 / fig_km_lineas) ----
km = per_max.dropna(subset=["NewCapacity"]).copy()
km = km[km["TECHNOLOGY"].map(classify_tech_type) == "Transmisión"]
km = km[(km["YEAR"] >= ACC_YEARS[0]) & (km["YEAR"] <= ACC_YEARS[-1])]
_iso2name = {iso3: name for name, iso3 in COUNTRY_ISO3_BY_NAME.items()}
km["Country"] = km["TECHNOLOGY"].str[6:9].map(_iso2name)
km = km.dropna(subset=["Country"])
km = km.merge(load_capacity_and_distances(), on=["Scenario", "Country"], how="inner")
is_rnw = km["TECHNOLOGY"].str.startswith("RNW")
dist = km["Distance RNW"].where(is_rnw, km["Distance NRNW"])
factor = km["TECHNOLOGY"].str.contains("RPO", regex=False).map({True: 1.25, False: 1.0})
km["km"] = (km["NewCapacity"] / km["Capacity"]) * dist * factor
km["LG"] = km["TECHNOLOGY"].map(classify_line_group_raw)
km = km[km["LG"].isin(KM_CATS)]
km_tbl = nested()
km_acc = {}
for (sc, lg, y), v in km.groupby(["Scenario", "LG", "YEAR"])["km"].sum().items():
    for pname, yrs in PERIODS.items():
        if y in yrs:
            cur = km_tbl[sc].setdefault(lg, {}).get(pname, 0.0)
            km_tbl[sc][lg][pname] = cur + float(v)
    km_acc[sc] = km_acc.get(sc, 0.0) + float(v)

# ---- T5 capacidad de almacenamiento [GW] (chart_03 / fig_almacenamiento) ----
st = per_max.dropna(subset=["TotalCapacityAnnual"]).copy()
st = st[st["TECHNOLOGY"].str.startswith(tuple(STO_TYPES))]
st["tipo"] = st["TECHNOLOGY"].str[:6].map(STO_TYPES)
st_tbl = nested()
gst = (st[st["YEAR"].isin(YEARS)]
       .groupby(["Scenario", "tipo", "YEAR"])["TotalCapacityAnnual"].sum())
for (sc, tipo, y), v in gst.items():
    put(st_tbl, sc, tipo, str(y), v)

# ---- T6 inversión anual promedio [MUSD/año] (chart_05: PWRTRN/RNWTRN ÷1.2) ----
# BCK EXCLUIDO como en chart_05/fig_inversion_*: su CapitalInvestment es una
# penalización big-M (en ETT/ETT-GP llega a ~1e8 MUSD, 100× la inversión real);
# el backstop se valora aparte como ENS al VOLL en T7.
inv = per_max.dropna(subset=["CapitalInvestment"]).copy()
inv = inv[~inv["TECHNOLOGY"].str.contains("BCK", regex=False)]
adj = inv["TECHNOLOGY"].str.startswith(("PWRTRN", "RNWTRN"))
inv["CapitalInvestment"] = inv["CapitalInvestment"].where(
    ~adj, inv["CapitalInvestment"] / 1.2)
inv["TechType"] = inv["TECHNOLOGY"].map(classify_tech_type)
inv = inv[inv["TechType"].notna()]
inv_tbl = nested()
inv_acc = {}
for (sc, tt, y), v in inv.groupby(["Scenario", "TechType", "YEAR"])["CapitalInvestment"].sum().items():
    for pname, yrs in PERIODS.items():
        if y in yrs:
            cur = inv_tbl[sc].setdefault(tt, {}).get(pname, 0.0)
            inv_tbl[sc][tt][pname] = cur + float(v)
    if y in ACC_YEARS:
        inv_acc[sc] = inv_acc.get(sc, 0.0) + float(v)
for sc in SCENARIOS:
    for tt, cols in inv_tbl[sc].items():
        for pname in list(cols):
            cols[pname] /= len(PERIODS[pname])

# ---- T7 costos del sistema [MUSD/año] (chart_14 nonsupplied) ----
# BCK excluido de CAPEX/O&M/Combustible (penalty big-M artificial); su energía
# se valora aparte como ENS al VOLL. CAPEX de PWRTRN/RNWTRN ÷1.2.
co = per_max.copy()
co["TechType"] = co["TECHNOLOGY"].map(classify_tech_type)
is_bck = co["TECHNOLOGY"].str.contains("BCK", regex=False)
is_infra = co["TechType"].notna() & ~is_bck
capex = co["CapitalInvestment"].fillna(0.0)
adj = co["TECHNOLOGY"].str.startswith(("PWRTRN", "RNWTRN"))
capex = capex.where(~adj, capex / 1.2)
op = co["OperatingCost"].fillna(0.0)
co["CAPEX"] = capex.where(is_infra, 0.0)
co["O&M"] = op.where(is_infra, 0.0)
co["Combustible"] = op.where(~is_infra & ~is_bck, 0.0)
gco = co.groupby(["Scenario", "YEAR"])[["CAPEX", "O&M", "Combustible"]].sum()

ens = per_sum[per_sum["TECHNOLOGY"].str.contains("BCK", regex=False)].copy()
ens["ens"] = ens["ProductionByTechnology"].fillna(0.0) * VOLL_MUSD_PJ
gens = ens.groupby(["Scenario", "YEAR"])["ens"].sum()

cost_tbl = nested()
cost_acc = {}
cost_rows = [((sc, y), cat, float(row[cat]))
             for (sc, y), row in gco.iterrows() for cat in COST_CATS[:3]]
cost_rows += [((sc, y), COST_CATS[3], float(v)) for (sc, y), v in gens.items()]
for (sc, y), cat, v in cost_rows:
    for pname, yrs in PERIODS.items():
        if y in yrs:
            cur = cost_tbl[sc].setdefault(cat, {}).get(pname, 0.0)
            cost_tbl[sc][cat][pname] = cur + v
    if y in ACC_YEARS:
        cost_acc[sc] = cost_acc.get(sc, 0.0) + v
for sc in SCENARIOS:
    for cat, cols in cost_tbl[sc].items():
        for pname in list(cols):
            cols[pname] /= len(PERIODS[pname])

# ---- T8 costo unitario [USD/MWh] (fig_costo_unitario / chart_11) ----
# (CAPEX anualizado al vuelo + OperatingCost) / producción, sin BCK en
# numerador ni denominador; promedio de los ratios anuales de cada periodo.
_no_bck = ~per_max["TECHNOLOGY"].str.contains("BCK", regex=False)
uc = per_max[_no_bck].copy()
inv_y = uc.groupby(["Scenario", "YEAR"])["CapitalInvestment"].sum(min_count=1).fillna(0.0)
op_y = uc.groupby(["Scenario", "YEAR"])["OperatingCost"].sum(min_count=1).fillna(0.0)
den_y = (per_sum[~per_sum["TECHNOLOGY"].str.contains("BCK", regex=False)]
         .groupby(["Scenario", "YEAR"])["ProductionByTechnology"].sum() * PJ_PER_TWH)
crf = calculate_crf(DISCOUNT_RATE, ASSET_LIFETIME)
uc_tbl = nested()
uc_avg = {}
for sc in SCENARIOS:
    ratios = {}
    for y in ACC_YEARS:
        if (sc, y) not in den_y.index or den_y.loc[(sc, y)] <= 0:
            continue
        win = [(sc, w) for w in range(y - ASSET_LIFETIME + 1, y + 1)]
        ann = sum(float(inv_y.get(k, 0.0)) for k in win) * crf
        num = ann + float(op_y.get((sc, y), 0.0))
        ratios[y] = num / float(den_y.loc[(sc, y)])
    for pname, yrs in PERIODS.items():
        vals = [ratios[y] for y in yrs if y in ratios]
        if vals:
            put(uc_tbl, sc, "Costo unitario", pname, sum(vals) / len(vals))
    if ratios:
        uc_avg[sc] = sum(ratios.values()) / len(ratios)

# ---- T9 consumo de combustibles fósiles [MBEP] (chart_08a / fig_combustibles) --
# Actividad MIN* fósil en PJ, convertida a millones de barriles equivalentes de
# petróleo con MBEP_PER_PJ (misma conversión IEA que la figura del reporte).
fo = per_max.dropna(subset=["TotalTechnologyAnnualActivity"]).copy()
fo["grp"] = fo["TECHNOLOGY"].map(classify_min_fossil_group)
fo = fo[fo["grp"].notna()]
fo_tbl = nested()
gfo = fo.groupby(["Scenario", "grp", "YEAR"])["TotalTechnologyAnnualActivity"].sum()
fo_acc = {}
for (sc, grp, y), v in gfo.items():
    if y in YEARS:
        put(fo_tbl, sc, grp, str(y), v * MBEP_PER_PJ)
    if y in ACC_YEARS:
        fo_acc[sc] = fo_acc.get(sc, 0.0) + float(v) * MBEP_PER_PJ

# ================================================================
# 3) Excel
# ================================================================
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TEAL = "23978E"
BAND = "DCEDEB"
GREY = "BFBFBF"
_side = Side(style="thin", color=GREY)
BORDER = Border(left=_side, right=_side, top=_side, bottom=_side)
F_BASE = Font(name="Arial", size=9)
F_BOLD = Font(name="Arial", size=9, bold=True)
F_HDR = Font(name="Arial", size=9, bold=True, color="FFFFFF")
F_TITLE = Font(name="Arial", size=10, bold=True)
F_NOTE = Font(name="Arial", size=8, italic=True)
FILL_HDR = PatternFill("solid", fgColor=TEAL)
FILL_BAND = PatternFill("solid", fgColor=BAND)
AL_C = Alignment(horizontal="center", vertical="center", wrap_text=True)
AL_L = Alignment(horizontal="left", vertical="center")
AL_R = Alignment(horizontal="right", vertical="center")


def numfmt(dec: int) -> str:
    return "#,##0" + ("." + "0" * dec if dec else "")


def _hdr(ws, r, c, text):
    cell = ws.cell(row=r, column=c, value=text)
    cell.font, cell.fill, cell.border, cell.alignment = F_HDR, FILL_HDR, BORDER, AL_C


def _val(ws, r, c, v, dec, bold=False):
    cell = ws.cell(row=r, column=c, value=round(float(v), 6))
    cell.font, cell.border, cell.alignment = (F_BOLD if bold else F_BASE), BORDER, AL_R
    cell.number_format = numfmt(dec)


def _lab(ws, r, c, text, bold=False, band=False):
    cell = ws.cell(row=r, column=c, value=text)
    cell.font = F_BOLD if bold else F_BASE
    cell.border = BORDER
    cell.alignment = AL_C if band else AL_L
    if band:
        cell.fill = FILL_BAND


def rows_with_totals(data_sc: dict, labels: list[str], cols: list[str],
                     totals: str | None) -> list[tuple[str, list[float], bool]]:
    """[(etiqueta, valores, es_total)] para una banda de escenario."""
    band = [(lab, [data_sc.get(lab, {}).get(c, 0.0) for c in cols], False)
            for lab in labels]
    if totals in ("sum", "sum_first"):
        tot = [sum(v[i] for _, v, _ in band) for i in range(len(cols))]
        item = ("Total", tot, True)
        band = [item] + band if totals == "sum_first" else band + [item]
    return band


def write_banded(wb, sheet, title, label_hdr, cols, data, labels, dec,
                 totals="sum", subtotal_groups=None, note=None):
    """Hoja con tabla de banda por escenario (col A vMerge = escenario)."""
    ws = wb.create_sheet(sheet)
    ws.sheet_view.showGridLines = False
    ncols = 2 + len(cols)
    ws.cell(row=1, column=1, value=title).font = F_TITLE
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    r = 3
    _hdr(ws, r, 1, "Escenario")
    _hdr(ws, r, 2, label_hdr)
    for j, cname in enumerate(cols):
        _hdr(ws, r, 3 + j, cname)
    r += 1
    for sc in SCENARIOS:
        band = rows_with_totals(data[sc], labels, cols, totals)
        if subtotal_groups:
            # subtotales intermedios (T1/T2: Renovable / No Renovable)
            band = []
            for gname, glabels in subtotal_groups:
                glabels = [l for l in glabels if l in labels]
                for lab in glabels:
                    band.append((lab, [data[sc].get(lab, {}).get(c, 0.0)
                                       for c in cols], False))
                band.append((f"Subtotal {gname}",
                             [sum(data[sc].get(l, {}).get(c, 0.0) for l in glabels)
                              for c in cols], True))
            band.append(("Total",
                         [sum(data[sc].get(l, {}).get(c, 0.0) for l in labels)
                          for c in cols], True))
        r0 = r
        for lab, vals, is_tot in band:
            _lab(ws, r, 2, lab, bold=is_tot)
            for j, v in enumerate(vals):
                _val(ws, r, 3 + j, v, dec, bold=is_tot)
            r += 1
        ws.merge_cells(start_row=r0, start_column=1, end_row=r - 1, end_column=1)
        for rr in range(r0, r):
            _lab(ws, rr, 1, ALIAS[sc] if rr == r0 else None, bold=True, band=True)
    ws.cell(row=r + 1, column=1, value="Fuente: Elaboración propia").font = F_NOTE
    if note:
        ws.cell(row=r + 2, column=1, value=note).font = F_NOTE
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 36
    for j in range(len(cols)):
        ws.column_dimensions[get_column_letter(3 + j)].width = 12.5
    ws.freeze_panes = "C4"
    return ws


def write_scenario_rows(wb, sheet, title, cols, data, rowkey, dec, note=None):
    """Hoja con una fila por escenario (sin banda) — T8 costo unitario."""
    ws = wb.create_sheet(sheet)
    ws.sheet_view.showGridLines = False
    ncols = 1 + len(cols)
    ws.cell(row=1, column=1, value=title).font = F_TITLE
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    r = 3
    _hdr(ws, r, 1, "Escenario")
    for j, cname in enumerate(cols):
        _hdr(ws, r, 2 + j, cname)
    r += 1
    for sc in SCENARIOS:
        _lab(ws, r, 1, ALIAS[sc], bold=True, band=True)
        for j, cname in enumerate(cols):
            _val(ws, r, 2 + j, data[sc].get(rowkey, {}).get(cname, 0.0), dec)
        r += 1
    ws.cell(row=r + 1, column=1, value="Fuente: Elaboración propia").font = F_NOTE
    if note:
        ws.cell(row=r + 2, column=1, value=note).font = F_NOTE
    ws.column_dimensions["A"].width = 28
    for j in range(len(cols)):
        ws.column_dimensions[get_column_letter(2 + j)].width = 12.5
    return ws


def fams_present(tbl) -> list[str]:
    present = {fam for sc in SCENARIOS for fam, cols in tbl[sc].items()
               if any(abs(v) > 5e-4 for v in cols.values())}
    return [f for f in FAMILY_ORDER if f in present]


wb = Workbook()

# --- Índice ---
ws = wb.active
ws.title = "Índice"
ws.sheet_view.showGridLines = False
ws.cell(row=1, column=1, value="Tablas completas de resultados — nivel regional "
        "(todos los países agregados)").font = F_TITLE
ws.cell(row=3, column=1, value="Escenarios").font = F_BOLD
# Descripciones alineadas con scenario_transforms de Config_MOMF_T1_AB.yaml.
_FA = "mayor costo de combustible fósil (VariableCost MIN* ×1,70)"
_FB = "menor costo de combustible fósil (VariableCost MIN* ×0,52)"
_RB = "menor costo de inversión en renovables (solar ×0,60, eólica ×0,75, baterías ×0,50)"
_RA = "mayor costo de inversión en renovables (multiplicadores pendientes de calibrar)"
_scen_desc = {
    # --- familia OPT ---
    "BAC": "OPT — Referencia: expansión óptima de mínimo costo (con tope en la "
           "rampa de inversión de transmisión).",
    "BFA": f"Sensibilidad de OPT con {_FA}.",
    "BFB": f"Sensibilidad de OPT con {_FB}.",
    "BRA": f"Sensibilidad de OPT con {_RA}.",
    "BRB": f"Sensibilidad de OPT con {_RB}.",
    "BSR": "OPT sin repotenciación de líneas.",
    # --- familia ETT ---
    "ISR": "ETT — Transmisión restringida (tendencia histórica), generación de "
           "menor costo, sin repotenciación.",
    "IFA": f"Sensibilidad de ETT con {_FA}.",
    "IFB": f"Sensibilidad de ETT con {_FB}.",
    "INV": "ETT con repotenciación de líneas.",
    "IRA": f"Sensibilidad de ETT con {_RA}.",
    "IRB": f"Sensibilidad de ETT con {_RB}.",
    # --- otros ---
    "OPC": "PLAN — Expansión planificada (con tope en la rampa de inversión de "
           "transmisión).",
    "VSR": "ETT-GP — Transmisión restringida, generación planificada (PEGs), sin "
           "repotenciación.",
}
r = 4
for sc in SCENARIOS:
    ws.cell(row=r, column=1, value=ALIAS[sc]).font = F_BOLD
    ws.cell(row=r, column=2, value=_scen_desc[sc]).font = F_BASE
    r += 1
r += 1
ws.cell(row=r, column=1, value="Tablas").font = F_BOLD
_INDEX = [
    ("T1", "Capacidad instalada de generación por fuente [GW] — 2030/2040/2050"),
    ("T2", "Generación anual por fuente [TWh] — 2030/2040/2050"),
    ("T3", "Capacidad instalada de transmisión por tipo de línea [GW] — 2030/2040/2050"),
    ("T4", "Kilómetros de línea instalados por tipo [km] — por periodo 2026-2050"),
    ("T5", "Capacidad instalada de almacenamiento [GW] — 2030/2040/2050"),
    ("T6", "Inversión anual promedio por tipo [MUSD/año] — por periodo 2026-2050"),
    ("T7", "Costo total del sistema, promedio anual [MUSD/año] — por periodo 2026-2050"),
    ("T8", "Costo unitario del sistema [USD/MWh] — por periodo 2026-2050"),
    ("T9", "Consumo de combustibles fósiles [MBEP] — 2030/2040/2050"),
    ("Resumen", "Comparativa entre escenarios (2050 y acumulados 2026-2050)"),
]
for tag, desc in _INDEX:
    r += 1
    ws.cell(row=r, column=1, value=tag).font = F_BOLD
    ws.cell(row=r, column=2, value=desc).font = F_BASE
ws.column_dimensions["A"].width = 12
ws.column_dimensions["B"].width = 95

# --- Tablas ---
fam_sub = [("Renovable", [f for f in FAMILY_ORDER if f in RENEWABLE_FAMS]),
           ("No Renovable", [f for f in FAMILY_ORDER if f not in RENEWABLE_FAMS])]
write_banded(wb, "T1 Capacidad Generación",
             "Tabla 1. Capacidad instalada de generación por fuente (GW).",
             "Fuente", YEAR_COLS, cap_tbl, fams_present(cap_tbl), 2,
             subtotal_groups=fam_sub)
write_banded(wb, "T2 Generación Anual",
             "Tabla 2. Generación anual por fuente (TWh).",
             "Fuente", YEAR_COLS, gen_tbl, fams_present(gen_tbl), 2,
             subtotal_groups=fam_sub)
write_banded(wb, "T3 Capacidad Transmisión",
             "Tabla 3. Capacidad instalada de transmisión por tipo de línea (GW).",
             "Tipo de línea", YEAR_COLS, tr_tbl, TRANS_CATS, 2,
             totals="sum_first",
             note="Nota: no incluye interconectores internacionales puros (TRN*).")
write_banded(wb, "T4 Km de Línea",
             "Tabla 4. Kilómetros de línea instalados por tipo (km).",
             "Tipo de línea", PERIOD_COLS, km_tbl, KM_CATS, 0,
             note="Nota: km instalados dentro de cada periodo; repotenciadas ×1,25.")
write_banded(wb, "T5 Almacenamiento",
             "Tabla 5. Capacidad instalada de almacenamiento (GW).",
             "Tipo", YEAR_COLS, st_tbl, list(STO_TYPES.values()), 2)
write_banded(wb, "T6 Inversión",
             "Tabla 6. Inversión anual promedio por tipo (MUSD/año).",
             "Tipo", PERIOD_COLS, inv_tbl, INV_CATS, 1,
             note="Nota: promedio anual del periodo; CAPEX de líneas planificadas "
                  "(PWRTRN/RNWTRN) ajustado ÷1,2.")
write_banded(wb, "T7 Costos Sistema",
             "Tabla 7. Costo total del sistema, promedio anual (MUSD/año).",
             "Componente", PERIOD_COLS, cost_tbl, COST_CATS, 1,
             note="Nota: Energía no Suministrada valorada al VOLL de 1.500 USD/MWh; "
                  "backstop excluido de CAPEX/O&M/Combustible.")
write_scenario_rows(wb, "T8 Costo Unitario",
                    "Tabla 8. Costo unitario del sistema (USD/MWh).",
                    PERIOD_COLS, uc_tbl, "Costo unitario", 2,
                    note="Nota: (CAPEX anualizado con CRF al "
                         f"{DISCOUNT_RATE * 100:.2f}% y vida útil {ASSET_LIFETIME} años "
                         "+ O&M + Combustible) / generación total; promedio de los "
                         "ratios anuales del periodo; backstop excluido.")
write_banded(wb, "T9 Combustibles Fósiles",
             "Tabla 9. Consumo anual de combustibles fósiles "
             "(millones de bep, MBEP).",
             "Combustible", YEAR_COLS, fo_tbl, FOSSIL_CATS, 1,
             note="Nota: actividad anual de las tecnologías de extracción/importación "
                  "MIN* fósiles (incluye importadas MIN*INT; excluye uranio), "
                  "convertida de PJ a millones de barriles equivalentes de petróleo "
                  "(IEA: 1 tep = 41,868 GJ; 1 tep ≈ 7,33 bep -> 1 bep ≈ 5,7119 GJ).")

# --- Resumen comparativo entre escenarios ---
def tot_col(tbl, labels, col):
    return {sc: sum(tbl[sc].get(l, {}).get(col, 0.0) for l in labels)
            for sc in SCENARIOS}


_SUMMARY = [
    ("Capacidad instalada de generación 2050 [GW]",
     tot_col(cap_tbl, fams_present(cap_tbl), "2050"), 1),
    ("Generación anual 2050 [TWh]",
     tot_col(gen_tbl, fams_present(gen_tbl), "2050"), 1),
    ("Capacidad de transmisión 2050 [GW]",
     tot_col(tr_tbl, TRANS_CATS, "2050"), 1),
    ("Kilómetros de línea instalados 2026-2050 [km]", km_acc, 0),
    ("Capacidad de almacenamiento 2050 [GW]",
     tot_col(st_tbl, list(STO_TYPES.values()), "2050"), 1),
    ("Inversión de capital acumulada 2026-2050 [MUSD]", inv_acc, 0),
    ("Costo total del sistema acumulado 2026-2050 [MUSD]", cost_acc, 0),
    ("Costo unitario promedio 2026-2050 [USD/MWh]", uc_avg, 2),
    ("Consumo de combustibles fósiles acumulado 2026-2050 [MBEP]", fo_acc, 0),
]
ws = wb.create_sheet("Resumen")
ws.sheet_view.showGridLines = False
ws.cell(row=1, column=1,
        value="Resumen comparativo entre escenarios (nivel regional).").font = F_TITLE
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2 + len(SCENARIOS))
_hdr(ws, 3, 1, "Indicador")
for j, sc in enumerate(SCENARIOS):
    _hdr(ws, 3, 2 + j, ALIAS[sc])
for i, (name, vals, dec) in enumerate(_SUMMARY):
    _lab(ws, 4 + i, 1, name, bold=True, band=True)
    ws.cell(row=4 + i, column=1).alignment = AL_L
    for j, sc in enumerate(SCENARIOS):
        _val(ws, 4 + i, 2 + j, vals.get(sc, 0.0), dec)
ws.cell(row=5 + len(_SUMMARY), column=1,
        value="Fuente: Elaboración propia").font = F_NOTE
ws.column_dimensions["A"].width = 52
ws.row_dimensions[3].height = 32
for j in range(len(SCENARIOS)):
    ws.column_dimensions[get_column_letter(2 + j)].width = 14

_targets = [OUT_XLSX] + [OUT_XLSX.replace(".xlsx", f"_nuevo{'' if i == 1 else i}.xlsx")
                         for i in (1, 2, 3)]
out_written = None
for _target in _targets:
    try:
        wb.save(_target)
        out_written = _target
        break
    except PermissionError:
        print(f"AVISO: {os.path.basename(_target)} bloqueado (¿abierto en Excel?)")
if out_written is None:
    raise SystemExit("ERROR: todos los destinos están bloqueados; cierra Excel y re-corre.")
if out_written == OUT_XLSX:
    # se pudo escribir el canónico: las copias *_nuevo* de corridas anteriores
    # quedan obsoletas; se eliminan (o se avisa si Excel las tiene bloqueadas).
    for _stale in _targets[1:]:
        if os.path.exists(_stale):
            try:
                os.remove(_stale)
                print("eliminada copia obsoleta:", os.path.basename(_stale))
            except PermissionError:
                print(f"AVISO: {os.path.basename(_stale)} quedó OBSOLETA pero está "
                      f"bloqueada; usa {os.path.basename(OUT_XLSX)}")
print("OK ->", out_written, flush=True)

# --- resumen de control por consola ---
for name, vals, dec in _SUMMARY:
    line = "  ".join(f"{ALIAS[sc]}={vals.get(sc, 0.0):,.{dec}f}" for sc in SCENARIOS)
    print(f"{name}: {line}")
