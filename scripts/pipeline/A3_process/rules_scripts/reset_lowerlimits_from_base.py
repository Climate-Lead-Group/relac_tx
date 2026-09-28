# -*- coding: utf-8 -*-
"""Reset the LowerLimit floor for PWR techs in a scenario back to the Base scenario's values.

Overwrites TotalTechnologyAnnualActivityLowerLimit (years 2023-2050) for every PWR row in a
scenario's "Secondary Techs" sheet with the Base scenario's values for the same tech. The Base
scenario only carries floors in the base years (2023/2024) and is empty for 2025-2050, so this
both restores the 2023/2024 floor to the Base value and clears any 2025-2050 extension
previously written (e.g. by extend_lowerlimits_pwr.py or B1b's V3 fix). Projection.Mode
(col G) is also copied from Base so the row's mode matches the reset state.

Use this as a manual reset of the LowerLimit data across scenarios before re-running the A3
workflow, so that only the scenarios configured under `lowerlimit_scenarios` in lid_rule.yaml
get their floors extended/adjusted again.

Inputs:
    <input_dir>/A-O_Parametrization.xlsx (sheet "Secondary Techs") — target scenario
    <base_dir>/A-O_Parametrization.xlsx (sheet "Secondary Techs") — Base scenario reference;
        --base-dir defaults to P.BASE_SCENARIO_REF / "Base" (inputs/reference/NO BORRAR
        A1_Outputs - Escenario Base/Base)
Outputs:
    same target workbook, in place (a timestamped backup is written first)
    <input_dir>/reset_lowerlimits_from_base_changes_<ts>.json — per-cell audit (row, tech,
        year, prev, new)
Usage:
    python reset_lowerlimits_from_base.py --input-dir <scenario_dir> [--base-dir <base_dir>]
Notes:
    - Not invoked by A3_process.py — a manual, standalone tool run before re-running A3.
Authors:
    CLG (A3_process integration)
"""
from __future__ import annotations

import argparse
import json
import shutil
import sys
import time
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[3]))  # -> scripts/
from common import relac_paths as P

import openpyxl

SHEET_NAME = "Secondary Techs"
TARGET_PARAMETER = "TotalTechnologyAnnualActivityLowerLimit"
TECH_PREFIX = "PWR"
BASE_YEAR = 2023
END_YEAR = 2050

# Column layout for "Secondary Techs":
#   B=Tech, E=Parameter, G=Projection.Mode, I=2023, J=2024, ..., AJ=2050.
COL_TECH = 2
COL_PARAMETER = 5
COL_PROJ_MODE = 7
COL_YEAR_2023 = 9  # column for year `y` = COL_YEAR_2023 + (y - 2023)

PARAM_FILENAME = "A-O_Parametrization.xlsx"

DEFAULT_BASE_DIR = P.BASE_SCENARIO_REF / "Base"


def year_to_col(year: int) -> int:
    """Return the 1-based column index of `year` in "Secondary Techs" (2023 -> COL_YEAR_2023)."""
    return COL_YEAR_2023 + (year - 2023)


def _norm(val):
    """Normalize a cell value for comparison (blank string -> None)."""
    if isinstance(val, str) and val.strip() == "":
        return None
    return val


def load_base_rows(base_paramfile: Path) -> tuple[dict, dict]:
    """Return (year_values_by_tech, proj_mode_by_tech) from the Base xlsx.

    year_values_by_tech[tech] = [v2023, v2024, ..., v2050]  (28 values)
    """
    wb = openpyxl.load_workbook(base_paramfile, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        sys.exit(f"ERROR: sheet '{SHEET_NAME}' not found in {base_paramfile}")
    ws = wb[SHEET_NAME]
    values: dict[str, list] = {}
    proj_mode: dict[str, object] = {}
    for row in range(2, ws.max_row + 1):
        tech = ws.cell(row=row, column=COL_TECH).value
        parameter = ws.cell(row=row, column=COL_PARAMETER).value
        if not isinstance(tech, str) or not tech.startswith(TECH_PREFIX):
            continue
        if parameter != TARGET_PARAMETER:
            continue
        years = [
            _norm(ws.cell(row=row, column=year_to_col(y)).value)
            for y in range(BASE_YEAR, END_YEAR + 1)
        ]
        values[tech] = years
        proj_mode[tech] = ws.cell(row=row, column=COL_PROJ_MODE).value
    wb.close()
    return values, proj_mode


def run(input_dir: Path, base_dir: Path) -> dict:
    """Reset PWR LowerLimit cells (2023-2050) and Projection.Mode in `input_dir` to Base's values.

    Args:
        input_dir: target scenario folder containing A-O_Parametrization.xlsx to reset.
        base_dir: Base scenario folder containing the reference A-O_Parametrization.xlsx.

    Returns:
        The audit log dict (also written to disk as JSON), including per-row `changes` and
        `rows_missing_in_base` (techs present in the target but absent from Base, left
        untouched).

    Side effects:
        Writes a timestamped backup of the target workbook, saves the workbook in place, and
        writes `reset_lowerlimits_from_base_changes_<ts>.json` next to it.
    """
    paramfile = input_dir / PARAM_FILENAME
    if not paramfile.is_file():
        sys.exit(f"ERROR: {PARAM_FILENAME} not found in {input_dir}")
    base_paramfile = base_dir / PARAM_FILENAME
    if not base_paramfile.is_file():
        sys.exit(f"ERROR: base {PARAM_FILENAME} not found in {base_dir}")

    base_values, base_proj_mode = load_base_rows(base_paramfile)
    if not base_values:
        sys.exit(f"ERROR: no {TARGET_PARAMETER} {TECH_PREFIX} rows found in "
                 f"base file {base_paramfile}")

    wb = openpyxl.load_workbook(paramfile)
    if SHEET_NAME not in wb.sheetnames:
        sys.exit(f"ERROR: sheet '{SHEET_NAME}' not found in {paramfile}")
    ws = wb[SHEET_NAME]

    changes: list[dict] = []
    rows_matched = 0
    rows_touched = 0
    rows_missing_in_base = []

    for row in range(2, ws.max_row + 1):
        tech = ws.cell(row=row, column=COL_TECH).value
        parameter = ws.cell(row=row, column=COL_PARAMETER).value
        if not isinstance(tech, str) or not tech.startswith(TECH_PREFIX):
            continue
        if parameter != TARGET_PARAMETER:
            continue
        rows_matched += 1
        if tech not in base_values:
            rows_missing_in_base.append(tech)
            continue

        base_years = base_values[tech]
        row_changes: list[dict] = []
        for idx, year in enumerate(range(BASE_YEAR, END_YEAR + 1)):
            col = year_to_col(year)
            prev = _norm(ws.cell(row=row, column=col).value)
            new = base_years[idx]
            if prev == new:
                continue
            ws.cell(row=row, column=col).value = new
            row_changes.append({"year": year, "prev": prev, "new": new})

        # Always realign Projection.Mode to base for this row.
        prev_pm = ws.cell(row=row, column=COL_PROJ_MODE).value
        new_pm = base_proj_mode.get(tech)
        pm_changed = (_norm(prev_pm) != _norm(new_pm))
        if pm_changed:
            ws.cell(row=row, column=COL_PROJ_MODE).value = new_pm

        if row_changes or pm_changed:
            rows_touched += 1
            changes.append({
                "row": row,
                "tech": tech,
                "proj_mode_prev": prev_pm if pm_changed else None,
                "proj_mode_new": new_pm if pm_changed else None,
                "cells": row_changes,
            })

    stamp = time.strftime("%Y%m%d_%H%M%S")
    backup = paramfile.with_name(
        f"{paramfile.stem}.backup_pre_reset_ll_{stamp}{paramfile.suffix}"
    )
    shutil.copy2(paramfile, backup)
    wb.save(paramfile)

    log_path = input_dir / f"reset_lowerlimits_from_base_changes_{stamp}.json"
    log = {
        "input_dir": str(input_dir),
        "base_dir": str(base_dir),
        "paramfile": str(paramfile),
        "base_paramfile": str(base_paramfile),
        "backup": str(backup),
        "sheet": SHEET_NAME,
        "parameter": TARGET_PARAMETER,
        "tech_prefix": TECH_PREFIX,
        "base_year": BASE_YEAR,
        "end_year": END_YEAR,
        "rows_matched": rows_matched,
        "rows_touched": rows_touched,
        "rows_missing_in_base": rows_missing_in_base,
        "total_cell_writes": sum(len(c["cells"]) for c in changes),
        "changes": changes,
    }
    log_path.write_text(json.dumps(log, indent=2, default=str))
    log["log_path"] = str(log_path)
    return log


def print_summary(log: dict) -> None:
    """Print a formatted summary of the `run` audit log dict to stdout."""
    print()
    print("=" * 70)
    print(f"  RESET LowerLimits from Base for {log['tech_prefix']} technologies")
    print("=" * 70)
    print(f"  paramfile            : {log['paramfile']}")
    print(f"  base_paramfile       : {log['base_paramfile']}")
    print(f"  backup               : {log['backup']}")
    print(f"  sheet                : {log['sheet']}")
    print(f"  parameter            : {log['parameter']}")
    print(f"  rows_matched         : {log['rows_matched']}")
    print(f"  rows_touched         : {log['rows_touched']}")
    print(f"  total_cell_writes    : {log['total_cell_writes']}")
    if log["rows_missing_in_base"]:
        print(f"  [WARN] {len(log['rows_missing_in_base'])} tech(s) absent in "
              f"base, left untouched: {log['rows_missing_in_base'][:10]}")
    print(f"  log                  : {log['log_path']}")


def main() -> int:
    """CLI entry point: parse args, run the LowerLimit reset, and print a summary.

    Returns:
        0 on success.
    """
    parser = argparse.ArgumentParser(
        description=__doc__,
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument(
        "--input-dir",
        type=Path,
        required=True,
        help="Scenario folder containing A-O_Parametrization.xlsx to reset.",
    )
    parser.add_argument(
        "--base-dir",
        type=Path,
        default=DEFAULT_BASE_DIR,
        help=f"Base scenario folder containing the reference "
             f"A-O_Parametrization.xlsx (default: {DEFAULT_BASE_DIR}).",
    )
    args = parser.parse_args()

    log = run(args.input_dir, args.base_dir)
    print_summary(log)
    return 0


if __name__ == "__main__":
    sys.exit(main())
