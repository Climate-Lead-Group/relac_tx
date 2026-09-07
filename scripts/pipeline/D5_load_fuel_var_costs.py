"""
Load Fuel VariableCosts into A-O_Parametrization.xlsx for every scenario.

Author: Climate Lead Group, Andrey Salazar-Vargas

Source: A1_Outputs/<scenario>/LAC_Fuel_Price_Projections_2026_v2_audit.xlsx
        - sheets matching pattern "Country_<ISO3>" (USD/L per fuel/year/scenario).
Target: A1_Outputs/<scenario>/A-O_Parametrization.xlsx, sheet "VariableCost".
        Rows whose TECHNOLOGY equals "MIN" + <fuel_code> + <ISO3> are updated.

Unit conversion (USD/L -> M$/PJ):
    1 USD/L / (GJ/L) = 1 USD/GJ = 1 M$/PJ
    so M$/PJ = price_USD_per_L / energy_content_GJ_per_L

Idempotent: rerunning with the same price_scenario yields the same workbook state.

Usage:
    python t1_confection/D5_load_fuel_var_costs.py
    python t1_confection/D5_load_fuel_var_costs.py --price-scenario HIGH
    python t1_confection/D5_load_fuel_var_costs.py --base A1_Outputs --price-scenario LOW
"""
import argparse
import re
import shutil
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import openpyxl


SOURCE_FILENAME = "LAC_Fuel_Price_Projections_2026_v2_audit.xlsx"
TARGET_FILENAME = "A-O_Parametrization.xlsx"
TARGET_SHEET = "VariableCost"
SCENARIOS = ("REF", "HIGH", "LOW")
COUNTRY_SHEET_RE = re.compile(r"^Country_([A-Z]{3})$")

# Column layout of the VariableCost sheet (1-based).
COL_TECH = 3          # C: Tech (e.g. MINOILHND)
COL_PROJ_MODE = 9     # I: Projection.Mode

# --- Configuration: extend these dicts to add new fuels --------------------

# Map source fuel header (row 4 in Country_XXX) -> OSeMOSYS FUEL code (3 chars).
# The model tech to update will be MIN + FUEL_CODE + ISO3.
FUEL_MAPPING: Dict[str, str] = {
    "Fuel Oil": "OIL",
    "Diesel": "PET",
}

# Country override: reassign a source ISO3 to a different model ISO3 when the
# country is not represented in the model but should be proxied by another.
# Example: Jamaica is in the source file but not in the model; use its data as
# a proxy for Barbados.
COUNTRY_OVERRIDE: Dict[str, str] = {
    "JAM": "BRB",
}

# Energy content (LHV) per source fuel, in GJ/L. Values from Anchors_Parameters
# in the source workbook (IPCC defaults, S026). Used to convert USD/L -> M$/PJ.
ENERGY_CONTENT_GJ_PER_L: Dict[str, float] = {
    "Fuel Oil": 0.0395,
    "Diesel": 0.0358,
    "Gasoline Reg": 0.0322,
    "Gasoline Sup": 0.0322,
    "LPG": 0.0259,
    "Kerosene": 0.0344,
    "Jet A-1": 0.0344,
}

# ---------------------------------------------------------------------------


def discover_scenarios(base_dir: Path) -> List[Path]:
    """Return scenario folders under ``base_dir`` that contain the source Excel.

    Scenarios without the source file are silently skipped here; the caller
    is expected to log which ones are processed vs omitted.
    """
    if not base_dir.exists():
        return []
    return sorted(
        p for p in base_dir.iterdir()
        if p.is_dir() and (p / SOURCE_FILENAME).exists()
    )


def build_source_index(source_path: Path, price_scenario: str) -> Dict[str, Dict[str, Dict[int, float]]]:
    """Read fuel prices from every ``Country_<ISO3>`` sheet of the source file.

    Returns a nested dict ``{ISO3: {source_fuel_name: {year: price_USD_per_L}}}``
    for the requested ``price_scenario`` (REF/HIGH/LOW).
    """
    if price_scenario not in SCENARIOS:
        raise ValueError(f"price_scenario must be one of {SCENARIOS}, got {price_scenario!r}")

    wb = openpyxl.load_workbook(source_path, data_only=True, read_only=True)
    data: Dict[str, Dict[str, Dict[int, float]]] = {}

    for sheet_name in wb.sheetnames:
        m = COUNTRY_SHEET_RE.match(sheet_name)
        if not m:
            continue
        iso3 = m.group(1)
        ws = wb[sheet_name]

        # Row 4 carries fuel labels (merged across 3 columns: REF, HIGH, LOW).
        # Row 5 carries the scenario tag under each fuel column.
        # Forward-fill row 4 because only the first column of each fuel block is
        # populated, then pair every column with its scenario from row 5.
        row4 = [c.value for c in ws[4]]
        row5 = [c.value for c in ws[5]]
        fuel_per_col: Dict[int, str] = {}
        current_fuel: Optional[str] = None
        for col_idx, val in enumerate(row4, start=1):
            if val is not None and str(val).strip() and str(val).strip().lower() != "year":
                current_fuel = str(val).strip()
            if current_fuel is None:
                continue
            scen = row5[col_idx - 1] if col_idx - 1 < len(row5) else None
            if scen is not None and str(scen).strip().upper() == price_scenario:
                fuel_per_col[col_idx] = current_fuel

        if not fuel_per_col:
            continue

        country_data: Dict[str, Dict[int, float]] = {}
        for row in ws.iter_rows(min_row=6, values_only=True):
            year_cell = row[0]
            if not isinstance(year_cell, int):
                continue
            for col_idx, fuel_name in fuel_per_col.items():
                value = row[col_idx - 1] if col_idx - 1 < len(row) else None
                if isinstance(value, (int, float)):
                    country_data.setdefault(fuel_name, {})[year_cell] = float(value)

        if country_data:
            data[iso3] = country_data

    wb.close()
    return data


def build_year_col_map(ws) -> Dict[int, int]:
    """Map year (int) -> column index (1-based) from the VariableCost header row."""
    year_cols: Dict[int, int] = {}
    for col_idx, cell in enumerate(ws[1], start=1):
        val = cell.value
        if isinstance(val, int) and 2000 <= val <= 2100:
            year_cols[val] = col_idx
        elif isinstance(val, str) and val.isdigit():
            yr = int(val)
            if 2000 <= yr <= 2100:
                year_cols[yr] = col_idx
    return year_cols


def build_tech_row_index(ws) -> Dict[str, int]:
    """Map TECHNOLOGY code (column C) -> row index (1-based)."""
    index: Dict[str, int] = {}
    for row_idx in range(2, ws.max_row + 1):
        tech = ws.cell(row=row_idx, column=COL_TECH).value
        if isinstance(tech, str) and tech:
            index[tech.strip()] = row_idx
    return index


def model_iso3_set(tech_index: Dict[str, int]) -> set:
    """ISO3 codes that appear in MIN<fuel><iso3> rows of the model."""
    iso3 = set()
    for tech in tech_index:
        if tech.startswith("MIN") and len(tech) == 9:
            iso3.add(tech[6:9])
    return iso3


def convert_price(price_usd_per_l: float, source_fuel: str) -> float:
    """Convert a USD/L price to M$/PJ using the fuel's energy content."""
    gj_per_l = ENERGY_CONTENT_GJ_PER_L[source_fuel]
    return price_usd_per_l / gj_per_l


def update_variable_cost(
    target_path: Path,
    source_data: Dict[str, Dict[str, Dict[int, float]]],
    scenario_label: str,
) -> Tuple[int, int]:
    """Apply fuel price updates to the VariableCost sheet of ``target_path``.

    ``source_data`` is the structure returned by :func:`build_source_index`.
    Returns ``(rows_updated, cells_updated)`` for reporting.
    """
    wb = openpyxl.load_workbook(target_path)
    if TARGET_SHEET not in wb.sheetnames:
        wb.close()
        raise KeyError(f"Sheet {TARGET_SHEET!r} not found in {target_path}")
    ws = wb[TARGET_SHEET]

    year_cols = build_year_col_map(ws)
    tech_index = build_tech_row_index(ws)
    model_iso3 = model_iso3_set(tech_index)
    model_years = set(year_cols.keys())

    src_iso3 = set(source_data.keys())

    # Resolve source ISO3 -> model ISO3 via COUNTRY_OVERRIDE. The override is
    # applied first so that countries absent from the model but present in
    # COUNTRY_OVERRIDE can still be written under their proxy ISO3.
    resolved: Dict[str, str] = {}
    for src in sorted(src_iso3):
        target = COUNTRY_OVERRIDE.get(src, src)
        if target == src and src not in model_iso3:
            print(f"    [WARN] Pais {src} no esta en el modelo, se omite.")
            continue
        if target != src:
            if target not in model_iso3:
                print(
                    f"    [WARN] Override {src}->{target}: destino tampoco existe en el modelo, se omite."
                )
                continue
            if target in src_iso3:
                print(
                    f"    [WARN] Override {src}->{target}: el destino {target} tambien esta en el Excel fuente; "
                    f"se priorizan los datos directos de {target} y se omite {src}."
                )
                continue
            print(f"    [INFO] Override {src}->{target}: usando datos de {src} como proxy de {target}.")
        resolved[src] = target

    rows_updated = 0
    cells_updated = 0
    skipped_years_logged = False
    skipped_fuels: set = set()

    for iso3 in sorted(resolved):
        country_block = source_data[iso3]
        target_iso3 = resolved[iso3]

        for source_fuel, year_prices in country_block.items():
            fuel_code = FUEL_MAPPING.get(source_fuel)
            if fuel_code is None:
                skipped_fuels.add(source_fuel)
                continue

            tech_code = f"MIN{fuel_code}{target_iso3}"
            row_idx = tech_index.get(tech_code)
            if row_idx is None:
                print(f"    [WARN] {tech_code} no esta en VariableCost, se omite.")
                continue

            src_years = set(year_prices.keys())
            usable_years = sorted(src_years & model_years)
            dropped_years = sorted(src_years - model_years)
            if dropped_years and not skipped_years_logged:
                print(
                    f"    [INFO] Anos del Excel fuera del modelo (se omiten): "
                    f"{dropped_years[0]}..{dropped_years[-1]}"
                )
                skipped_years_logged = True

            if not usable_years:
                continue

            row_cells_written = 0
            for year in usable_years:
                col_idx = year_cols[year]
                new_val = convert_price(year_prices[year], source_fuel)
                ws.cell(row=row_idx, column=col_idx).value = new_val
                row_cells_written += 1

            ws.cell(row=row_idx, column=COL_PROJ_MODE).value = "User defined"

            if row_cells_written:
                rows_updated += 1
                cells_updated += row_cells_written
                origin = f"{source_fuel} {scenario_label}"
                if target_iso3 != iso3:
                    origin += f" (proxy {iso3}->{target_iso3})"
                print(
                    f"    {tech_code}: {row_cells_written} celdas "
                    f"({usable_years[0]}-{usable_years[-1]}) <- {origin}"
                )

    for fuel in sorted(skipped_fuels):
        print(f"    [INFO] Fuel sin mapeo, se omite: {fuel!r}")

    wb.save(target_path)
    wb.close()
    return rows_updated, cells_updated


def process_scenario(scenario_dir: Path, price_scenario: str, do_backup: bool) -> Tuple[int, int]:
    """Process a single scenario folder. Returns ``(rows_updated, cells_updated)``."""
    source_path = scenario_dir / SOURCE_FILENAME
    target_path = scenario_dir / TARGET_FILENAME

    if not target_path.exists():
        print(f"  [SKIP] {scenario_dir.name}: no existe {TARGET_FILENAME}")
        return (0, 0)

    print(f"  Scenario: {scenario_dir.name}")
    print(f"    Source: {source_path.name}")
    print(f"    Target: {target_path.name}  (price_scenario={price_scenario})")

    source_data = build_source_index(source_path, price_scenario)
    if not source_data:
        print(f"    [WARN] No se hallaron hojas Country_XXX en {source_path.name}")
        return (0, 0)

    print(f"    Paises en Excel fuente: {sorted(source_data.keys())}")

    if do_backup:
        ts = datetime.now().strftime("%Y%m%d-%H%M%S")
        backup_path = target_path.with_suffix(f".backup-fuelvc-{ts}.xlsx")
        shutil.copy2(target_path, backup_path)
        print(f"    Backup: {backup_path.name}")

    return update_variable_cost(target_path, source_data, price_scenario)


def run(base_dir: Path, price_scenario: str = "REF", do_backup: bool = True) -> None:
    """Iterate every scenario folder under ``base_dir`` and update fuel costs.

    A scenario folder is processed only if it contains :data:`SOURCE_FILENAME`.
    """
    print(f"Base dir: {base_dir}")
    print(f"Price scenario: {price_scenario}")
    print(f"Fuel mapping: {FUEL_MAPPING}")

    candidates = sorted(p for p in base_dir.iterdir() if p.is_dir())
    if not candidates:
        print(f"[WARN] No hay subcarpetas en {base_dir}")
        return

    processed = []
    omitted = []
    total_rows = 0
    total_cells = 0

    for scen_dir in candidates:
        if not (scen_dir / SOURCE_FILENAME).exists():
            omitted.append(scen_dir.name)
            continue
        rows, cells = process_scenario(scen_dir, price_scenario, do_backup)
        processed.append((scen_dir.name, rows, cells))
        total_rows += rows
        total_cells += cells

    print("\n=== SUMMARY ===")
    if omitted:
        print(f"Escenarios omitidos (sin {SOURCE_FILENAME}): {omitted}")
    for name, rows, cells in processed:
        print(f"  {name}: {rows} filas, {cells} celdas actualizadas")
    print(f"TOTAL: {total_rows} filas, {total_cells} celdas actualizadas")


def parse_args() -> argparse.Namespace:
    """Parse CLI arguments."""
    default_base = Path(__file__).parent / "A1_Outputs"
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument(
        "--base",
        type=Path,
        default=default_base,
        help=f"Base dir with scenario subfolders (default: {default_base})",
    )
    parser.add_argument(
        "--price-scenario",
        choices=SCENARIOS,
        default="REF",
        help="Source price scenario selector (default: REF)",
    )
    parser.add_argument(
        "--no-backup",
        action="store_true",
        help="Skip the .backup-fuelvc-<timestamp>.xlsx copy.",
    )
    return parser.parse_args()


def main() -> None:
    """CLI entry point."""
    args = parse_args()
    run(args.base, price_scenario=args.price_scenario, do_backup=not args.no_backup)


if __name__ == "__main__":
    main()
