# -*- coding: utf-8 -*-
"""
Z_AUX_update_transmission_iar.py

Actualiza el Input Activity Ratio (IAR) de las tecnologias de transmision por
pais en los archivos AR de cada escenario, preservando el formato de los .xlsx.

Las tecnologias de transmision (6 prefijos x 19 paises = 114 techs) viven en la
hoja "Demand Techs" de ambos archivos:
    A1_Outputs/A1_Outputs_<SCEN>/A-O_AR_Model_Base_Year.xlsx
    A1_Outputs/A1_Outputs_<SCEN>/A-O_AR_Projections.xlsx

Cada codigo es <PREFIJO 6 chars><PAIS 3 chars>XX  (p.ej. RNWTRNARGXX).
Mapeo prefijo -> categoria de IAR (η = 1 - α·d ; IAR = 1/η):
    No Renovable : PWRTRN, TRNNLI -> columna "IAR No-Renovable"
    Renovable    : RNWTRN, RNWNLI -> columna "IAR Renovable"
    Repotenciada : RNWRPO, TRNRPO -> columna "IAR Repotenciada"

Que se edita:
    Base_Year   : columna D "Value.Fuel.I" (1 fila por tech, mode 1).
    Projections : columnas de año 2023..2050 (I..AJ) SOLO en las filas con
                  Direction == "Input" (las filas "Output" quedan intactas).

Patron de edicion (openpyxl, save() preserva formato) siguiendo
reset_lowerlimits_from_base.py / extend_lowerlimits_pwr.py.

@author: Climate Lead Group
"""

import argparse
import shutil
import sys
from pathlib import Path

import openpyxl

# ---------------------------------------------------------------------------
# Datos: IAR por pais (de la Tabla Principal del requerimiento)
#   ren   = IAR Renovable      (RNWTRN, RNWNLI)
#   noren = IAR No-Renovable   (PWRTRN, TRNNLI)
#   repo  = IAR Repotenciada   (RNWRPO, TRNRPO)
# ---------------------------------------------------------------------------
IAR = {
    "ARG": {"ren": 1.042, "noren": 1.012, "repo": 1.042},
    "BOL": {"ren": 1.031, "noren": 1.020, "repo": 1.031},
    "BRA": {"ren": 1.050, "noren": 1.012, "repo": 1.050},
    "BRB": {"ren": 1.006, "noren": 1.005, "repo": 1.006},
    "CHL": {"ren": 1.037, "noren": 1.010, "repo": 1.037},
    "COL": {"ren": 1.033, "noren": 1.008, "repo": 1.033},
    "CRI": {"ren": 1.031, "noren": 1.010, "repo": 1.031},
    "DOM": {"ren": 1.033, "noren": 1.016, "repo": 1.033},
    "ECU": {"ren": 1.034, "noren": 1.020, "repo": 1.034},
    "GTM": {"ren": 1.026, "noren": 1.013, "repo": 1.026},
    "HND": {"ren": 1.026, "noren": 1.015, "repo": 1.026},
    "HTI": {"ren": 1.018, "noren": 1.009, "repo": 1.018},
    "MEX": {"ren": 1.044, "noren": 1.025, "repo": 1.044},
    "NIC": {"ren": 1.020, "noren": 1.010, "repo": 1.020},
    "PAN": {"ren": 1.039, "noren": 1.010, "repo": 1.039},
    "PER": {"ren": 1.029, "noren": 1.006, "repo": 1.029},
    "PRY": {"ren": 1.039, "noren": 1.013, "repo": 1.039},
    "SLV": {"ren": 1.015, "noren": 1.006, "repo": 1.015},
    "URY": {"ren": 1.016, "noren": 1.003, "repo": 1.016},
}

PREFIX_TO_KIND = {
    "PWRTRN": "noren",
    "TRNNLI": "noren",
    "RNWTRN": "ren",
    "RNWNLI": "ren",
    "RNWRPO": "repo",
    "TRNRPO": "repo",
}

SHEET = "Demand Techs"
SCENARIOS = ["BAU", "INV", "OPT", "VGB"]

SCRIPT_DIR = Path(__file__).resolve().parent
A1_BASE = SCRIPT_DIR / "A1_Outputs"

BASE_YEAR_FILE = "A-O_AR_Model_Base_Year.xlsx"
PROJECTIONS_FILE = "A-O_AR_Projections.xlsx"

YEAR_START, YEAR_END = 2023, 2050

# Total esperado de techs de transmision por archivo: 6 prefijos x 19 paises.
EXPECTED_TECHS = len(PREFIX_TO_KIND) * len(IAR)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def iar_for(tech):
    """Devuelve el IAR para un codigo de tech de transmision, o None si no aplica.

    tech = <PREFIJO 6 chars><PAIS 3 chars>XX  (p.ej. 'RNWTRNARGXX')
    """
    if not isinstance(tech, str) or len(tech) < 9:
        return None
    prefix = tech[:6]
    kind = PREFIX_TO_KIND.get(prefix)
    if kind is None:
        return None
    country = tech[6:9]
    row = IAR.get(country)
    if row is None:
        # Tech de transmision con pais fuera de la tabla -> avisar arriba.
        return None
    return row[kind]


def find_cols(ws):
    """Lee la fila 1 (headers) y devuelve un dict header->indice (1-based).

    Tambien agrega 'YEARS' -> {anio: col} para los headers numericos.
    """
    cols = {}
    years = {}
    for idx, cell in enumerate(ws[1], start=1):
        val = cell.value
        if val is None:
            continue
        key = str(val).strip()
        cols[key] = idx
        # Headers de año pueden venir como int o str numerico.
        try:
            y = int(float(key))
            if YEAR_START <= y <= YEAR_END:
                years[y] = idx
        except (ValueError, TypeError):
            pass
    cols["YEARS"] = years
    return cols


def require(cols, name, path):
    if name not in cols:
        raise KeyError(f"Columna '{name}' no encontrada en {path} (hoja '{SHEET}')")
    return cols[name]


# ---------------------------------------------------------------------------
# Actualizadores
# ---------------------------------------------------------------------------
def update_base_year(path, dry_run, make_backup):
    """Escribe Value.Fuel.I (col 'Value.Fuel.I') = IAR en cada fila de transmision."""
    wb = openpyxl.load_workbook(path)
    if SHEET not in wb.sheetnames:
        wb.close()
        raise ValueError(f"Hoja '{SHEET}' no existe en {path}")
    ws = wb[SHEET]
    cols = find_cols(ws)
    c_tech = require(cols, "Tech", path)
    c_val = require(cols, "Value.Fuel.I", path)

    changed = 0
    seen = set()
    unmapped = []
    for row in range(2, ws.max_row + 1):
        tech = ws.cell(row=row, column=c_tech).value
        if not isinstance(tech, str):
            continue
        if tech[:6] not in PREFIX_TO_KIND:
            continue
        val = iar_for(tech)
        if val is None:
            unmapped.append(tech)
            continue
        seen.add(tech)
        if not dry_run:
            ws.cell(row=row, column=c_val).value = val
        changed += 1

    _report_file("Base_Year", path, changed, len(seen), unmapped)

    if not dry_run:
        if make_backup:
            shutil.copy2(path, str(path) + ".bak")
        wb.save(path)
    wb.close()
    return changed


def update_projections(path, dry_run, make_backup):
    """Escribe IAR en 2023..2050 SOLO en filas de transmision con Direction='Input'."""
    wb = openpyxl.load_workbook(path)
    if SHEET not in wb.sheetnames:
        wb.close()
        raise ValueError(f"Hoja '{SHEET}' no existe en {path}")
    ws = wb[SHEET]
    cols = find_cols(ws)
    c_tech = require(cols, "Tech", path)
    c_dir = require(cols, "Direction", path)
    years = cols["YEARS"]
    missing_years = [y for y in range(YEAR_START, YEAR_END + 1) if y not in years]
    if missing_years:
        wb.close()
        raise ValueError(f"Faltan columnas de año {missing_years} en {path}")

    changed_cells = 0
    input_rows = set()
    unmapped = []
    for row in range(2, ws.max_row + 1):
        tech = ws.cell(row=row, column=c_tech).value
        if not isinstance(tech, str):
            continue
        if tech[:6] not in PREFIX_TO_KIND:
            continue
        direction = ws.cell(row=row, column=c_dir).value
        if not (isinstance(direction, str) and direction.strip().lower() == "input"):
            continue
        val = iar_for(tech)
        if val is None:
            unmapped.append(tech)
            continue
        input_rows.add(tech)
        if not dry_run:
            for y in range(YEAR_START, YEAR_END + 1):
                ws.cell(row=row, column=years[y]).value = val
        changed_cells += (YEAR_END - YEAR_START + 1)

    _report_file("Projections", path, changed_cells, len(input_rows), unmapped,
                 unit="celdas", techs_label="filas Input")

    if not dry_run:
        if make_backup:
            shutil.copy2(path, str(path) + ".bak")
        wb.save(path)
    wb.close()
    return changed_cells


def _report_file(tag, path, changed, n_techs, unmapped, unit="celdas",
                 techs_label="techs"):
    print(f"  [{tag}] {Path(path).name}: {changed} {unit} en {n_techs} {techs_label}", end="")
    if n_techs != EXPECTED_TECHS:
        print(f"  !! esperaba {EXPECTED_TECHS} {techs_label}", end="")
    if unmapped:
        uniq = sorted(set(unmapped))
        print(f"  !! sin mapear ({len(uniq)}): {uniq[:5]}{'...' if len(uniq) > 5 else ''}", end="")
    print()


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(
        description="Actualiza el IAR de tecnologias de transmision en los archivos AR.")
    ap.add_argument("--scenario", action="append", choices=SCENARIOS,
                    help="Escenario(s) a procesar. Repetible. Default: todos.")
    ap.add_argument("--dry-run", action="store_true",
                    help="No escribe: solo reporta lo que cambiaria.")
    ap.add_argument("--no-backup", action="store_true",
                    help="No crear copia .bak antes de escribir.")
    args = ap.parse_args()

    scenarios = args.scenario or SCENARIOS
    make_backup = not args.no_backup

    if args.dry_run:
        print(">>> DRY-RUN: no se escribira ningun archivo.\n")

    total = 0
    errors = 0
    for scen in scenarios:
        folder = A1_BASE / f"A1_Outputs_{scen}"
        print(f"Escenario {scen}  ({folder})")
        for fname, fn in ((BASE_YEAR_FILE, update_base_year),
                          (PROJECTIONS_FILE, update_projections)):
            path = folder / fname
            if not path.exists():
                print(f"  !! NO existe: {path}")
                errors += 1
                continue
            try:
                total += fn(path, args.dry_run, make_backup)
            except Exception as exc:  # noqa: BLE001
                print(f"  !! ERROR en {fname}: {exc}")
                errors += 1
        print()

    print(f"Total {'(simuladas) ' if args.dry_run else ''}celdas: {total}")
    if errors:
        print(f"Hubo {errors} error(es).")
        sys.exit(1)


if __name__ == "__main__":
    main()
