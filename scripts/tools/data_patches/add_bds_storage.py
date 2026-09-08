#!/usr/bin/env python3
"""
add_bds_storage.py — Clona la familia de almacenamiento SDS como BDS sobre
el nodo no renovable ELC<ISO3>XX02.

Spec: docs/superpowers/specs/2026-07-23-bds-storage-clone-design.md

Transformación por país C:
    PWRSDS{C}XX  -> PWRBDS{C}XX      (tecnología, lado potencia)
    SDS{C}XX01   -> BDS{C}XX01       (instalación, lado energía)
    ELC{C}XX00   -> ELC{C}XX02       (nodo eléctrico)

Idempotente por reemplazo: borra toda fila PWRBDS*/BDS* de las hojas
objetivo y reinserta según COUNTRIES. Re-correr converge al estado de la
lista (sirve para agregar y quitar países).

WATCHOUTS:
  * Correr DESPUÉS de A3 (A3 regenera los A-O y borra BDS) y ANTES de B1.
    Cadena: A3 -> add_bds_storage.py -> B1 -> B2.
  * Layout nuevo (2026-09): rutas relativas a inputs/ (A1_Outputs/, A2_Extra_Inputs/,
    config/) via scripts/common/relac_paths.py. YA APLICADO a inputs/ (ver README.md).
  * D2_update_secondary_techs es seguro (edita por match de Tech).
    B1b es seguro (los parámetros que valida están vacíos en PWRBDS).
  * sync_historical_from_bau exige orden de filas idéntico entre
    escenarios: este script itera sorted(COUNTRIES) SIEMPRE.
  * Los residuales y el piso de inversión BRA-2031 NO se copian (BDS nace
    sin stock).
  * Restore: copiar de vuelta los archivos desde _bds_backups/<timestamp>/.
"""
from __future__ import annotations

import shutil
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))  # -> scripts/
from common import relac_paths as P  # noqa: E402

# =====================================================================
# CONFIGURACIÓN — editar aquí antes de correr
# =====================================================================
COUNTRIES = [
    "ARG", "BOL", "BRA", "BRB", "CHL", "COL", "CRI", "DOM", "ECU",
    "GTM", "HND", "HTI", "MEX", "NIC", "PAN", "PER", "PRY", "SLV", "URY",
]
DRY_RUN = False    # True = imprime el plan de cambios, no escribe nada
BACKUP = True      # copia archivos a _bds_backups/<timestamp>/ antes de tocar
BASE_DIR = P.INPUTS   # raíz inputs/ del layout nuevo (el test la sobreescribe)

# =====================================================================
# CONSTANTES (no editar)
# =====================================================================
SRC_TECH_PREFIX = "PWRSDS"
DST_TECH_PREFIX = "PWRBDS"
SRC_STO_PREFIX = "SDS"
DST_STO_PREFIX = "BDS"

TECH_NAME_SRC = "Short duration storage (Power generator)"
TECH_NAME_DST = "Short duration storage (Power generator, non-renewable)"
STO_NAME_SRC = "Short duration storage "
STO_NAME_DST = "Short duration storage non-renewable "

# Parámetros que nacen en cero en BDS (decisión D3 del spec)
BLANK_PARAMS = {"ResidualCapacity", "TotalAnnualMinCapacityInvestment"}

A1_GLOB = "A1_Outputs/A1_Outputs_*"
XTRA_PATH = "A2_Extra_Inputs/A-Xtra_Storage.xlsx"
CONFIG_A = "config/Config_MOMF_T1_A.yaml"
CONFIG_AB = "config/Config_MOMF_T1_AB.yaml"
BACKUP_DIR = "_bds_backups"

BY_FILE = "A-O_AR_Model_Base_Year.xlsx"
PR_FILE = "A-O_AR_Projections.xlsx"
PM_FILE = "A-O_Parametrization.xlsx"


# =====================================================================
# Transforms puras
# =====================================================================
def src_tech(c: str) -> str:
    return f"{SRC_TECH_PREFIX}{c}XX"


def dst_tech(c: str) -> str:
    return f"{DST_TECH_PREFIX}{c}XX"


def src_storage(c: str) -> str:
    return f"{SRC_STO_PREFIX}{c}XX01"


def dst_storage(c: str) -> str:
    return f"{DST_STO_PREFIX}{c}XX01"


def src_fuel(c: str) -> str:
    return f"ELC{c}XX00"


def dst_fuel(c: str) -> str:
    return f"ELC{c}XX02"


def transform_text(v, c: str):
    """Aplica los renames SDS->BDS a un valor de celda (solo strings).

    NO transforma Fuel.Name (los escritores lo fijan explícitamente con el
    nombre real del nodo destino leído de Base_Year).
    """
    if not isinstance(v, str):
        return v
    v = v.replace(src_tech(c), dst_tech(c))
    v = v.replace(src_storage(c), dst_storage(c))
    v = v.replace(src_fuel(c), dst_fuel(c))
    if TECH_NAME_SRC in v:
        v = v.replace(TECH_NAME_SRC, TECH_NAME_DST)
    elif v.startswith(STO_NAME_SRC):
        v = STO_NAME_DST + v[len(STO_NAME_SRC):]
    return v


# =====================================================================
# Helpers openpyxl
# =====================================================================
def header_map(ws) -> dict[str, int]:
    """{str(header): índice de columna 1-based} de la fila 1."""
    return {str(c.value): i for i, c in enumerate(ws[1], start=1)
            if c.value is not None}


def year_columns(hmap: dict[str, int]) -> list[str]:
    """Headers que son años (dígitos), en orden de columna."""
    return sorted((h for h in hmap if h.isdigit()), key=lambda h: hmap[h])


def rows_where(ws, col_idx: int, value) -> list[int]:
    """Índices de fila 1-based (sin header) donde col_idx == value."""
    out = []
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=col_idx).value == value:
            out.append(r)
    return out


def purge_rows(ws, col_idx: int, prefixes: tuple[str, ...]) -> int:
    """Borra (de abajo hacia arriba) filas cuyo valor en col_idx empiece
    con alguno de los prefijos. Devuelve cuántas borró."""
    removed = 0
    for r in range(ws.max_row, 1, -1):
        v = ws.cell(row=r, column=col_idx).value
        if isinstance(v, str) and v.startswith(prefixes):
            ws.delete_rows(r)
            removed += 1
    return removed


def next_id(ws, col_idx: int) -> int:
    """max(valores enteros de la columna) + 1 (ignora no numéricos)."""
    best = 0
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=col_idx).value
        if isinstance(v, (int, float)) and int(v) == v:
            best = max(best, int(v))
    return best + 1


# =====================================================================
# Descubrimiento, fuel names y pre-validación
# =====================================================================
def discover_scenarios(base_dir: Path) -> list[Path]:
    """Carpetas A1_Outputs/A1_Outputs_* ordenadas (todos los escenarios)."""
    return sorted(p for p in Path(base_dir).glob(A1_GLOB) if p.is_dir())


def fuel_names_from_base_year(path: Path) -> dict[str, str]:
    """{fuel_code: Fuel.Name de la PRIMERA ocurrencia} en todo el workbook.

    Criterio first-occurrence de B1: hojas en orden del workbook; dentro de
    cada hoja, columna Fuel.I top-down y luego Fuel.O top-down.
    """
    names: dict[str, str] = {}
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        h = header_map(ws)
        pairs = []
        if 'Fuel.I' in h and 'Fuel.I.Name' in h:
            pairs.append((h['Fuel.I'], h['Fuel.I.Name']))
        if 'Fuel.O' in h and 'Fuel.O.Name' in h:
            pairs.append((h['Fuel.O'], h['Fuel.O.Name']))
        for code_col, name_col in pairs:
            for row in ws.iter_rows(min_row=2, values_only=True):
                code = row[code_col - 1] if code_col - 1 < len(row) else None
                name = row[name_col - 1] if name_col - 1 < len(row) else None
                if isinstance(code, str) and code not in names and name:
                    names[code] = name
    wb.close()
    return names


def _count_rows(path: Path, sheet: str, col_name: str, value: str) -> int:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    h = header_map(ws)
    n = 0
    if col_name in h:
        ci = h[col_name] - 1
        for row in ws.iter_rows(min_row=2, values_only=True):
            if ci < len(row) and row[ci] == value:
                n += 1
    wb.close()
    return n


def validate_inputs(base_dir: Path, scen_dirs: list[Path],
                    countries: list[str]) -> list[str]:
    """Valida TODO antes de escribir NADA. Devuelve lista de errores."""
    errors: list[str] = []
    base_dir = Path(base_dir)

    for sd in scen_dirs:
        scen = sd.name.replace('A1_Outputs_', '')
        missing = [f for f in (BY_FILE, PR_FILE, PM_FILE) if not (sd / f).exists()]
        if missing:
            errors.extend(f"[{scen}] falta {f}" for f in missing)
            continue
        fuel_names = fuel_names_from_base_year(sd / BY_FILE)
        for c in sorted(countries):
            t = src_tech(c)
            for path, sheet, col, val, want, op in [
                (sd / BY_FILE, 'Secondary', 'Tech', t, 2, '=='),
                (sd / PR_FILE, 'Secondary', 'Tech', t, 2, '=='),
                (sd / PM_FILE, 'Fixed Horizon Parameters', 'Tech', t, 2, '=='),
                (sd / PM_FILE, 'Secondary Techs', 'Tech', t, 1, '>='),
                (sd / PM_FILE, 'VariableCost', 'Tech', t, 1, '>='),
            ]:
                n = _count_rows(path, sheet, col, val)
                ok = (n == want) if op == '==' else (n >= want)
                if not ok:
                    errors.append(
                        f"[{scen}] {path.name}/{sheet}: {val} tiene {n} filas"
                        f" (se esperaba {op}{want})")
            if dst_fuel(c) not in fuel_names:
                errors.append(
                    f"[{scen}] Base_Year no contiene el fuel destino"
                    f" {dst_fuel(c)} — no se puede conectar BDS")

    xtra = base_dir / XTRA_PATH
    if not xtra.exists():
        errors.append(f"falta {XTRA_PATH}")
    else:
        for c in sorted(countries):
            for sheet, col, val, want, op in [
                ('Fixed Horizon Parameters', 'STORAGE', src_storage(c), 2, '=='),
                ('CapitalCostStorage', 'STORAGE', src_storage(c), 1, '>='),
                ('TechnologyStorage', 'TECHNOLOGY', src_tech(c), 4, '=='),
            ]:
                n = _count_rows(xtra, sheet, col, val)
                ok = (n == want) if op == '==' else (n >= want)
                if not ok:
                    errors.append(
                        f"[XTRA] {sheet}: {val} tiene {n} filas"
                        f" (se esperaba {op}{want})")

    for cfg in (CONFIG_A, CONFIG_AB):
        if not (base_dir / cfg).exists():
            errors.append(f"falta {cfg}")

    # Anclajes YAML requeridos por patch_config_a/patch_config_ab. Se validan
    # aquí (sin escribir nada) para que un anclaje faltante no tumbe la
    # corrida DESPUÉS de haber guardado los 13 xlsx (BDS presente en xlsx
    # pero ausente en configs -> GLPK "out of domain").
    cfg_a_path = base_dir / CONFIG_A
    if cfg_a_path.exists():
        text_a = cfg_a_path.read_text(encoding='utf-8')
        if not any(l.rstrip() == '  Storage:' for l in text_a.splitlines()):
            errors.append(
                "[CFG A] no se encontró el bloque '  Storage:' — anclaje"
                " requerido")

    cfg_ab_path = base_dir / CONFIG_AB
    if cfg_ab_path.exists():
        text_ab = cfg_ab_path.read_text(encoding='utf-8')
        if not any(l.strip() == 'storage_delay_storage_prefixes:'
                   for l in text_ab.splitlines()):
            errors.append(
                "[CFG AB] no se encontró 'storage_delay_storage_prefixes:'"
                " — anclaje requerido")
        if '["PWRSDS"' not in text_ab and '"PWRBDS"' not in text_ab:
            errors.append(
                "[CFG AB] no se encontró '[\"PWRSDS\"' ni '\"PWRBDS\"' —"
                " anclaje requerido")
    return errors


# =====================================================================
# Escritores por escenario (A-O files)
# =====================================================================
def _clone_rows(ws, h: dict[str, int], key_col_name: str,
                countries: list[str], id_col: str | None = None,
                post=None) -> tuple[int, int]:
    """Patrón compartido purge-and-reinsert.

    - Borra filas cuyo key_col empiece con el prefijo destino (PWRBDS/BDS).
    - Por país (sorted): localiza filas fuente (key_col == código SDS),
      clona aplicando transform_text y el hook `post(new_vals, c)` si existe.
    - id_col: nombre de columna de ID a reasignar (max+1+índice de país).
    """
    key_col = h[key_col_name]
    is_storage_key = key_col_name == 'STORAGE'
    src_key = src_storage if is_storage_key else src_tech
    prefix = (DST_STO_PREFIX,) if is_storage_key else (DST_TECH_PREFIX,)
    removed = purge_rows(ws, key_col, prefix)
    base_id = next_id(ws, h[id_col]) if id_col else None
    ncols = len(ws[1])
    added = 0
    for i, c in enumerate(sorted(countries)):
        for r in rows_where(ws, key_col, src_key(c)):
            vals = [ws.cell(row=r, column=j).value for j in range(1, ncols + 1)]
            new = [transform_text(v, c) for v in vals]
            if id_col:
                new[h[id_col] - 1] = base_id + i
            if post:
                post(new, c)
            ws.append(new)
            added += 1
    return removed, added


def write_base_year(wb, countries: list[str],
                    fuel_names: dict[str, str]) -> tuple[int, int]:
    """A-O_AR_Model_Base_Year / Secondary: 2 filas por país (IAR m1, OAR m2)."""
    ws = wb['Secondary']
    h = header_map(ws)

    def post(new, c):
        fname = fuel_names[dst_fuel(c)]
        if new[h['Fuel.I'] - 1] == dst_fuel(c):
            new[h['Fuel.I.Name'] - 1] = fname
        if new[h['Fuel.O'] - 1] == dst_fuel(c):
            new[h['Fuel.O.Name'] - 1] = fname

    return _clone_rows(ws, h, 'Tech', countries, post=post)


def write_projections(wb, countries: list[str],
                      fuel_names: dict[str, str]) -> tuple[int, int]:
    """A-O_AR_Projections / Secondary: 2 filas por país (Input 1, Output 0.85)."""
    ws = wb['Secondary']
    h = header_map(ws)

    def post(new, c):
        new[h['Fuel.Name'] - 1] = fuel_names[dst_fuel(c)]

    return _clone_rows(ws, h, 'Tech', countries, post=post)


def write_param_fhp(wb, countries: list[str]) -> tuple[int, int]:
    """A-O_Parametrization / Fixed Horizon Parameters: 2 filas por país."""
    ws = wb['Fixed Horizon Parameters']
    return _clone_rows(ws, header_map(ws), 'Tech', countries, id_col='Tech.ID')


def write_param_sectechs(wb, countries: list[str]) -> tuple[int, int]:
    """A-O_Parametrization / Secondary Techs: 11 filas por país.

    ResidualCapacity y TotalAnnualMinCapacityInvestment nacen en cero (D3).
    """
    ws = wb['Secondary Techs']
    h = header_map(ws)
    years = year_columns(h)

    def post(new, c):
        if new[h['Parameter'] - 1] in BLANK_PARAMS:
            new[h['Projection.Mode'] - 1] = 'EMPTY'
            new[h['Projection.Parameter'] - 1] = None
            for y in years:
                new[h[y] - 1] = None

    return _clone_rows(ws, h, 'Tech', countries, id_col='Tech.ID', post=post)


def write_param_varcost(wb, countries: list[str]) -> tuple[int, int]:
    """A-O_Parametrization / VariableCost: 1 fila por país (modo 2, 0)."""
    ws = wb['VariableCost']
    return _clone_rows(ws, header_map(ws), 'Tech', countries, id_col='Tech.ID')


# =====================================================================
# Escritores globales
# =====================================================================
def write_xtra(wb, countries: list[str]) -> dict[str, tuple[int, int]]:
    """A-Xtra_Storage.xlsx: 3 hojas (FHP, CapitalCostStorage, TechnologyStorage)."""
    out: dict[str, tuple[int, int]] = {}

    ws = wb['Fixed Horizon Parameters']
    out['Fixed Horizon Parameters'] = _clone_rows(
        ws, header_map(ws), 'STORAGE', countries, id_col='STORAGE.ID')

    ws = wb['CapitalCostStorage']
    h = header_map(ws)
    years = year_columns(h)

    def post_cc(new, c):
        # Residual de energía nace en cero (D3); CapitalCostStorage se
        # clona tal cual con Projection.Mode='EMPTY' (D4, bug-compatible).
        if new[h['Parameter'] - 1] == 'ResidualStorageCapacity':
            new[h['Projection.Mode'] - 1] = 'EMPTY'
            new[h['Projection.Parameter'] - 1] = None
            for y in years:
                new[h[y] - 1] = None

    out['CapitalCostStorage'] = _clone_rows(
        ws, h, 'STORAGE', countries, id_col='STORAGE.ID', post=post_cc)

    ws = wb['TechnologyStorage']
    out['TechnologyStorage'] = _clone_rows(
        ws, header_map(ws), 'TECHNOLOGY', countries)
    return out


def patch_config_a(text: str, countries: list[str]) -> tuple[str, int]:
    """Inserta BDS{C}XX01 en xtra_scen.Storage de Config_MOMF_T1_A.yaml.

    Edición TEXTUAL (preserva comentarios/formato). Idempotente: primero
    elimina toda línea BDS existente y reinserta según `countries`.
    """
    lines = text.splitlines(keepends=True)
    lines = [l for l in lines if not l.startswith("  - 'BDS")]
    for i, l in enumerate(lines):
        if l.rstrip() == '  Storage:':
            break
    else:
        raise ValueError("No se encontró el bloque '  Storage:' en Config A")
    j = i + 1
    while j < len(lines) and lines[j].startswith("  - '"):
        j += 1
    new = [f"  - 'BDS{c}XX01'\n" for c in sorted(countries)]
    lines[j:j] = new
    return ''.join(lines), len(new)


def patch_config_ab(text: str) -> tuple[str, int]:
    """Config_MOMF_T1_AB.yaml: BDS en delay prefixes + PWRBDS en excludes."""
    edits = 0
    lines = text.splitlines(keepends=True)
    lines = [l for l in lines if l.strip() != '- BDS']
    for i, l in enumerate(lines):
        if l.strip() == 'storage_delay_storage_prefixes:':
            j = i + 1
            while j < len(lines) and lines[j].lstrip().startswith('- '):
                j += 1
            lines[j:j] = ['  - BDS\n']
            edits += 1
            break
    else:
        raise ValueError(
            "No se encontró 'storage_delay_storage_prefixes:' en Config AB")
    text = ''.join(lines)
    if '"PWRBDS"' not in text:
        if '["PWRSDS"' not in text:
            raise ValueError(
                "No se encontró activity_upper_limit_exclude_prefixes en Config AB")
        text = text.replace('["PWRSDS"', '["PWRSDS", "PWRBDS"', 1)
        edits += 1
    return text, edits


# =====================================================================
# Orquestación
# =====================================================================
def run(countries=None, dry_run=None, backup=None, base_dir=None) -> int:
    """Aplica la clonación SDS->BDS. None = usar la config del módulo.

    Devuelve 0 si OK, 1 si la pre-validación falla (nada escrito).
    """
    countries = sorted(COUNTRIES if countries is None else countries)
    dry_run = DRY_RUN if dry_run is None else dry_run
    backup = BACKUP if backup is None else backup
    base_dir = Path(BASE_DIR if base_dir is None else base_dir)

    scen_dirs = discover_scenarios(base_dir)
    if not scen_dirs:
        print(f"ERROR: no hay carpetas {A1_GLOB} bajo {base_dir}")
        return 1

    print(f"Escenarios: {[d.name for d in scen_dirs]}")
    print(f"Países ({len(countries)}): {countries}")
    if dry_run:
        print("*** DRY RUN: no se escribirá nada ***")

    errors = validate_inputs(base_dir, scen_dirs, countries)
    if errors:
        print(f"\nPRE-VALIDACIÓN FALLÓ ({len(errors)} errores) — nada escrito:")
        for e in errors:
            print(f"  ✗ {e}")
        return 1

    # YAML-first: computa el parcheo de ambos configs ANTES del primer save
    # de xlsx. Si patch_config_a/patch_config_ab lanzan (anclaje ausente que
    # la pre-validación no haya detectado), abortamos sin haber escrito nada.
    try:
        text_a = (base_dir / CONFIG_A).read_text(encoding='utf-8')
        new_text_a, n_a = patch_config_a(text_a, countries)
        text_ab = (base_dir / CONFIG_AB).read_text(encoding='utf-8')
        new_text_ab, n_ab = patch_config_ab(text_ab)
    except Exception as exc:  # noqa: BLE001
        print(f"\nERROR calculando el parcheo YAML — nada escrito: {exc!r}")
        return 1

    stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    report: list[tuple[str, str, str, int, int]] = []

    def backup_file(p: Path) -> None:
        if dry_run or not backup:
            return
        dst = base_dir / BACKUP_DIR / stamp / p.relative_to(base_dir)
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(p, dst)

    # --- Por escenario ---
    for sd in scen_dirs:
        scen = sd.name.replace('A1_Outputs_', '')
        fuel_names = fuel_names_from_base_year(sd / BY_FILE)

        p = sd / BY_FILE
        wb = openpyxl.load_workbook(p)
        r, a = write_base_year(wb, countries, fuel_names)
        report.append((scen, BY_FILE, 'Secondary', r, a))
        backup_file(p)
        if not dry_run:
            wb.save(p)

        p = sd / PR_FILE
        wb = openpyxl.load_workbook(p)
        r, a = write_projections(wb, countries, fuel_names)
        report.append((scen, PR_FILE, 'Secondary', r, a))
        backup_file(p)
        if not dry_run:
            wb.save(p)

        p = sd / PM_FILE
        wb = openpyxl.load_workbook(p)
        for label, fn in [('Fixed Horizon Parameters', write_param_fhp),
                          ('Secondary Techs', write_param_sectechs),
                          ('VariableCost', write_param_varcost)]:
            r, a = fn(wb, countries)
            report.append((scen, PM_FILE, label, r, a))
        backup_file(p)
        if not dry_run:
            wb.save(p)

    # --- Globales ---
    p = base_dir / XTRA_PATH
    wb = openpyxl.load_workbook(p)
    for label, (r, a) in write_xtra(wb, countries).items():
        report.append(('GLOBAL', p.name, label, r, a))
    backup_file(p)
    if not dry_run:
        wb.save(p)

    for fname, new_text, n in [(CONFIG_A, new_text_a, n_a),
                               (CONFIG_AB, new_text_ab, n_ab)]:
        p = base_dir / fname
        report.append(('GLOBAL', fname, '-', 0, n))
        backup_file(p)
        if not dry_run:
            p.write_text(new_text, encoding='utf-8')

    # --- Reporte ---
    print(f"\n{'Escenario':<10} {'Archivo':<32} {'Hoja':<26} {'-filas':>7} {'+filas':>7}")
    for scen, fname, sheet, r, a in report:
        print(f"{scen:<10} {fname:<32} {sheet:<26} {r:>7} {a:>7}")
    if not dry_run and backup:
        print(f"\nBackups en: {base_dir / BACKUP_DIR / stamp}")
    print("\nOK" + (" (dry run)" if dry_run else ""))
    return 0


if __name__ == '__main__':
    sys.exit(run())
