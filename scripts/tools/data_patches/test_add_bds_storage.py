#!/usr/bin/env python3
"""
Test pre para add_bds_storage.py.

Corre el clonador contra un SANDBOX temporal (copia de los archivos reales)
y verifica su comportamiento sin tocar nada real.

Uso:
    python scripts/tools/data_patches/test_add_bds_storage.py
"""
import shutil
import sys
import tempfile
from pathlib import Path

import openpyxl
import yaml

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE))
import add_bds_storage as bds  # noqa: E402
import validate_bds_structure as vbs  # noqa: E402

REAL = Path(bds.BASE_DIR)  # inputs/ real (layout nuevo); solo se LEE para armar el sandbox

FAILS: list[str] = []


def check(cond: bool, msg: str) -> None:
    if not cond:
        FAILS.append(msg)


# ---------------------------------------------------------------------------
# T0 — transforms puras
# ---------------------------------------------------------------------------
def t0_transforms() -> None:
    check(bds.src_tech('ARG') == 'PWRSDSARGXX', "src_tech('ARG')")
    check(bds.dst_tech('ARG') == 'PWRBDSARGXX', "dst_tech('ARG')")
    check(bds.src_storage('CRI') == 'SDSCRIXX01', "src_storage('CRI')")
    check(bds.dst_storage('CRI') == 'BDSCRIXX01', "dst_storage('CRI')")
    check(bds.src_fuel('MEX') == 'ELCMEXXX00', "src_fuel('MEX')")
    check(bds.dst_fuel('MEX') == 'ELCMEXXX02', "dst_fuel('MEX')")
    # strings con códigos
    check(bds.transform_text('PWRSDSARGXX', 'ARG') == 'PWRBDSARGXX',
          "transform_text tech code")
    check(bds.transform_text('SDSARGXX01', 'ARG') == 'BDSARGXX01',
          "transform_text storage code")
    check(bds.transform_text('ELCARGXX00', 'ARG') == 'ELCARGXX02',
          "transform_text fuel code")
    # nombres humanos
    check(bds.transform_text(
        'Short duration storage (Power generator) Argentina, region XX', 'ARG')
        == 'Short duration storage (Power generator, non-renewable) Argentina, region XX',
        "transform_text Tech.Name")
    check(bds.transform_text(
        'Short duration storage Argentina, region XX', 'ARG')
        == 'Short duration storage non-renewable Argentina, region XX',
        "transform_text STORAGE.Name")
    # no-strings y strings ajenos intactos
    check(bds.transform_text(0.85, 'ARG') == 0.85, "transform_text float intacto")
    check(bds.transform_text(None, 'ARG') is None, "transform_text None intacto")
    check(bds.transform_text('User defined', 'ARG') == 'User defined',
          "transform_text string ajeno intacto")
    # otro país no se toca
    check(bds.transform_text('PWRSDSCHLXX', 'ARG') == 'PWRSDSCHLXX',
          "transform_text respeta el país del contexto")


# ---------------------------------------------------------------------------
# Sandbox
# ---------------------------------------------------------------------------
ALL = list(bds.COUNTRIES)

FILES_PER_SCEN = [bds.BY_FILE, bds.PR_FILE, bds.PM_FILE]

TARGET_SHEETS = {
    bds.BY_FILE: ['Secondary'],
    bds.PR_FILE: ['Secondary'],
    bds.PM_FILE: ['Fixed Horizon Parameters', 'Secondary Techs', 'VariableCost'],
}
XTRA_SHEETS = ['Fixed Horizon Parameters', 'CapitalCostStorage', 'TechnologyStorage']
CONFIGS = [bds.CONFIG_A, bds.CONFIG_AB]


def new_sandbox() -> Path:
    """Copia los archivos reales a un tmpdir con el mismo layout relativo."""
    sb = Path(tempfile.mkdtemp(prefix='bds_sandbox_'))
    for sd in sorted((REAL / 'A1_Outputs').glob('A1_Outputs_*')):
        if not sd.is_dir():
            continue
        tgt = sb / 'A1_Outputs' / sd.name
        tgt.mkdir(parents=True)
        for f in FILES_PER_SCEN:
            shutil.copy2(sd / f, tgt / f)
    (sb / 'A2_Extra_Inputs').mkdir()
    shutil.copy2(REAL / bds.XTRA_PATH, sb / bds.XTRA_PATH)
    (sb / 'config').mkdir()
    for cfg in CONFIGS:
        shutil.copy2(REAL / cfg, sb / cfg)
    return sb


def read_sheet(path: Path, sheet: str) -> tuple[list[str], list[dict]]:
    """(headers como str, filas como dicts str(header)->valor)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    rows_iter = ws.iter_rows(values_only=True)
    raw_header = next(rows_iter)
    header = [str(h) for h in raw_header if h is not None]
    out = []
    for row in rows_iter:
        out.append({header[i]: row[i] for i in range(len(header))})
    wb.close()
    return header, out


def snapshot(sb: Path) -> dict:
    """Estado comparable de todas las hojas objetivo + textos YAML."""
    snap = {}
    for sd in sorted((sb / 'A1_Outputs').glob('A1_Outputs_*')):
        for fname, sheets in TARGET_SHEETS.items():
            wb = openpyxl.load_workbook(sd / fname, read_only=True, data_only=True)
            for sh in sheets:
                snap[(sd.name, fname, sh)] = [
                    tuple(r) for r in wb[sh].iter_rows(values_only=True)]
            wb.close()
    wb = openpyxl.load_workbook(sb / bds.XTRA_PATH, read_only=True, data_only=True)
    for sh in XTRA_SHEETS:
        snap[('XTRA', 'A-Xtra_Storage.xlsx', sh)] = [
            tuple(r) for r in wb[sh].iter_rows(values_only=True)]
    wb.close()
    for cfg in CONFIGS:
        snap[('CFG', cfg, '-')] = (sb / cfg).read_text(encoding='utf-8')
    return snap


# ---------------------------------------------------------------------------
# T-sandbox — smoke test de la infraestructura
# ---------------------------------------------------------------------------
def t_sandbox() -> None:
    sb = new_sandbox()
    try:
        scen_dirs = sorted((sb / 'A1_Outputs').glob('A1_Outputs_*'))
        check(len(scen_dirs) >= 1, "sandbox: al menos 1 escenario copiado")
        for sd in scen_dirs:
            for f in FILES_PER_SCEN:
                check((sd / f).exists(), f"sandbox: falta {sd.name}/{f}")
        check((sb / bds.XTRA_PATH).exists(), "sandbox: falta Xtra_Storage")
        for cfg in CONFIGS:
            check((sb / cfg).exists(), f"sandbox: falta {cfg}")
        snap = snapshot(sb)
        check(len(snap) > 0, "snapshot no vacío")
        hdr, _rows = read_sheet(scen_dirs[0] / bds.PR_FILE, 'Secondary')
        check('2023' in hdr and '2050' in hdr,
              "read_sheet normaliza años a str en Projections")
        hdr2, _ = read_sheet(scen_dirs[0] / bds.PM_FILE, 'Secondary Techs')
        check('2023' in hdr2, "read_sheet normaliza años int a str en Secondary Techs")
    finally:
        shutil.rmtree(sb, ignore_errors=True)


# ---------------------------------------------------------------------------
# T-preval — descubrimiento, fuel names y pre-validación
# ---------------------------------------------------------------------------
def t_preval() -> None:
    sb = new_sandbox()
    try:
        scen_dirs = bds.discover_scenarios(sb)
        check(len(scen_dirs) >= 1, "discover_scenarios encuentra escenarios")
        check(scen_dirs == sorted(scen_dirs), "discover_scenarios ordenado")

        fn = bds.fuel_names_from_base_year(scen_dirs[0] / bds.BY_FILE)
        for c in ALL:
            check(bds.dst_fuel(c) in fn,
                  f"fuel_names: falta {bds.dst_fuel(c)} en Base_Year")
        name02 = fn.get(bds.dst_fuel('ARG'), '')
        check(isinstance(name02, str) and len(name02) > 0,
              "fuel_names: nombre no vacío para ELCARGXX02")

        errs = bds.validate_inputs(sb, scen_dirs, ALL)
        check(errs == [], f"validate_inputs happy path devolvió: {errs[:3]}")

        errs = bds.validate_inputs(sb, scen_dirs, ['ZZZ'])
        check(len(errs) > 0, "validate_inputs detecta país sin bloque SDS")
    finally:
        shutil.rmtree(sb, ignore_errors=True)


# ---------------------------------------------------------------------------
# T-writers — escritores por escenario aplicados a un escenario del sandbox
# ---------------------------------------------------------------------------
def t_writers() -> None:
    sb = new_sandbox()
    try:
        sd = bds.discover_scenarios(sb)[0]
        fuel_names = bds.fuel_names_from_base_year(sd / bds.BY_FILE)
        two = ['ARG', 'CRI']

        # --- Base_Year ---
        wb = openpyxl.load_workbook(sd / bds.BY_FILE)
        removed, added = bds.write_base_year(wb, two, fuel_names)
        check(removed == 0 and added == 4, f"BY: removed={removed} added={added}")
        wb.save(sd / bds.BY_FILE)
        _hdr, rows = read_sheet(sd / bds.BY_FILE, 'Secondary')
        arg = [r for r in rows if r['Tech'] == 'PWRBDSARGXX']
        check(len(arg) == 2, "BY: 2 filas PWRBDSARGXX")
        m1 = [r for r in arg if r['Mode.Operation'] == 1]
        m2 = [r for r in arg if r['Mode.Operation'] == 2]
        check(len(m1) == 1 and m1[0]['Fuel.I'] == 'ELCARGXX02'
              and m1[0]['Value.Fuel.I'] == 1 and m1[0]['Fuel.O'] is None,
              "BY: modo 1 = carga desde ELCARGXX02")
        check(len(m2) == 1 and m2[0]['Fuel.O'] == 'ELCARGXX02'
              and m2[0]['Value.Fuel.O'] == 1 and m2[0]['Fuel.I'] is None,
              "BY: modo 2 = descarga hacia ELCARGXX02")
        check(m1[0]['Fuel.I.Name'] == fuel_names['ELCARGXX02'],
              "BY: Fuel.I.Name = nombre real del nodo 02")
        check(m2[0]['Tech.Name'] ==
              'Short duration storage (Power generator, non-renewable)'
              ' Argentina, region XX', "BY: Tech.Name transformado")

        # --- Projections ---
        wb = openpyxl.load_workbook(sd / bds.PR_FILE)
        removed, added = bds.write_projections(wb, two, fuel_names)
        check(removed == 0 and added == 4, f"PR: removed={removed} added={added}")
        wb.save(sd / bds.PR_FILE)
        _hdr, rows = read_sheet(sd / bds.PR_FILE, 'Secondary')
        arg = [r for r in rows if r['Tech'] == 'PWRBDSARGXX']
        years = [str(y) for y in range(2023, 2051)]
        inp = [r for r in arg if r['Direction'] == 'Input']
        out = [r for r in arg if r['Direction'] == 'Output']
        check(len(inp) == 1 and len(out) == 1, "PR: 1 Input + 1 Output")
        check(all(inp[0][y] == 1 for y in years), "PR: Input=1 plano")
        check(all(abs(out[0][y] - 0.85) < 1e-9 for y in years),
              "PR: Output=0.85 plano")
        check(inp[0]['Fuel'] == 'ELCARGXX02'
              and inp[0]['Fuel.Name'] == fuel_names['ELCARGXX02'],
              "PR: Fuel y Fuel.Name del nodo 02")
        check(inp[0]['Projection.Mode'] == 'User defined',
              "PR: Projection.Mode preservado")

        # --- Parametrization: FHP ---
        wb = openpyxl.load_workbook(sd / bds.PM_FILE)
        removed, added = bds.write_param_fhp(wb, two)
        check(removed == 0 and added == 4, f"FHP: removed={removed} added={added}")
        r2, a2 = bds.write_param_sectechs(wb, two)
        check(r2 == 0 and a2 == 22, f"ST: removed={r2} added={a2}")
        r3, a3 = bds.write_param_varcost(wb, two)
        check(r3 == 0 and a3 == 2, f"VC: removed={r3} added={a3}")
        wb.save(sd / bds.PM_FILE)

        _hdr, rows = read_sheet(sd / bds.PM_FILE, 'Fixed Horizon Parameters')
        arg = [r for r in rows if r['Tech'] == 'PWRBDSARGXX']
        params = {r['Parameter']: r['Value'] for r in arg}
        check(params.get('CapacityToActivityUnit') == 31.536, "FHP: C2A=31.536")
        check(params.get('OperationalLife') == 30, "FHP: OperationalLife=30")
        ids = {r['Tech.ID'] for r in arg}
        check(len(ids) == 1 and all(isinstance(i, int) for i in ids),
              "FHP: Tech.ID único y entero por tech nueva")

        # --- Secondary Techs: clonado + blanking ---
        _hdr, rows = read_sheet(sd / bds.PM_FILE, 'Secondary Techs')
        arg = [r for r in rows if r['Tech'] == 'PWRBDSARGXX']
        check(len(arg) == 11, f"ST: 11 filas por país, hay {len(arg)}")
        by_param = {r['Parameter']: r for r in arg}
        cc = by_param['CapitalCost']
        check(cc['Projection.Mode'] == 'User defined'
              and all(cc[y] == 944 for y in years), "ST: CapitalCost=944 UD")
        rm = by_param['ReserveMarginTagTechnology']
        check(all(abs(rm[y] - 0.7) < 1e-9 for y in years), "ST: RMTag=0.7")
        for p in ('ResidualCapacity', 'TotalAnnualMinCapacityInvestment'):
            row = by_param[p]
            check(row['Projection.Mode'] == 'EMPTY'
                  and all(row[y] is None for y in years),
                  f"ST: {p} nace en cero (EMPTY + años vacíos)")

        # BRA: el piso 2031 NO se hereda
        wb = openpyxl.load_workbook(sd / bds.PM_FILE)
        bds.write_param_sectechs(wb, ['BRA'])
        wb.save(sd / bds.PM_FILE)
        _hdr, rows = read_sheet(sd / bds.PM_FILE, 'Secondary Techs')
        bra = [r for r in rows if r['Tech'] == 'PWRBDSBRAXX'
               and r['Parameter'] == 'TotalAnnualMinCapacityInvestment']
        check(len(bra) == 1 and bra[0]['2031'] is None,
              "ST: piso BRA-2031 no heredado")

        # --- VariableCost ---
        _hdr, rows = read_sheet(sd / bds.PM_FILE, 'VariableCost')
        arg = [r for r in rows if r['Tech'] == 'PWRBDSARGXX']
        check(len(arg) == 1 and arg[0]['Mode.Operation'] == 2
              and all(arg[0][y] == 0 for y in years),
              "VC: modo 2, 0 plano")

        # --- Idempotencia a nivel writer (purge borra lo previo) ---
        wb = openpyxl.load_workbook(sd / bds.BY_FILE)
        removed, added = bds.write_base_year(wb, two, fuel_names)
        check(removed == 4 and added == 4,
              f"BY re-run: removed={removed} added={added}")
    finally:
        shutil.rmtree(sb, ignore_errors=True)


# ---------------------------------------------------------------------------
# T-globals — Xtra_Storage y parcheo YAML
# ---------------------------------------------------------------------------
def t_globals() -> None:
    sb = new_sandbox()
    try:
        two = ['ARG', 'CRI']
        xtra = sb / bds.XTRA_PATH

        wb = openpyxl.load_workbook(xtra)
        res = bds.write_xtra(wb, two)
        wb.save(xtra)
        check(res['Fixed Horizon Parameters'] == (0, 4),
              f"XTRA FHP: {res['Fixed Horizon Parameters']}")
        check(res['CapitalCostStorage'] == (0, 4),
              f"XTRA CC: {res['CapitalCostStorage']}")
        check(res['TechnologyStorage'] == (0, 8),
              f"XTRA TS: {res['TechnologyStorage']}")

        _hdr, rows = read_sheet(xtra, 'Fixed Horizon Parameters')
        arg = [r for r in rows if r['STORAGE'] == 'BDSARGXX01']
        params = {r['Parameter']: r['Value'] for r in arg}
        check(params.get('StorageLevelStart') == 0, "XTRA: StorageLevelStart=0")
        check(params.get('OperationalLifeStorage') == 30,
              "XTRA: OperationalLifeStorage=30")
        check(all(isinstance(r['STORAGE.ID'], int) and r['STORAGE.ID'] >= 39
                  for r in arg), "XTRA: STORAGE.ID nuevo >= 39")
        check(arg[0]['STORAGE.Name'] ==
              'Short duration storage non-renewable Argentina, region XX',
              "XTRA: STORAGE.Name transformado")

        years = [str(y) for y in range(2023, 2051)]
        _hdr, rows = read_sheet(xtra, 'CapitalCostStorage')
        arg = {r['Parameter']: r for r in rows if r['STORAGE'] == 'BDSARGXX01'}
        cc = arg['CapitalCostStorage']
        check(cc['Projection.Mode'] == 'EMPTY'
              and all(abs(cc[y] - 32777.77) < 1e-6 for y in years),
              "XTRA: CapitalCostStorage=32777.77 con modo EMPTY (bug-compatible)")
        rsc = arg['ResidualStorageCapacity']
        check(rsc['Projection.Mode'] == 'EMPTY'
              and all(rsc[y] is None for y in years),
              "XTRA: ResidualStorageCapacity vacío")

        # CHL tiene residual >0 en SDS: en BDS debe quedar vacío
        wb = openpyxl.load_workbook(xtra)
        bds.write_xtra(wb, ['CHL'])
        # Check en memoria (sin guardar, para que xtra siga con ARG/CRI en disco)
        ws = wb['CapitalCostStorage']
        h = bds.header_map(ws)
        chl = []
        for r in range(2, ws.max_row + 1):
            if ws.cell(row=r, column=h['STORAGE']).value == 'BDSCHLXX01':
                param = ws.cell(row=r, column=h['Parameter']).value
                if param == 'ResidualStorageCapacity':
                    row_dict = {}
                    for col_name in h:
                        col_idx = h[col_name]
                        row_dict[col_name] = ws.cell(row=r, column=col_idx).value
                    chl.append(row_dict)
        check(len(chl) == 1 and all(chl[0].get(str(y)) is None for y in years),
              "XTRA: residual CHL no heredado")

        _hdr, rows = read_sheet(xtra, 'TechnologyStorage')
        arg = [r for r in rows if r['TECHNOLOGY'] == 'PWRBDSARGXX']
        check(len(arg) == 4, "XTRA TS: 4 filas por país")
        check(all(r['STORAGE'] == 'BDSARGXX01' for r in arg),
              "XTRA TS: acople 1:1 con BDSARGXX01")
        flags = {(r['MODE_OF_OPERATION'], r['Parameter']): r['Value.STORAGE']
                 for r in arg}
        check(flags[(1, 'TechnologyToStorage')] == 1
              and flags[(1, 'TechnologyFromStorage')] == 0
              and flags[(2, 'TechnologyToStorage')] == 0
              and flags[(2, 'TechnologyFromStorage')] == 1,
              "XTRA TS: patrón de flags m1-carga / m2-descarga")

        # --- YAML A (puro, sobre strings) ---
        text_a = (sb / bds.CONFIG_A).read_text(encoding='utf-8')
        new_a, n = bds.patch_config_a(text_a, two)
        check(n == 2, f"CFG A: insertó {n} (esperado 2)")
        check("  - 'BDSARGXX01'\n" in new_a and "  - 'BDSCRIXX01'\n" in new_a,
              "CFG A: entradas BDS presentes")
        parsed = yaml.safe_load(new_a)
        storage_list = parsed['xtra_scen']['Storage']
        check('BDSARGXX01' in storage_list and 'SDSURYXX01' in storage_list,
              "CFG A: parsea y conserva las 38 originales")
        check(len([l for l in new_a.splitlines() if l.startswith('#')])
              == len([l for l in text_a.splitlines() if l.startswith('#')]),
              "CFG A: comentarios intactos")
        # idempotente
        again, n2 = bds.patch_config_a(new_a, two)
        check(again == new_a and n2 == 2, "CFG A: idempotente")
        # convergencia al reducir
        one, _ = bds.patch_config_a(new_a, ['ARG'])
        check("'BDSCRIXX01'" not in one, "CFG A: quitar país converge")

        # --- YAML AB ---
        text_ab = (sb / bds.CONFIG_AB).read_text(encoding='utf-8')
        new_ab, n = bds.patch_config_ab(text_ab)
        check(n == 2, f"CFG AB: {n} ediciones (esperado 2)")
        parsed = yaml.safe_load(new_ab)
        check(parsed['storage_delay_storage_prefixes'] == ['SDS', 'LDS', 'BDS'],
              "CFG AB: BDS en delay prefixes")
        check('PWRBDS' in parsed['activity_upper_limit_exclude_prefixes'],
              "CFG AB: PWRBDS en exclude prefixes")
        again, n2 = bds.patch_config_ab(new_ab)
        check(again == new_ab, "CFG AB: idempotente")
    finally:
        shutil.rmtree(sb, ignore_errors=True)


# ---------------------------------------------------------------------------
# Batería E2E sobre run()
# ---------------------------------------------------------------------------
YEARS = [str(y) for y in range(2023, 2051)]


def _bds_count(rows: list[dict], col: str, prefix: str) -> int:
    return sum(1 for r in rows
               if isinstance(r.get(col), str) and r[col].startswith(prefix))


def t_run_full() -> None:
    """T1 conteos + T2 equivalencia + T9 SDS intacto + backups, en una corrida."""
    sb = new_sandbox()
    try:
        before = snapshot(sb)
        rc = bds.run(countries=ALL, dry_run=False, backup=True, base_dir=sb)
        check(rc == 0, f"run() devolvió {rc}")
        n = len(ALL)

        for sd in bds.discover_scenarios(sb):
            scen = sd.name
            # --- T1: conteos por escenario ---
            _h, rows = read_sheet(sd / bds.BY_FILE, 'Secondary')
            check(_bds_count(rows, 'Tech', 'PWRBDS') == 2 * n,
                  f"T1 {scen} BY: 2*{n} filas")
            _h, prows = read_sheet(sd / bds.PR_FILE, 'Secondary')
            check(_bds_count(prows, 'Tech', 'PWRBDS') == 2 * n,
                  f"T1 {scen} PR: 2*{n} filas")
            _h, fhp = read_sheet(sd / bds.PM_FILE, 'Fixed Horizon Parameters')
            check(_bds_count(fhp, 'Tech', 'PWRBDS') == 2 * n,
                  f"T1 {scen} FHP: 2*{n} filas")
            _h, st = read_sheet(sd / bds.PM_FILE, 'Secondary Techs')
            check(_bds_count(st, 'Tech', 'PWRBDS') == 11 * n,
                  f"T1 {scen} ST: 11*{n} filas")
            _h, vc = read_sheet(sd / bds.PM_FILE, 'VariableCost')
            check(_bds_count(vc, 'Tech', 'PWRBDS') == 1 * n,
                  f"T1 {scen} VC: {n} filas")

            # --- T2: equivalencia mecánica SDS->BDS (Projections + ST) ---
            for c in ALL:
                sds = [r for r in prows if r['Tech'] == bds.src_tech(c)]
                bd = [r for r in prows if r['Tech'] == bds.dst_tech(c)]
                for direction in ('Input', 'Output'):
                    s = next(r for r in sds if r['Direction'] == direction)
                    b = next(r for r in bd if r['Direction'] == direction)
                    check(all(b[y] == s[y] for y in YEARS),
                          f"T2 {scen} {c} PR {direction}: años iguales a SDS")
                    check(b['Fuel'] == bds.dst_fuel(c),
                          f"T2 {scen} {c} PR {direction}: fuel 02")
                sds_st = {r['Parameter']: r for r in st
                          if r['Tech'] == bds.src_tech(c)}
                bds_st = {r['Parameter']: r for r in st
                          if r['Tech'] == bds.dst_tech(c)}
                check(set(bds_st) == set(sds_st),
                      f"T2 {scen} {c} ST: mismos 11 parámetros")
                for p, srow in sds_st.items():
                    brow = bds_st[p]
                    if p in bds.BLANK_PARAMS:
                        check(all(brow[y] is None for y in YEARS)
                              and brow['Projection.Mode'] == 'EMPTY',
                              f"T2 {scen} {c} ST {p}: nace en cero")
                    else:
                        check(all(brow[y] == srow[y] for y in YEARS)
                              and brow['Projection.Mode'] == srow['Projection.Mode'],
                              f"T2 {scen} {c} ST {p}: igual a SDS")

            # --- T9: SDS/LDS intactos (prefix-equality: solo se apendea) ---
            for fname, sheets in TARGET_SHEETS.items():
                for sh in sheets:
                    old = before[(scen, fname, sh)]
                    wb = openpyxl.load_workbook(sd / fname, read_only=True,
                                                data_only=True)
                    new = [tuple(r) for r in wb[sh].iter_rows(values_only=True)]
                    wb.close()
                    check(new[:len(old)] == old,
                          f"T9 {scen} {fname}/{sh}: filas originales intactas")

        # --- T1 global: Xtra + YAML ---
        xtra = sb / bds.XTRA_PATH
        _h, rows = read_sheet(xtra, 'Fixed Horizon Parameters')
        check(_bds_count(rows, 'STORAGE', 'BDS') == 2 * n, "T1 XTRA FHP")
        _h, rows = read_sheet(xtra, 'CapitalCostStorage')
        check(_bds_count(rows, 'STORAGE', 'BDS') == 2 * n, "T1 XTRA CC")
        _h, rows = read_sheet(xtra, 'TechnologyStorage')
        check(_bds_count(rows, 'TECHNOLOGY', 'PWRBDS') == 4 * n, "T1 XTRA TS")
        old_xtra = before[('XTRA', 'A-Xtra_Storage.xlsx', 'TechnologyStorage')]
        wb = openpyxl.load_workbook(xtra, read_only=True, data_only=True)
        new_xtra = [tuple(r) for r in
                    wb['TechnologyStorage'].iter_rows(values_only=True)]
        wb.close()
        check(new_xtra[:len(old_xtra)] == old_xtra, "T9 XTRA TS: intacto")

        # --- T10: YAML sanos ---
        text_a = (sb / bds.CONFIG_A).read_text(encoding='utf-8')
        parsed = yaml.safe_load(text_a)
        sto = parsed['xtra_scen']['Storage']
        check(len([s for s in sto if s.startswith('BDS')]) == n,
              "T10 CFG A: n entradas BDS")
        check(len(sto) == 38 + n, "T10 CFG A: 38 originales + n")
        parsed_ab = yaml.safe_load((sb / bds.CONFIG_AB).read_text(encoding='utf-8'))
        check('BDS' in parsed_ab['storage_delay_storage_prefixes'],
              "T10 CFG AB: BDS en delay")
        check('PWRBDS' in parsed_ab['activity_upper_limit_exclude_prefixes'],
              "T10 CFG AB: PWRBDS en excludes")

        # --- Backups: 4 escenarios × 3 archivos + Xtra + 2 YAML = 15 ---
        bdirs = list((sb / bds.BACKUP_DIR).iterdir())
        check(len(bdirs) == 1, "backup: 1 carpeta timestamped")
        n_files = sum(1 for p in bdirs[0].rglob('*') if p.is_file())
        n_scen = len(bds.discover_scenarios(sb))
        check(n_files == n_scen * 3 + 3,
              f"backup: {n_files} archivos (esperado {n_scen * 3 + 3})")

        # --- T5: idempotencia (2ª corrida = mismo estado) ---
        snap1 = snapshot(sb)
        rc = bds.run(countries=ALL, dry_run=False, backup=False, base_dir=sb)
        check(rc == 0, "run() 2ª vez OK")
        check(snapshot(sb) == snap1, "T5: 2ª corrida idéntica (sin duplicados)")

        # --- T8: convergencia 19 -> 2 países ---
        rc = bds.run(countries=['CRI', 'PAN'], dry_run=False, backup=False,
                     base_dir=sb)
        check(rc == 0, "run() con 2 países OK")
        sd0 = bds.discover_scenarios(sb)[0]
        _h, rows = read_sheet(sd0 / bds.BY_FILE, 'Secondary')
        check(_bds_count(rows, 'Tech', 'PWRBDS') == 4,
              "T8: solo quedan 2 países (4 filas BY)")
        check(not any(r['Tech'] == 'PWRBDSARGXX' for r in rows),
              "T8: ARG eliminado")
        text_a = (sb / bds.CONFIG_A).read_text(encoding='utf-8')
        check("'BDSARGXX01'" not in text_a and "'BDSCRIXX01'" in text_a,
              "T8: CFG A convergió a 2 países")
    finally:
        shutil.rmtree(sb, ignore_errors=True)


def t_dry_run() -> None:
    """T6: DRY_RUN no modifica NADA."""
    sb = new_sandbox()
    try:
        before = snapshot(sb)
        rc = bds.run(countries=ALL, dry_run=True, backup=True, base_dir=sb)
        check(rc == 0, f"dry-run devolvió {rc}")
        check(snapshot(sb) == before, "T6: dry-run no cambió archivos")
        check(not (sb / bds.BACKUP_DIR).exists(), "T6: dry-run no crea backups")
    finally:
        shutil.rmtree(sb, ignore_errors=True)


def t_country_filter() -> None:
    """T7: solo los países pedidos, en todos los escenarios."""
    sb = new_sandbox()
    try:
        rc = bds.run(countries=['CRI', 'PAN'], dry_run=False, backup=False,
                     base_dir=sb)
        check(rc == 0, f"run() filtrado devolvió {rc}")
        for sd in bds.discover_scenarios(sb):
            _h, rows = read_sheet(sd / bds.PM_FILE, 'Secondary Techs')
            techs = {r['Tech'] for r in rows
                     if isinstance(r.get('Tech'), str)
                     and r['Tech'].startswith('PWRBDS')}
            check(techs == {'PWRBDSCRIXX', 'PWRBDSPANXX'},
                  f"T7 {sd.name}: techs = {techs}")
    finally:
        shutil.rmtree(sb, ignore_errors=True)


def t_validation_aborts() -> None:
    """La pre-validación aborta sin escrituras parciales."""
    sb = new_sandbox()
    try:
        before = snapshot(sb)
        rc = bds.run(countries=['ZZZ'], dry_run=False, backup=False, base_dir=sb)
        check(rc == 1, f"run() con país inválido devolvió {rc} (esperado 1)")
        check(snapshot(sb) == before, "abort: nada escrito")
    finally:
        shutil.rmtree(sb, ignore_errors=True)

    # Ancla YAML corrupta (Config A sin '  Storage:') -> debe abortar en la
    # pre-validación SIN haber tocado ningún xlsx (agujero que Fix 1 cierra).
    sb = new_sandbox()
    try:
        cfg_a_path = sb / bds.CONFIG_A
        text = cfg_a_path.read_text(encoding='utf-8')
        corrupted = text.replace('  Storage:', '  StorageX:')
        check(corrupted != text, "sandbox: corrupción de '  Storage:' aplicada")
        cfg_a_path.write_text(corrupted, encoding='utf-8')

        before = snapshot(sb)
        rc = bds.run(countries=ALL, dry_run=False, backup=False, base_dir=sb)
        check(rc == 1,
              f"run() con ancla '  Storage:' corrupta devolvió {rc} (esperado 1)")
        check(snapshot(sb) == before, "abort YAML: nada escrito (ni xlsx ni configs)")
    finally:
        shutil.rmtree(sb, ignore_errors=True)


# ---------------------------------------------------------------------------
# T-validator — el auditor post detecta estados buenos y malos
# ---------------------------------------------------------------------------
def t_validator() -> None:
    # 1) sandbox SIN aplicar: debe fallar (estructura BDS ausente)
    sb = new_sandbox()
    try:
        rc = vbs.audit(countries=ALL, base_dir=sb)
        check(rc == 1, "validator: falla sobre sandbox sin aplicar")
    finally:
        shutil.rmtree(sb, ignore_errors=True)

    # 2) sandbox aplicado: debe pasar completo
    sb = new_sandbox()
    try:
        check(bds.run(countries=ALL, dry_run=False, backup=False,
                      base_dir=sb) == 0, "validator-prep: run OK")
        rc = vbs.audit(countries=ALL, base_dir=sb)
        check(rc == 0, "validator: pasa sobre sandbox aplicado")

        # 3) corrupción puntual: CapitalCost de un BDS alterado -> V6 falla
        sd = bds.discover_scenarios(sb)[0]
        p = sd / bds.PM_FILE
        wb = openpyxl.load_workbook(p)
        ws = wb['Secondary Techs']
        h = {str(c.value): i for i, c in enumerate(ws[1], start=1)
             if c.value is not None}
        for r in range(2, ws.max_row + 1):
            if (ws.cell(row=r, column=h['Tech']).value == 'PWRBDSARGXX'
                    and ws.cell(row=r, column=h['Parameter']).value
                    == 'CapitalCost'):
                ws.cell(row=r, column=h['2023'], value=999)
                break
        wb.save(p)
        rc = vbs.audit(countries=ALL, base_dir=sb)
        check(rc == 1, "validator: detecta CapitalCost corrupto (V6)")
    finally:
        shutil.rmtree(sb, ignore_errors=True)


TESTS = [
    ('T0 transforms puras', t0_transforms),
    ('T-sandbox infraestructura', t_sandbox),
    ('T-preval descubrimiento y validación', t_preval),
    ('T-writers escritores por escenario', t_writers),
    ('T-globals Xtra y YAML', t_globals),
    ('T1/T2/T5/T8/T9/T10+backup run() completo', t_run_full),
    ('T6 dry-run inocuo', t_dry_run),
    ('T7 filtro de países', t_country_filter),
    ('T-abort pre-validación', t_validation_aborts),
    ('T-validator auditor post', t_validator),
]


def main() -> int:
    any_fail = False
    for name, fn in TESTS:
        FAILS.clear()
        try:
            fn()
        except Exception as exc:  # noqa: BLE001
            FAILS.append(f"EXCEPCIÓN: {exc!r}")
        if FAILS:
            any_fail = True
            print(f"❌ {name}")
            for f in FAILS:
                print(f"     ✗ {f}")
        else:
            print(f"✅ {name}")
    print()
    print("RESULTADO:", "FALLÓ" if any_fail else "OK")
    return 1 if any_fail else 0


if __name__ == '__main__':
    sys.exit(main())
