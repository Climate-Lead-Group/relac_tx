"""Open TotalAnnualMaxCapacityInvestment in OPT only where the new floors require it.

Already applied to inputs/ on 2026-08-13 14:25 — log
A1_Outputs/open_maxcaps_pisos_OPT_20260813_142555.json (82 cells, see
`scripts/tools/data_patches/README.md`). Manual complement to a step A3 would otherwise have
done: in OPT's A-O_Parametrization.xlsx (sheet "Secondary Techs"), for the 24 techs of the
Nueva_Capacidad_2040+ dataset (TECHS24), raises the Max cap ONLY on the (tech, year) cells
where the new floor (TotalAnnualMinCapacityInvestment) exceeds the current cap and would make
the model infeasible. Rule: new_max = floor * 1.01 (SLACK, the VGB precedent of Max=Min*1.01).
Max rows with Projection.Mode='EMPTY' or empty cells mean no cap (unlimited in OSeMOSYS) and
are left untouched; years without a conflict keep their current cap.

Inputs:
    A1_Outputs_OPT/A-O_Parametrization.xlsx / Secondary Techs
Outputs:
    same file, in place; a "*_backup_<ts>.xlsx" backup and a cell-level old->new change log
        at A1_Outputs/open_maxcaps_pisos_OPT_<ts>.json
Usage:
    python scripts/tools/data_patches/patch_open_maxcaps_pisos_OPT.py [--dry-run]
    --dry-run reports the conflicts found without writing (today reports 0 conflicts).
"""
import argparse
import json
import shutil
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))  # -> scripts/
from common import relac_paths as P  # noqa: E402

SCRIPT_DIR = Path(__file__).parent
TARGET = P.scenario_dir("OPT") / "A-O_Parametrization.xlsx"   # inputs/A1_Outputs/A1_Outputs_OPT

MIN_PARAM = "TotalAnnualMinCapacityInvestment"
MAX_PARAM = "TotalAnnualMaxCapacityInvestment"
Y0, Y1 = 2040, 2050
SLACK = 1.01

TECHS24 = [
    "PWRBIOARGXX", "PWRGEOGTMXX", "PWRHYDARGXX", "PWRHYDBRAXX", "PWRHYDGTMXX",
    "PWRLDSCHLXX", "PWRNGSARGXX", "PWROILURYXX", "PWRSDSARGXX", "PWRSDSCHLXX",
    "PWRSDSPRYXX", "PWRSPVARGXX", "PWRSPVBRAXX", "PWRSPVCHLXX", "PWRSPVCRIXX",
    "PWRSPVPRYXX", "PWRSPVURYXX", "PWRURNARGXX", "PWRWONARGXX", "PWRWONBRAXX",
    "PWRWONCHLXX", "PWRWONCRIXX", "PWRWONGTMXX", "PWRWONURYXX",
]

COL_TECH, COL_PARAM, COL_MODE = 2, 5, 7


def main():
    """Raise OPT's Max cap wherever the new floor exceeds it, then verify no conflict remains.

    Reads Min/Max rows for TECHS24, raises Max to floor*SLACK on every (tech, year) in
    2040-2050 where floor > cap (skipping rows not in "User defined" mode or with no cap),
    saves the workbook with a backup and change log unless --dry-run, then rereads the file
    and returns 1 if any Min > Max conflict remains.
    """
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(TARGET)
    ws = wb["Secondary Techs"]
    ycols = {}
    for c in range(1, ws.max_column + 1):
        h = ws.cell(1, c).value
        if isinstance(h, (int, float)) and 2000 <= int(h) <= 2100:
            ycols[int(h)] = c

    min_rows, max_rows = {}, {}
    for r in range(2, ws.max_row + 1):
        tech = ws.cell(r, COL_TECH).value
        if tech not in TECHS24:
            continue
        param = ws.cell(r, COL_PARAM).value
        if param == MIN_PARAM:
            min_rows[tech] = r
        elif param == MAX_PARAM:
            max_rows[tech] = r

    missing = [t for t in TECHS24 if t not in min_rows]
    if missing:
        raise SystemExit(f"Faltan filas Min para: {missing}")

    changes = []
    for tech in TECHS24:
        rmin = min_rows[tech]
        if ws.cell(rmin, COL_MODE).value != "User defined":
            continue
        rmax = max_rows.get(tech)
        # sin fila Max o fila EMPTY -> ilimitado -> no hay conflicto posible
        if rmax is None or ws.cell(rmax, COL_MODE).value != "User defined":
            continue
        for y in range(Y0, Y1 + 1):
            floor = ws.cell(rmin, ycols[y]).value
            if floor in (None, ""):
                continue
            cap = ws.cell(rmax, ycols[y]).value
            if cap in (None, ""):  # celda vacia = sin tope ese anio
                continue
            if float(floor) > float(cap):
                new_cap = round(float(floor) * SLACK, 9)
                ws.cell(rmax, ycols[y]).value = new_cap
                changes.append({
                    "scenario": "OPT", "tech": tech, "parameter": MAX_PARAM,
                    "year": y, "old": cap, "new": new_cap, "floor": floor,
                })

    print(f"Celdas en conflicto (Min > Max) abiertas: {len(changes)}")
    per_tech = {}
    for c in changes:
        per_tech.setdefault(c["tech"], []).append(c)
    for tech, cs in sorted(per_tech.items()):
        yrs = [c["year"] for c in cs]
        print(f"  {tech}: {len(cs)} anios ({min(yrs)}-{max(yrs)}), "
              f"ej. {cs[0]['year']}: {cs[0]['old']} -> {cs[0]['new']}")

    if args.dry_run:
        print("(dry-run, sin guardar)")
        wb.close()
        return 0

    if changes:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup = TARGET.with_name(f"{TARGET.stem}_backup_{ts}{TARGET.suffix}")
        shutil.copy2(TARGET, backup)
        wb.save(TARGET)
        log_path = P.A1_OUTPUTS / f"open_maxcaps_pisos_OPT_{ts}.json"
        log_path.write_text(json.dumps(changes, indent=2, ensure_ascii=False))
        print(f"Backup: {backup.name}\nLog: {log_path}")
    wb.close()

    # verificacion: releer y confirmar Max >= Min en 2040-2050 para las 24 techs
    wb2 = openpyxl.load_workbook(TARGET, read_only=True, data_only=True)
    ws2 = wb2["Secondary Techs"]
    hdr = None
    mins, maxs, modes = {}, {}, {}
    for row in ws2.iter_rows(values_only=True):
        if hdr is None:
            hdr = row
            yc = {int(h): i for i, h in enumerate(hdr) if isinstance(h, (int, float))}
            continue
        if row[COL_TECH - 1] in TECHS24:
            key = (row[COL_TECH - 1], row[COL_PARAM - 1])
            modes[key] = row[COL_MODE - 1]
            for y in range(Y0, Y1 + 1):
                v = row[yc[y]]
                if v not in (None, ""):
                    if row[COL_PARAM - 1] == MIN_PARAM:
                        mins[(row[COL_TECH - 1], y)] = float(v)
                    elif row[COL_PARAM - 1] == MAX_PARAM:
                        maxs[(row[COL_TECH - 1], y)] = float(v)
    wb2.close()
    bad = []
    for (tech, y), f in mins.items():
        if modes.get((tech, MIN_PARAM)) != "User defined":
            continue
        if modes.get((tech, MAX_PARAM)) != "User defined":
            continue
        cap = maxs.get((tech, y))
        if cap is not None and f > cap:
            bad.append(f"{tech} {y}: Min {f} > Max {cap}")
    if bad:
        print("\nVERIFICACION: QUEDAN CONFLICTOS")
        for b in bad:
            print("  -", b)
        return 1
    print("\nVERIFICACION: OK — ningun piso supera su tope en OPT 2040-2050")
    return 0


if __name__ == "__main__":
    sys.exit(main())
