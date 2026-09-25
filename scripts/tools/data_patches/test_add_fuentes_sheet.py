#!/usr/bin/env python3
"""
Test de add_fuentes_sheet.py sobre copias temporales (no toca inputs/ real).

Verifica:
  1. La hoja `Fuentes` se crea al final, con las filas del CSV que corresponden a ese libro.
  2. Ninguna otra hoja cambia (valores celda a celda).
  3. Idempotencia: aplicar dos veces deja el mismo resultado y no duplica hojas.
  4. El filtro por `Archivo` (con `|`) reparte bien las filas entre los 3 tipos de libro.
  5. La lectura estilo B1 (pandas ExcelFile) sigue funcionando al excluir `Fuentes`, y
     sync_historical_from_bau.year_columns() no ve columnas de año en la hoja nueva.
  6. Las plantillas de Miscellaneous reciben la hoja con encabezado pero sin filas de datos.

Uso:
    python scripts/tools/data_patches/test_add_fuentes_sheet.py
"""
from __future__ import annotations

import gc
import shutil
import sys
import tempfile
import time
from pathlib import Path

import openpyxl
import pandas as pd

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parents[1]))          # scripts/
sys.path.insert(0, str(HERE.parents[1] / "pipeline"))
import add_fuentes_sheet as afs  # noqa: E402
from common import relac_paths as P  # noqa: E402

FAILS: list[str] = []


def check(cond: bool, msg: str) -> None:
    print(("  OK   " if cond else "  FAIL ") + msg)
    if not cond:
        FAILS.append(msg)


def snapshot(path: Path, skip: str) -> dict[str, list[tuple]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    snap = {ws.title: list(ws.iter_rows(values_only=True)) for ws in wb.worksheets if ws.title != skip}
    wb.close()
    return snap


def main() -> int:
    rows = afs.load_rows()
    check(len(rows) > 50, f"CSV maestro con {len(rows)} filas")
    by_key = {k: afs.rows_for(rows, k) for k in afs.WORKBOOK_FILES}
    check(all(len(v) > 0 for v in by_key.values()), f"filas por libro: { {k: len(v) for k, v in by_key.items()} }")
    shared = [r for r in rows if "|" in r["Archivo"]]
    check(shared and all(r in by_key["A-O_Parametrization"] and r in by_key["A-O_Demand"] for r in shared),
          f"{len(shared)} filas compartidas (Archivo con '|') llegan a ambos libros A-O")

    tmp = Path(tempfile.mkdtemp(prefix="fuentes_sandbox_"))
    try:
        sources = {
            "A-O_Parametrization": P.scenario_dir("BAU") / "A-O_Parametrization.xlsx",
            "A-O_Demand": P.scenario_dir("BAU") / "A-O_Demand.xlsx",
            "A-Xtra_Storage": P.A2_EXTRA_INPUTS / "A-Xtra_Storage.xlsx",
        }
        for key, src in sources.items():
            dst = tmp / src.name
            shutil.copy2(src, dst)
            before = snapshot(dst, skip=afs.SHEET_NAME)
            sheets_before = openpyxl.load_workbook(dst, read_only=True).sheetnames

            afs.apply_to_workbook(dst, by_key[key], dry_run=False)
            afs.apply_to_workbook(dst, by_key[key], dry_run=False)   # segunda vez: idempotencia

            wb = openpyxl.load_workbook(dst)
            names = wb.sheetnames
            check(names.count(afs.SHEET_NAME) == 1 and names[-1] == afs.SHEET_NAME,
                  f"{src.name}: una sola hoja '{afs.SHEET_NAME}' y va al final")
            check([n for n in names if n != afs.SHEET_NAME] == [n for n in sheets_before if n != afs.SHEET_NAME],
                  f"{src.name}: las demás hojas conservan nombre y orden")
            ws = wb[afs.SHEET_NAME]
            header = [ws.cell(row=2, column=c).value for c in range(1, len(afs.COLUMNS) + 1)]
            check(header == afs.COLUMNS, f"{src.name}: encabezado {afs.COLUMNS}")
            data = [[ws.cell(row=r, column=c).value for c in range(1, len(afs.COLUMNS) + 1)]
                    for r in range(3, ws.max_row + 1)]
            expected = [[(row[c] or None) for c in afs.COLUMNS] for row in by_key[key]]
            check(data == expected, f"{src.name}: {len(expected)} filas idénticas al CSV")
            fcol = afs.COLUMNS.index("Fuente") + 1
            marked = [ws.cell(row=r, column=fcol) for r in range(3, ws.max_row + 1)
                      if afs.ASSUMPTION_TAG.lower() in str(ws.cell(row=r, column=fcol).value or "").lower()]
            others = [ws.cell(row=r, column=fcol) for r in range(3, ws.max_row + 1)
                      if afs.ASSUMPTION_TAG.lower() not in str(ws.cell(row=r, column=fcol).value or "").lower()]
            if key == "A-O_Parametrization":
                check(marked and all(c.font.bold for c in marked),
                      f"{src.name}: {len(marked)} celdas 'Supuesto propio' en negrita")
                check(all(c.fill.patternType is None for c in marked),
                      f"{src.name}: celdas 'Supuesto propio' sin relleno")
            check(all(not c.font.underline for c in marked + others),
                  f"{src.name}: ninguna fuente va subrayada")
            check(all(c.fill.patternType is None for c in others), f"{src.name}: las demás fuentes sin relleno")
            wb.close()

            after = snapshot(dst, skip=afs.SHEET_NAME)
            check(after == before, f"{src.name}: ninguna otra hoja cambió (valores celda a celda)")

            # Lectura estilo B1: pandas ExcelFile, excluyendo 'Fuentes' como hace B1 con growth_formula
            xls = pd.ExcelFile(dst)
            param_sheets = [s for s in xls.sheet_names if s != afs.SHEET_NAME]
            if key == "A-O_Parametrization":
                ok = all("Parameter" in xls.parse(s).columns for s in param_sheets)
                check(ok, f"{src.name}: todas las hojas que B1 itera siguen teniendo columna 'Parameter'")
            elif key == "A-O_Demand":
                ok = all("Demand/Share" in xls.parse(s).columns for s in param_sheets)
                check(ok, f"{src.name}: todas las hojas que B1 itera siguen teniendo 'Demand/Share'")
            xls.close()

            # sync_historical_from_bau.year_columns no debe ver años en la hoja nueva
            try:
                import sync_historical_from_bau as sh  # noqa: E402
                wb2 = openpyxl.load_workbook(dst, read_only=True)
                yc = sh.year_columns(wb2[afs.SHEET_NAME], set(range(2023, 2051)))
                wb2.close()
                check(yc == {}, f"{src.name}: sync_historical no detecta columnas de año en '{afs.SHEET_NAME}'")
            except ImportError as e:  # pragma: no cover
                print(f"  SKIP sync_historical no importable: {e}")
        # Plantillas: la hoja debe quedar solo con título y encabezado (0 filas de datos)
        tmpl = tmp / "plantilla_A-O_Demand.xlsx"
        shutil.copy2(P.MISCELLANEOUS / "A-O_Demand.xlsx", tmpl)
        afs.apply_to_workbook(tmpl, [], dry_run=False)
        wbt = openpyxl.load_workbook(tmpl, read_only=True)
        wst = wbt[afs.SHEET_NAME]
        data_rows = [r for r in wst.iter_rows(min_row=3, values_only=True) if any(v is not None for v in r)]
        hdr_t = [c.value for c in next(wst.iter_rows(min_row=2, max_row=2))][: len(afs.COLUMNS)]
        wbt.close()
        check(hdr_t == afs.COLUMNS and not data_rows, "plantilla: hoja Fuentes con encabezado y 0 filas de datos")
        tmpl_targets = [(k, p, tp) for k, p, tp in afs.targets(True) if tp]
        check(len(tmpl_targets) == 3 and all("Miscellaneous" in str(p) for _, p, _ in tmpl_targets),
              "targets(): las 3 plantillas de Miscellaneous van marcadas como plantilla")
    finally:
        gc.collect()
        time.sleep(0.2)
        shutil.rmtree(tmp, ignore_errors=True)

    print("\nRESULTADO:", "TODO OK" if not FAILS else f"{len(FAILS)} fallos")
    return 0 if not FAILS else 1


if __name__ == "__main__":
    sys.exit(main())
