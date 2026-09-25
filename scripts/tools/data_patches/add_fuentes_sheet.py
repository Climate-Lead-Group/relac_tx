#!/usr/bin/env python3
"""
add_fuentes_sheet.py — agrega (o reemplaza) la hoja `Fuentes` en los libros de entrada del modelo.

La hoja documenta de dónde viene cada dato (fuente, archivo/hoja de origen y transformación). Su contenido vive en UNA sola tabla, `inputs/config/data_sources.csv`, para que los
8 libros A-O_* de los 4 escenarios y `A-Xtra_Storage.xlsx` queden siempre idénticos entre sí.

Libros que toca (filtrando las filas del CSV por la columna `Archivo`, que admite varios nombres
separados por `|`):
  - inputs/A1_Outputs/A1_Outputs_<esc>/A-O_Parametrization.xlsx   (BAU, INV, OPT, VGB)
  - inputs/A1_Outputs/A1_Outputs_<esc>/A-O_Demand.xlsx            (BAU, INV, OPT, VGB)
  - inputs/A2_Extra_Inputs/A-Xtra_Storage.xlsx
  - inputs/Miscellaneous/{A-O_Parametrization,A-O_Demand,A-Xtra_Storage}.xlsx  (plantillas: solo título y
    encabezado, SIN filas, porque las plantillas no tienen datos; --no-templates para omitirlas)

La hoja se llama `Fuentes`, va al final del libro y NO tiene columnas de años, de modo que
sync_historical_from_bau.py la salta ("no year columns"). B1_Compiler.py la excluye explícitamente
(igual que hace con `growth_formula`). Ninguna otra hoja se modifica.

Uso:
    python scripts/tools/data_patches/add_fuentes_sheet.py            # dry-run: muestra qué haría
    python scripts/tools/data_patches/add_fuentes_sheet.py --apply    # escribe
    python scripts/tools/data_patches/add_fuentes_sheet.py --apply --no-templates

Idempotente: si la hoja ya existe se elimina y se vuelve a crear desde el CSV.
"""
from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from common import relac_paths as P  # noqa: E402

SHEET_NAME = "Fuentes"
CSV_PATH = P.CONFIG / "data_sources.csv"
SCENARIOS = ["BAU", "INV", "OPT", "VGB"]
COLUMNS = ["Hoja", "Tecnologias", "Parametro", "Años", "Fuente", "Detalle"]
COL_WIDTHS = {"Hoja": 28, "Tecnologias": 42, "Parametro": 40, "Años": 18,
              "Fuente": 55, "Detalle": 90}
TITLE_ROW = ("Fuentes de los datos de este libro. Las fuentes marcadas como 'Supuesto propio' son "
             "estimaciones internas del equipo de modelación (no provienen de una fuente externa). "
             "Hoja informativa: el modelo no la lee.")
ASSUMPTION_TAG = "Supuesto propio"

# (nombre lógico en el CSV, nombre de archivo)
WORKBOOK_FILES = {
    "A-O_Parametrization": "A-O_Parametrization.xlsx",
    "A-O_Demand": "A-O_Demand.xlsx",
    "A-Xtra_Storage": "A-Xtra_Storage.xlsx",
}


def load_rows(csv_path: Path = CSV_PATH) -> list[dict[str, str]]:
    with open(csv_path, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    expected = ["Archivo", *COLUMNS]
    if not rows or list(rows[0].keys()) != expected:
        raise ValueError(f"{csv_path}: columnas esperadas {expected}, hay {list(rows[0].keys()) if rows else []}")
    return rows


def rows_for(rows: list[dict[str, str]], workbook_key: str) -> list[dict[str, str]]:
    return [r for r in rows if workbook_key in [a.strip() for a in r["Archivo"].split("|")]]


def targets(include_templates: bool = True) -> list[tuple[str, Path, bool]]:
    """(clave del CSV, ruta, es_plantilla). Las plantillas reciben la hoja vacía (solo encabezado)."""
    out: list[tuple[str, Path, bool]] = []
    for scen in SCENARIOS:
        for key in ("A-O_Parametrization", "A-O_Demand"):
            out.append((key, P.scenario_dir(scen) / WORKBOOK_FILES[key], False))
    out.append(("A-Xtra_Storage", P.A2_EXTRA_INPUTS / WORKBOOK_FILES["A-Xtra_Storage"], False))
    if include_templates:
        for key, fname in WORKBOOK_FILES.items():
            out.append((key, P.MISCELLANEOUS / fname, True))
    return out


def write_sheet(wb: openpyxl.Workbook, rows: list[dict[str, str]]) -> None:
    """Elimina la hoja `Fuentes` si existe y la recrea al final del libro."""
    if SHEET_NAME in wb.sheetnames:
        del wb[SHEET_NAME]
    ws = wb.create_sheet(SHEET_NAME)  # al final

    ws.cell(row=1, column=1, value=TITLE_ROW).font = Font(italic=True, color="555555")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))

    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for c, name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=2, column=c, value=name)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = COL_WIDTHS[name]

    for r_idx, row in enumerate(rows, start=3):
        for c, name in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=r_idx, column=c, value=row[name] or None)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A3"


def _rel(path: Path) -> str:
    try:
        return str(path.relative_to(P.REPO_ROOT))
    except ValueError:  # p. ej. copias en un sandbox temporal
        return str(path)


def apply_to_workbook(path: Path, rows: list[dict[str, str]], dry_run: bool) -> int:
    if not path.exists():
        print(f"  [omitido] no existe: {path}")
        return 0
    if dry_run:
        print(f"  [dry-run] {_rel(path)}: {len(rows)} filas en hoja '{SHEET_NAME}'" + (" (plantilla: solo encabezado)" if not rows else ""))
        return len(rows)
    wb = openpyxl.load_workbook(path)
    write_sheet(wb, rows)
    wb.save(path)
    wb.close()
    print(f"  [ok] {_rel(path)}: hoja '{SHEET_NAME}' con {len(rows)} filas" + (" (plantilla: solo encabezado)" if not rows else ""))
    return len(rows)


def run(apply: bool, include_templates: bool = True, csv_path: Path = CSV_PATH) -> dict[str, int]:
    rows = load_rows(csv_path)
    summary: dict[str, int] = {}
    print(f"Tabla maestra: {csv_path} ({len(rows)} filas)")
    for key, path, is_template in targets(include_templates):
        # Plantillas de Miscellaneous: no tienen datos, así que la hoja va sin filas (solo título y encabezado).
        wb_rows = [] if is_template else rows_for(rows, key)
        summary[str(path)] = apply_to_workbook(path, wb_rows, dry_run=not apply)
    if not apply:
        print("\nDry-run: nada escrito. Use --apply para escribir la hoja.")
    return summary


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--apply", action="store_true", help="escribir la hoja (por defecto solo dry-run)")
    ap.add_argument("--no-templates", action="store_true",
                    help="no tocar las plantillas de inputs/Miscellaneous")
    args = ap.parse_args()
    run(apply=args.apply, include_templates=not args.no_templates)
    return 0


if __name__ == "__main__":
    sys.exit(main())
