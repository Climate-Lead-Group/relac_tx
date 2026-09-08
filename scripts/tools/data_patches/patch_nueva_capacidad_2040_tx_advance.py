"""
Patch: Nueva Capacidad 2040+ (solo OPT) + adelanto de interconexiones (4 escenarios)

Aplica dos cambios sobre A-O_Parametrization.xlsx (hoja 'Secondary Techs'):

1. PISOS (solo escenario OPT):
   Lee Nueva_Capacidad_2040+_LAC.xlsx (hoja ReLAC_Gen_SecondaryTechs) y escribe
   los valores 2040-2050 de TotalAnnualMinCapacityInvestment para las 24 techs.
   Fija Projection.Mode='User defined' (corrige las 7 filas EMPTY del dataset).
   Solo escribe celdas con valor en el Excel; no toca años previos a 2040.

2. INTERCONEXIONES (BAU, INV, OPT, VGB):
   - TRNCOLXXPANXX: ResidualCapacity=0.4  y UpperLimit=12.614 en 2029-2050
     (antes: entrada 2033, limite decreciente 10.6->8.19)
   - TRNBOLXXBRAXX: ResidualCapacity=0.42 y UpperLimit=13.245 en 2032-2050
     (antes: entrada 2036, limite decreciente 13.35->10.89)
   Los ceros previos al anio de entrada se conservan (linea bloqueada antes).

Uso:
    python scripts/tools/data_patches/patch_nueva_capacidad_2040_tx_advance.py [--dry-run]
        [--xlsx RUTA] [--only-floors | --only-tx]

Genera backup *_backup_<ts>.xlsx por archivo modificado y un JSON de cambios
(old->new por celda) en A1_Outputs/nueva_capacidad_tx_changes_<ts>.json.
"""
import argparse
import json
import shutil
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))  # -> scripts/
from common import relac_paths as P  # noqa: E402

SCRIPT_DIR = Path(__file__).parent
BASE_PATH = P.A1_OUTPUTS   # inputs/A1_Outputs (layout nuevo)
DEFAULT_XLSX = Path(r"C:\Users\ClimateLeadGroup\Downloads\_Dataset_Transmission\Nueva_Capacidad_2040+_LAC - copia.xlsx")

FLOORS_SCENARIOS = ["OPT"]
TX_SCENARIOS = ["BAU", "INV", "OPT", "VGB"]

FLOORS_PARAM = "TotalAnnualMinCapacityInvestment"
FLOORS_SHEET = "ReLAC_Gen_SecondaryTechs"
FLOORS_YEAR_MIN, FLOORS_YEAR_MAX = 2040, 2050

# tech -> {param: (start_year, value)} ; se escribe start_year..2050 constante
TX_CHANGES = {
    "TRNCOLXXPANXX": {
        "ResidualCapacity": (2029, 0.4),
        "TotalTechnologyAnnualActivityUpperLimit": (2029, 12.614),
    },
    "TRNBOLXXBRAXX": {
        "ResidualCapacity": (2032, 0.42),
        "TotalTechnologyAnnualActivityUpperLimit": (2032, 13.245),
    },
}
TX_LAST_YEAR = 2050

COL_TECH = 2        # B: Tech
COL_PARAM = 5       # E: Parameter
COL_MODE = 7        # G: Projection.Mode
COL_PROJPARAM = 8   # H: Projection.Parameter


def read_floors(xlsx_path):
    """Lee el Excel de nueva capacidad -> {tech: {year: value}} (solo 2040-2050)."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[FLOORS_SHEET]
    hdr = [c.value for c in ws[1]]
    year_cols = {i: int(h) for i, h in enumerate(hdr) if isinstance(h, (int, float))}
    floors = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        tech, param = row[1], row[4]
        if not tech:
            continue
        if param != FLOORS_PARAM:
            raise ValueError(f"Fila de {tech}: Parameter inesperado {param!r}")
        vals = {}
        for i, yr in year_cols.items():
            v = row[i]
            if v is None or v == "":
                continue
            if not (FLOORS_YEAR_MIN <= yr <= FLOORS_YEAR_MAX):
                raise ValueError(f"{tech}: valor fuera de rango 2040-2050 en {yr}")
            vals[yr] = float(v)
        if vals:
            floors[tech] = vals
    wb.close()
    return floors


def find_row(ws, tech, param):
    """Devuelve el indice de la unica fila (tech, param); error si 0 o >1."""
    matches = [
        r for r in range(2, ws.max_row + 1)
        if ws.cell(r, COL_TECH).value == tech and ws.cell(r, COL_PARAM).value == param
    ]
    if len(matches) != 1:
        raise ValueError(f"{tech}/{param}: {len(matches)} filas encontradas (se esperaba 1)")
    return matches[0]


def year_columns(ws):
    """{year:int -> col_idx} a partir de la fila 1."""
    out = {}
    for c in range(1, ws.max_column + 1):
        h = ws.cell(1, c).value
        if isinstance(h, (int, float)) and 2000 <= int(h) <= 2100:
            out[int(h)] = c
    return out


def apply_cells(ws, row_idx, target, ycols, changes, scen, tech, param):
    """Escribe {year: value} en la fila; registra old->new. Devuelve n escritos."""
    n = 0
    for yr, val in sorted(target.items()):
        col = ycols.get(yr)
        if col is None:
            raise ValueError(f"{tech}/{param}: no existe columna para el anio {yr}")
        old = ws.cell(row_idx, col).value
        if old == val:
            continue
        ws.cell(row_idx, col).value = val
        changes.append({
            "scenario": scen, "tech": tech, "parameter": param,
            "year": yr, "old": old, "new": val,
        })
        n += 1
    # Projection.Mode -> 'User defined' (y Projection.Parameter=0 si esta vacio)
    mode_cell = ws.cell(row_idx, COL_MODE)
    if mode_cell.value != "User defined":
        changes.append({
            "scenario": scen, "tech": tech, "parameter": param,
            "field": "Projection.Mode", "old": mode_cell.value, "new": "User defined",
        })
        mode_cell.value = "User defined"
        n += 1
    if ws.cell(row_idx, COL_PROJPARAM).value is None:
        ws.cell(row_idx, COL_PROJPARAM).value = 0
    return n


def process_scenario(scen, floors, do_floors, do_tx, dry_run, changes):
    path = BASE_PATH / f"A1_Outputs_{scen}" / "A-O_Parametrization.xlsx"
    if not path.exists():
        raise FileNotFoundError(path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Secondary Techs"]
    ycols = year_columns(ws)
    n = 0

    if do_floors and scen in FLOORS_SCENARIOS:
        for tech, vals in sorted(floors.items()):
            r = find_row(ws, tech, FLOORS_PARAM)
            n += apply_cells(ws, r, vals, ycols, changes, scen, tech, FLOORS_PARAM)

    if do_tx and scen in TX_SCENARIOS:
        for tech, params in TX_CHANGES.items():
            for param, (y0, val) in params.items():
                r = find_row(ws, tech, param)
                target = {yr: val for yr in range(y0, TX_LAST_YEAR + 1)}
                n += apply_cells(ws, r, target, ycols, changes, scen, tech, param)

    if n and not dry_run:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup = path.with_name(f"{path.stem}_backup_{ts}{path.suffix}")
        shutil.copy2(path, backup)
        wb.save(path)
        print(f"  {scen}: {n} celdas modificadas | backup: {backup.name}")
    else:
        print(f"  {scen}: {n} celdas {'(dry-run, sin guardar)' if dry_run else 'modificadas'}")
    wb.close()
    return n


def verify(floors, do_floors, do_tx):
    """Relee los archivos y confirma los valores finales. Devuelve n errores."""
    errors = []
    for scen in TX_SCENARIOS:
        path = BASE_PATH / f"A1_Outputs_{scen}" / "A-O_Parametrization.xlsx"
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb["Secondary Techs"]
        rows = {}
        hdr = None
        for row in ws.iter_rows(values_only=True):
            if hdr is None:
                hdr = row
                ycols = {int(h): i for i, h in enumerate(hdr) if isinstance(h, (int, float))}
                continue
            key = (row[COL_TECH - 1], row[COL_PARAM - 1])
            rows[key] = row

        def check(tech, param, yr, expected):
            row = rows.get((tech, param))
            if row is None:
                errors.append(f"{scen}: falta fila {tech}/{param}")
                return
            got = row[ycols[yr]]
            if got != expected:
                errors.append(f"{scen} {tech}/{param} {yr}: esperado {expected}, hay {got!r}")
            if row[COL_MODE - 1] != "User defined":
                errors.append(f"{scen} {tech}/{param}: Projection.Mode={row[COL_MODE - 1]!r}")

        if do_floors:
            for tech, vals in floors.items():
                if scen in FLOORS_SCENARIOS:
                    for yr, v in vals.items():
                        check(tech, FLOORS_PARAM, yr, v)
                else:
                    # los pisos NO deben aparecer fuera de OPT: comparamos contra
                    # el valor que el Excel escribiria; si coincide exacto, alerta
                    row = rows.get((tech, FLOORS_PARAM))
                    if row:
                        hits = [yr for yr, v in vals.items() if row[ycols[yr]] == v]
                        if len(hits) == len(vals):
                            errors.append(f"{scen}: {tech} tiene TODOS los pisos del Excel (deberian ser solo de OPT)")
        if do_tx:
            for tech, params in TX_CHANGES.items():
                for param, (y0, val) in params.items():
                    for yr in (y0, y0 + 2, TX_LAST_YEAR):
                        check(tech, param, yr, val)
        wb.close()
    return errors


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[1])
    ap.add_argument("--dry-run", action="store_true", help="No guarda nada; solo reporta")
    ap.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX, help="Excel Nueva_Capacidad_2040+")
    ap.add_argument("--only-floors", action="store_true", help="Solo pisos OPT")
    ap.add_argument("--only-tx", action="store_true", help="Solo interconexiones")
    args = ap.parse_args()

    do_floors = not args.only_tx
    do_tx = not args.only_floors

    floors = read_floors(args.xlsx) if do_floors else {}
    if do_floors:
        ncells = sum(len(v) for v in floors.values())
        total = sum(sum(v.values()) for v in floors.values())
        print(f"Excel: {len(floors)} techs, {ncells} celdas, {total:.2f} GW acumulados")

    changes = []
    total_n = 0
    for scen in TX_SCENARIOS:
        total_n += process_scenario(scen, floors, do_floors, do_tx, args.dry_run, changes)

    if not args.dry_run and changes:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        log_path = BASE_PATH / f"nueva_capacidad_tx_changes_{ts}.json"
        log_path.write_text(json.dumps(changes, indent=2, ensure_ascii=False))
        print(f"Log de cambios: {log_path}")

    if not args.dry_run:
        errs = verify(floors, do_floors, do_tx)
        if errs:
            print("\nVERIFICACION: ERRORES")
            for e in errs:
                print("  -", e)
            return 1
        print("\nVERIFICACION: OK (valores confirmados releyendo los 4 archivos)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
