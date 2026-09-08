#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Parche puntual: ResidualCapacity de RNWTRNBRAXX en A-O_Parametrization.xlsx.

Edita la hoja "Demand Techs" de los escenarios BAU/INV/OPT/VGB:
  - Fila objetivo: Tech == RNWTRNBRAXX y Parameter == ResidualCapacity.
  - Verifica que las columnas I..AJ (2023-2050) tengan el valor plano 52.990319556.
  - Solo si la verificacion pasa, sustituye por 85.40 en esas 28 celdas.

Por defecto corre en DRY-RUN (no guarda nada). Con --apply escribe los cambios,
creando antes un backup <archivo>.bak_<timestamp> junto al original.

Uso:
    python patch_rnwtrnbraxx_residualcapacity.py                 # dry-run, 4 escenarios
    python patch_rnwtrnbraxx_residualcapacity.py --apply         # aplica y guarda
    python patch_rnwtrnbraxx_residualcapacity.py --scenarios BAU OPT
"""

import argparse
import shutil
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))  # -> scripts/
from common import relac_paths as P  # noqa: E402

SCRIPT_DIR = Path(__file__).resolve().parent
SCENARIOS = ["BAU", "INV", "OPT", "VGB"]
SHEET = "Demand Techs"
TECH = "RNWTRNBRAXX"
PARAM = "ResidualCapacity"
OLD_VALUE = 52.990319556
NEW_VALUE = 85.40
COL_INI = column_index_from_string("I")   # anio 2023
COL_FIN = column_index_from_string("AJ")  # anio 2050
TOL = 1e-9

COL_TECH = column_index_from_string("B")   # Tech
COL_PARAM = column_index_from_string("E")  # Parameter
COL_MODE = column_index_from_string("G")   # Projection.Mode


def ruta_escenario(scen):
    return P.scenario_dir(scen) / "A-O_Parametrization.xlsx"   # inputs/A1_Outputs/A1_Outputs_<scen>


def validar_encabezado(ws):
    """La estructura debe ser identica en los 4 archivos; si no, abortar."""
    esperado = [(COL_TECH, "Tech"), (COL_PARAM, "Parameter"), (COL_INI, 2023), (COL_FIN, 2050)]
    for col, valor in esperado:
        real = ws.cell(row=1, column=col).value
        if real != valor:
            raise ValueError(
                f"encabezado inesperado en {get_column_letter(col)}1: "
                f"se esperaba {valor!r} y hay {real!r}"
            )


def fila_objetivo(ws):
    filas = [
        r for r in range(2, ws.max_row + 1)
        if ws.cell(row=r, column=COL_TECH).value == TECH
        and ws.cell(row=r, column=COL_PARAM).value == PARAM
    ]
    if len(filas) != 1:
        raise ValueError(f"se esperaba 1 fila {TECH}/{PARAM} y se encontraron {len(filas)}")
    return filas[0]


def celdas_fuera_de_valor(ws, fila, valor):
    fuera = []
    for col in range(COL_INI, COL_FIN + 1):
        v = ws.cell(row=fila, column=col).value
        if not isinstance(v, (int, float)) or abs(v - valor) > TOL:
            fuera.append((get_column_letter(col), repr(v)))
    return fuera


def procesar(scen, apply_changes):
    path = ruta_escenario(scen)
    if not path.exists():
        print(f"[{scen}] ERROR: no existe {path}")
        return False

    wb = openpyxl.load_workbook(path)
    try:
        ws = wb[SHEET]
        validar_encabezado(ws)
        fila = fila_objetivo(ws)
        modo = ws.cell(row=fila, column=COL_MODE).value

        fuera = celdas_fuera_de_valor(ws, fila, OLD_VALUE)
        if fuera:
            print(f"[{scen}] NO SE SUSTITUYE: {len(fuera)} celda(s) de I..AJ (fila {fila}) "
                  f"no tienen el valor esperado {OLD_VALUE}:")
            for letra, valor in fuera:
                print(f"    {letra}{fila} = {valor}")
            return False

        n = COL_FIN - COL_INI + 1
        print(f"[{scen}] verificacion OK: fila {fila}, {n} celdas I..AJ (2023-2050) "
              f"= {OLD_VALUE} | Projection.Mode = {modo!r}")

        if not apply_changes:
            print(f"[{scen}] DRY-RUN: se sustituiria {OLD_VALUE} -> {NEW_VALUE} (no se guarda nada)")
            return True

        backup = path.with_name(path.name + ".bak_" + datetime.now().strftime("%Y%m%d_%H%M%S"))
        shutil.copy2(path, backup)
        for col in range(COL_INI, COL_FIN + 1):
            ws.cell(row=fila, column=col).value = NEW_VALUE
        wb.save(path)
        print(f"[{scen}] APLICADO: {OLD_VALUE} -> {NEW_VALUE} en {n} celdas | backup: {backup.name}")
    finally:
        wb.close()

    # Releer el archivo guardado para confirmar que el cambio quedo escrito.
    wb2 = openpyxl.load_workbook(path, read_only=True)
    try:
        ws2 = wb2[SHEET]
        fuera = celdas_fuera_de_valor(ws2, fila, NEW_VALUE)
        if fuera:
            print(f"[{scen}] ERROR: tras guardar, {len(fuera)} celda(s) no quedaron en {NEW_VALUE}")
            return False
        print(f"[{scen}] releido y confirmado: I..AJ = {NEW_VALUE}")
        return True
    finally:
        wb2.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--apply", action="store_true",
                        help="escribe y guarda los cambios (por defecto: dry-run)")
    parser.add_argument("--scenarios", nargs="+", choices=SCENARIOS, default=SCENARIOS,
                        metavar="ESC", help=f"subconjunto de escenarios (default: {' '.join(SCENARIOS)})")
    args = parser.parse_args()

    modo = "APPLY" if args.apply else "DRY-RUN"
    print(f"== Parche {TECH}/{PARAM}: {OLD_VALUE} -> {NEW_VALUE} | modo {modo} ==")

    exitosos = [s for s in args.scenarios if procesar(s, args.apply)]
    fallidos = [s for s in args.scenarios if s not in exitosos]

    etiqueta = "efectiva" if args.apply else "posible (dry-run)"
    print(f"\nEscenarios con sustitucion {etiqueta}: {', '.join(exitosos) or 'ninguno'}")
    if fallidos:
        print(f"Escenarios NO modificados: {', '.join(fallidos)}")
    return 0 if not fallidos else 1


if __name__ == "__main__":
    sys.exit(main())
