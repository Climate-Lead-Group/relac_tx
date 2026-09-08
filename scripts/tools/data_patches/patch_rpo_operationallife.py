#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Parche: OperationalLife de las tecnologias RPO en A-O_Parametrization.xlsx.

Edita la hoja "Fixed Horizon Parameters" de los escenarios BAU/INV/OPT/VGB:
  - Filas objetivo: Tech contiene "RPO" y Parameter == OperationalLife.
  - Verifica que TODAS esas filas tengan Value == 20; si alguna difiere,
    el escenario se reporta y NO se modifica.
  - Solo si la verificacion pasa, sustituye 20 -> 50 en la columna H (Value).

Por defecto corre en DRY-RUN (no guarda nada). Con --apply escribe los cambios,
creando antes un backup <archivo>.bak_<timestamp> junto al original.

Uso:
    python patch_rpo_operationallife.py                 # dry-run, 4 escenarios
    python patch_rpo_operationallife.py --apply         # aplica y guarda
    python patch_rpo_operationallife.py --scenarios BAU OPT
"""

import argparse
import shutil
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))  # -> scripts/
from common import relac_paths as P  # noqa: E402

SCRIPT_DIR = Path(__file__).resolve().parent
SCENARIOS = ["BAU", "INV", "OPT", "VGB"]
SHEET = "Fixed Horizon Parameters"
TECH_SUBSTR = "RPO"
PARAM = "OperationalLife"
OLD_VALUE = 20
NEW_VALUE = 50
TOL = 1e-9

COL_TECH = column_index_from_string("C")   # Tech
COL_PARAM = column_index_from_string("F")  # Parameter
COL_VALUE = column_index_from_string("H")  # Value


def ruta_escenario(scen):
    return P.scenario_dir(scen) / "A-O_Parametrization.xlsx"   # inputs/A1_Outputs/A1_Outputs_<scen>


def validar_encabezado(ws):
    """La estructura debe ser identica en los 4 archivos; si no, abortar."""
    esperado = [(COL_TECH, "Tech"), (COL_PARAM, "Parameter"), (COL_VALUE, "Value")]
    for col, valor in esperado:
        real = ws.cell(row=1, column=col).value
        if real != valor:
            raise ValueError(
                f"encabezado inesperado en columna {col} fila 1: "
                f"se esperaba {valor!r} y hay {real!r}"
            )


def filas_objetivo(ws):
    filas = []
    for r in range(2, ws.max_row + 1):
        tech = ws.cell(row=r, column=COL_TECH).value
        param = ws.cell(row=r, column=COL_PARAM).value
        if tech and TECH_SUBSTR in str(tech) and param == PARAM:
            filas.append((r, str(tech)))
    return filas


def filas_fuera_de_valor(ws, filas, valor):
    fuera = []
    for r, tech in filas:
        v = ws.cell(row=r, column=COL_VALUE).value
        if not isinstance(v, (int, float)) or abs(v - valor) > TOL:
            fuera.append((r, tech, repr(v)))
    return fuera


def procesar(scen, apply_changes):
    path = ruta_escenario(scen)
    if not path.exists():
        print(f"[{scen}] ERROR: no existe {path}")
        return False

    wb = openpyxl.load_workbook(path)
    try:
        ws = wb[SHEET]
        validar_encabezado(ws)
        filas = filas_objetivo(ws)
        if not filas:
            print(f"[{scen}] NO SE SUSTITUYE: no hay filas {TECH_SUBSTR}/{PARAM} en '{SHEET}'")
            return False

        fuera = filas_fuera_de_valor(ws, filas, OLD_VALUE)
        if fuera:
            print(f"[{scen}] NO SE SUSTITUYE: {len(fuera)} de {len(filas)} fila(s) "
                  f"no tienen el valor esperado {OLD_VALUE}:")
            for r, tech, valor in fuera:
                print(f"    fila {r} ({tech}): Value = {valor}")
            return False

        print(f"[{scen}] verificacion OK: {len(filas)} filas {TECH_SUBSTR}/{PARAM} con Value = {OLD_VALUE}")

        if not apply_changes:
            print(f"[{scen}] DRY-RUN: se sustituiria {OLD_VALUE} -> {NEW_VALUE} (no se guarda nada)")
            return True

        backup = path.with_name(path.name + ".bak_" + datetime.now().strftime("%Y%m%d_%H%M%S"))
        shutil.copy2(path, backup)
        for r, _tech in filas:
            ws.cell(row=r, column=COL_VALUE).value = NEW_VALUE
        wb.save(path)
        print(f"[{scen}] APLICADO: {OLD_VALUE} -> {NEW_VALUE} en {len(filas)} filas | backup: {backup.name}")
    finally:
        wb.close()

    # Releer el archivo guardado para confirmar que el cambio quedo escrito.
    wb2 = openpyxl.load_workbook(path, read_only=True)
    try:
        ws2 = wb2[SHEET]
        filas2 = filas_objetivo(ws2)
        fuera = filas_fuera_de_valor(ws2, filas2, NEW_VALUE)
        if fuera:
            print(f"[{scen}] ERROR: tras guardar, {len(fuera)} fila(s) no quedaron en {NEW_VALUE}")
            return False
        print(f"[{scen}] releido y confirmado: {len(filas2)} filas con Value = {NEW_VALUE}")
        return True
    finally:
        wb2.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--apply", action="store_true",
                        help="escribe y guarda los cambios (por defecto: dry-run)")
    parser.add_argument("--scenarios", nargs="+", choices=SCENARIOS, default=SCENARIOS,
                        metavar="ESC", help=f"subconjunto de escenarios (default: {' '.join(SCENARIOS)})")
    args = parser.parse_args()

    modo = "APPLY" if args.apply else "DRY-RUN"
    print(f"== Parche {TECH_SUBSTR}*/{PARAM}: {OLD_VALUE} -> {NEW_VALUE} | modo {modo} ==")

    exitosos = [s for s in args.scenarios if procesar(s, args.apply)]
    fallidos = [s for s in args.scenarios if s not in exitosos]

    etiqueta = "efectiva" if args.apply else "posible (dry-run)"
    print(f"\nEscenarios con sustitucion {etiqueta}: {', '.join(exitosos) or 'ninguno'}")
    if fallidos:
        print(f"Escenarios NO modificados: {', '.join(fallidos)}")
    return 0 if not fallidos else 1


if __name__ == "__main__":
    sys.exit(main())
