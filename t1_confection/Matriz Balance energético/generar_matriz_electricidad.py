"""
Script Maestro para Generar Matriz de Electricidad
===================================================

Este script ejecuta todo el proceso completo:
1. Extrae datos de PRODUCCIÓN, IMPORTACIÓN y EXPORTACIÓN del Excel de OLADE
2. Detecta las unidades de cada hoja y convierte a GWh si es necesario
3. Genera la matriz final de electricidad por país y año

IMPORTANTE: El archivo de OLADE tiene diferentes unidades según el año:
- 2021: GWh
- 2023: 10¹⁵ J (PJ - Petajoules)

Conversión: 1 PJ = 277.778 GWh

Uso:
    python generar_matriz_electricidad.py

Archivos generados:
    - Matriz_Completa_Con_Produccion.xlsx (intermedio)
    - Matriz_ImportExport_PorPais.xlsx (resultado final)
"""

import pandas as pd
import openpyxl
import sys
from pathlib import Path

# Factor de conversión
PJ_TO_GWH = 277.778  # 1 PJ = 277.778 GWh

print("=" * 100)
print("SCRIPT MAESTRO: GENERACIÓN DE MATRIZ DE ELECTRICIDAD")
print("=" * 100)
print()

# ============================================================================
# PASO 1: EXTRACCIÓN DE DATOS DEL EXCEL DE OLADE
# ============================================================================

print("PASO 1/2: Extrayendo datos del Excel de OLADE...")
print("-" * 100)

file_path = 'OLADE - Matriz de balance energético - Anual.xlsx'

# Verificar que existe el archivo
if not Path(file_path).exists():
    print(f"❌ ERROR: No se encontró el archivo '{file_path}'")
    print("   Por favor, asegúrate de que el archivo existe en este directorio.")
    sys.exit(1)

wb = openpyxl.load_workbook(file_path)
data_list = []

print(f"Procesando {len(wb.sheetnames)} hojas...\n")

for sheet_name in wb.sheetnames:
    try:
        parts = sheet_name.split(' - ')
        if len(parts) != 2:
            print(f"⚠ Saltando hoja con formato inesperado: '{sheet_name}'")
            continue

        year = parts[0].strip()
        country = parts[1].strip()

        print(f"  Procesando: {country} ({year})...", end=" ")

        df = pd.read_excel(file_path, sheet_name=sheet_name, header=None)

        # Identificar filas clave
        produccion_row = None
        importacion_row = None
        exportacion_row = None

        for idx, val in df[0].items():
            if pd.notna(val):
                val_clean = str(val).strip()
                if 'PRODUCCIÓN' in val_clean.upper():
                    produccion_row = idx
                elif 'IMPORTACIÓN' in val_clean.upper():
                    importacion_row = idx
                elif 'EXPORTACIÓN' in val_clean.upper():
                    exportacion_row = idx

        if produccion_row is None or importacion_row is None or exportacion_row is None:
            print("❌ Faltan filas clave")
            continue

        # Buscar fila de headers
        header_productos_row = None
        for potential_header in range(produccion_row - 1, max(0, produccion_row - 5), -1):
            non_empty = sum(1 for i in range(1, min(30, len(df.columns))) if pd.notna(df.iloc[potential_header, i]))
            if non_empty > 10:
                sample_values = [str(df.iloc[potential_header, i]) for i in range(1, min(10, len(df.columns))) if pd.notna(df.iloc[potential_header, i])]
                if any('PETRÓLEO' in v or 'GAS' in v or 'CARBÓN' in v or 'ELECTRICIDAD' in v for v in sample_values):
                    header_productos_row = potential_header
                    break

        if header_productos_row is None:
            print("❌ No se encontró header")
            continue

        # Fila de unidades (una fila después de los headers)
        units_row = header_productos_row + 1

        # =====================================================================
        # DETECTAR UNIDAD DE ELECTRICIDAD Y FACTOR DE CONVERSIÓN
        # =====================================================================
        elec_col = None
        for col_idx in range(1, len(df.columns)):
            val = df.iloc[header_productos_row, col_idx]
            if pd.notna(val) and 'ELECTRICIDAD' in str(val).upper():
                elec_col = col_idx
                break

        # Determinar factor de conversión para electricidad
        factor_elec = 1.0  # Por defecto, sin conversión
        unidad_elec = "GWh"

        if elec_col is not None:
            unidad_raw = df.iloc[units_row, elec_col]
            if pd.notna(unidad_raw):
                unidad_str = str(unidad_raw).upper()
                # Detectar si está en PJ (10¹⁵ J, PJ, Petajoules)
                if 'PJ' in unidad_str or '10¹⁵' in unidad_str or '10^15' in unidad_str or '1015' in unidad_str:
                    factor_elec = PJ_TO_GWH
                    unidad_elec = "PJ→GWh"

        # Obtener nombres de productos
        productos = []
        for col_idx in range(1, len(df.columns)):
            producto = df.iloc[header_productos_row, col_idx]
            if pd.notna(producto):
                producto_str = str(producto).strip()
                if producto_str and not producto_str.isdigit() and '.' not in producto_str:
                    productos.append(producto_str)
                else:
                    productos.append(f"Producto_Col_{col_idx}")
            else:
                productos.append(f"Producto_Col_{col_idx}")

        # Extraer valores
        produccion = df.iloc[produccion_row, 1:].values
        importaciones = df.iloc[importacion_row, 1:].values
        exportaciones = df.iloc[exportacion_row, 1:].values

        row_data = {'País': country, 'Año': year}

        for i, producto in enumerate(productos):
            # Determinar si este producto es ELECTRICIDAD para aplicar factor
            es_electricidad = 'ELECTRICIDAD' in producto.upper()
            factor = factor_elec if es_electricidad else 1.0

            if i < len(produccion):
                prod_val = produccion[i]
                prod_val = prod_val if pd.notna(prod_val) and isinstance(prod_val, (int, float)) else 0
                row_data[f'PROD_{producto}'] = prod_val * factor

            if i < len(importaciones):
                imp_val = importaciones[i]
                imp_val = imp_val if pd.notna(imp_val) and isinstance(imp_val, (int, float)) else 0
                row_data[f'IMP_{producto}'] = imp_val * factor

            if i < len(exportaciones):
                exp_val = exportaciones[i]
                exp_val = exp_val if pd.notna(exp_val) and isinstance(exp_val, (int, float)) else 0
                row_data[f'EXP_{producto}'] = exp_val * factor

        total_prod = sum([v for k, v in row_data.items() if k.startswith('PROD_') and isinstance(v, (int, float))])
        total_imp = sum([v for k, v in row_data.items() if k.startswith('IMP_') and isinstance(v, (int, float))])
        total_exp = sum([v for k, v in row_data.items() if k.startswith('EXP_') and isinstance(v, (int, float))])

        row_data['TOTAL_PRODUCCIÓN'] = total_prod
        row_data['TOTAL_IMPORTACIONES'] = total_imp
        row_data['TOTAL_EXPORTACIONES'] = total_exp

        data_list.append(row_data)

        # Mostrar información de conversión
        if factor_elec != 1.0:
            print(f"✓ [{unidad_elec}]")
        else:
            print("✓")

    except Exception as e:
        print(f"❌ Error: {str(e)}")
        continue

print()
print("-" * 100)
print(f"✓ Extracción completada: {len(data_list)} registros procesados")
print()

if len(data_list) == 0:
    print("❌ ERROR: No se extrajeron datos")
    sys.exit(1)

# Crear DataFrame consolidado
df_consolidated = pd.DataFrame(data_list)
base_cols = ['País', 'Año', 'TOTAL_PRODUCCIÓN', 'TOTAL_IMPORTACIONES', 'TOTAL_EXPORTACIONES']
other_cols = [col for col in df_consolidated.columns if col not in base_cols]
df_consolidated = df_consolidated[base_cols + other_cols]
df_consolidated = df_consolidated.sort_values(['País', 'Año'])

# Guardar archivo intermedio
intermediate_file = 'Matriz_Completa_Con_Produccion.xlsx'
df_consolidated.to_excel(intermediate_file, index=False, sheet_name='Consolidado')
print(f"✓ Archivo intermedio guardado: {intermediate_file}")
print(f"  Dimensiones: {df_consolidated.shape[0]} filas x {df_consolidated.shape[1]} columnas")
print()

# ============================================================================
# PASO 2: GENERAR MATRIZ DE ELECTRICIDAD
# ============================================================================

print("PASO 2/2: Generando matriz de electricidad...")
print("-" * 100)

# Extraer columnas de electricidad
columnas_electricidad = ['País', 'Año']

if 'PROD_ELECTRICIDAD' in df_consolidated.columns:
    columnas_electricidad.append('PROD_ELECTRICIDAD')
    print("✓ Columna PROD_ELECTRICIDAD encontrada")

if 'IMP_ELECTRICIDAD' in df_consolidated.columns:
    columnas_electricidad.append('IMP_ELECTRICIDAD')
    print("✓ Columna IMP_ELECTRICIDAD encontrada")

if 'EXP_ELECTRICIDAD' in df_consolidated.columns:
    columnas_electricidad.append('EXP_ELECTRICIDAD')
    print("✓ Columna EXP_ELECTRICIDAD encontrada")

df_electricidad = df_consolidated[columnas_electricidad].copy()

# Calcular balance
if 'IMP_ELECTRICIDAD' in df_consolidated.columns and 'EXP_ELECTRICIDAD' in df_consolidated.columns:
    df_electricidad['BALANCE_ELECTRICIDAD'] = (
        df_electricidad['EXP_ELECTRICIDAD'] - df_electricidad['IMP_ELECTRICIDAD']
    )
    print("✓ Columna BALANCE_ELECTRICIDAD calculada")

# Renombrar columnas
df_electricidad = df_electricidad.rename(columns={
    'PROD_ELECTRICIDAD': 'Producción_Electricidad_GWh',
    'IMP_ELECTRICIDAD': 'Importación_Electricidad_GWh',
    'EXP_ELECTRICIDAD': 'Exportación_Electricidad_GWh',
    'BALANCE_ELECTRICIDAD': 'Balance_Neto_GWh'
})

df_electricidad = df_electricidad.sort_values(['País', 'Año'])

# Crear resumen por país
df_resumen = df_electricidad.groupby('País').agg({
    'Producción_Electricidad_GWh': 'sum',
    'Importación_Electricidad_GWh': 'sum',
    'Exportación_Electricidad_GWh': 'sum',
    'Balance_Neto_GWh': 'sum'
}).reset_index()

df_resumen = df_resumen.sort_values('Balance_Neto_GWh', ascending=False)

# Guardar archivo final
output_file = 'Matriz_ImportExport_PorPais.xlsx'
with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
    df_electricidad.to_excel(writer, sheet_name='Matriz_Completa', index=False)
    df_resumen.to_excel(writer, sheet_name='Resumen_Por_País', index=False)

print()
print("-" * 100)
print(f"✓ Archivo final guardado: {output_file}")
print(f"  - Hoja 'Matriz_Completa': {df_electricidad.shape[0]} filas x {df_electricidad.shape[1]} columnas")
print(f"  - Hoja 'Resumen_Por_País': {df_resumen.shape[0]} filas x {df_resumen.shape[1]} columnas")
print()

# ============================================================================
# RESUMEN DE RESULTADOS
# ============================================================================

print("=" * 100)
print("RESUMEN DE RESULTADOS")
print("=" * 100)
print()

print(f"Total de países procesados: {len(df_electricidad['País'].unique())}")
print(f"Años disponibles: {', '.join(sorted(df_electricidad['Año'].unique()))}")
print()

print("NOTA: Todos los valores están en GWh")
print("      Los datos de 2023 (originalmente en PJ) fueron convertidos usando: 1 PJ = 277.778 GWh")
print()

print("Top 5 productores de electricidad:")
top_productores = df_resumen.nlargest(5, 'Producción_Electricidad_GWh')
for i, (_, row) in enumerate(top_productores.iterrows(), 1):
    print(f"  {i}. {row['País']:25s} {row['Producción_Electricidad_GWh']:>15,.2f} GWh")

print()
print("Top 5 exportadores netos de electricidad:")
exportadores = df_resumen[df_resumen['Balance_Neto_GWh'] > 0].nlargest(5, 'Balance_Neto_GWh')
for i, (_, row) in enumerate(exportadores.iterrows(), 1):
    print(f"  {i}. {row['País']:25s} {row['Balance_Neto_GWh']:>15,.2f} GWh")

print()
print("Top 5 importadores netos de electricidad:")
importadores = df_resumen[df_resumen['Balance_Neto_GWh'] < 0].nsmallest(5, 'Balance_Neto_GWh')
for i, (_, row) in enumerate(importadores.iterrows(), 1):
    print(f"  {i}. {row['País']:25s} {row['Balance_Neto_GWh']:>15,.2f} GWh")

print()
print("=" * 100)
print("✓ PROCESO COMPLETADO EXITOSAMENTE")
print("=" * 100)
print()
print(f"Archivo final: {output_file}")
print()
