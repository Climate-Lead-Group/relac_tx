"""
sync_historical_from_bau.py
===========================

Copia los valores de los anos historicos (2023, 2024, 2025) desde el escenario
BAU hacia los escenarios INV y OPT en A-O_Parametrization.xlsx.

Motivacion
----------
Los anos 2023-2025 son historicos/observados y deben ser identicos en todos los
escenarios. Solo a partir de 2026 los escenarios divergen (es ahi donde A3 y D4
aplican sus reglas). Este script garantiza el punto de partida comun: para cada
hoja con columnas de ano, copia las celdas 2023/2024/2025 de BAU a INV/OPT.

Alcance
-------
- Solo A-O_Parametrization.xlsx (no toca Demand, Projections, etc.).
- Todas las hojas que tengan columnas de ano (encabezado int 2023 o string
  "2023"). Las hojas sin columnas de ano (p.ej. 'Fixed Horizon Parameters')
  se omiten automaticamente.
- Solo los anos en --years (default 2023,2024,2025).
- EXCEPCION: en la hoja 'Demand Techs', las filas de
  TotalAnnualMaxCapacityInvestment de tecnologias de transmision (TRN*) NO se
  copian para anos >= 2026. Ahi manda el tope que escribe
  D4_load_dsptrn_max_cap_inv.py (la Transmision de INV diverge de BAU desde
  2026). Los anos <= 2025 (historicos) si se copian. OPT solo sincroniza
  2023-2025, asi que esta excepcion nunca le aplica.
- EXCEPCION GENERICA (--protect-params): los parametros nombrados en
  --protect-params NO se copian para anos >= 2026 (en cualquier hoja). Sirve
  para escenarios que divergen deliberadamente de BAU en un parametro desde
  2026 y cuya divergencia debe sobrevivir al sync (que de otro modo lo
  sobre-escribe en la ventana sincronizada). Los anos <= 2025 (historicos)
  siempre se copian. Ej.: VGB protege TotalTechnologyAnnualActivityLowerLimit
  (piso solo en renovables) y TotalAnnualMinCapacityInvestment (PEGs = OPT).

Seguridad
---------
- Las tres workbooks (BAU/INV/OPT) tienen el mismo numero de filas por hoja, asi
  que la copia es POSICIONAL. Antes de copiar se verifica que (a) la hoja exista
  en ambos, (b) el numero de filas y el encabezado coincidan, y (c) la columna
  'Parameter' coincida fila a fila (guard contra desalineacion). Si algo no
  cuadra, el script aborta sin escribir.
- --apply hace backup timestamped de cada xlsx destino antes de escribir.

Uso
---
    python t1_confection/sync_historical_from_bau.py --dry-run     # preview (default)
    python t1_confection/sync_historical_from_bau.py --apply
    python t1_confection/sync_historical_from_bau.py --apply --scenarios INV
    python t1_confection/sync_historical_from_bau.py --years 2023,2024,2025
"""
from __future__ import annotations

import argparse
import shutil
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
A1_OUTPUTS = HERE / "A1_Outputs"
PARAM_FILENAME = "A-O_Parametrization.xlsx"
SOURCE_SCENARIO = "BAU"
DEFAULT_TARGET_SCENARIOS = ["INV", "OPT"]
DEFAULT_YEARS = [2023, 2024, 2025]

# Transmission MaxCapInv divergence. D4_load_dsptrn_max_cap_inv.py escribe el
# tope anual de transmision para INV desde TX_DIVERGE_YEAR en adelante (cap =
# NewCapacity_BAU * factor). Esas celdas NO deben volver a BAU por este sync, o
# INV perderia su tope de transmision. Por eso, para anos >= TX_DIVERGE_YEAR se
# saltan las filas de transmision (TX_PARAM sobre techs TX_PREFIXES) de la hoja
# TX_SHEET. Los anos < TX_DIVERGE_YEAR (historicos) siempre se copian. OPT solo
# sincroniza 2023-2025, asi que esta exclusion nunca le aplica.
TX_DIVERGE_YEAR = 2026
TX_SHEET = "Demand Techs"
TX_PARAM = "TotalAnnualMaxCapacityInvestment"
TX_PREFIXES = ("PWRTRN", "TRNNLI", "TRNRPO", "RNWTRN", "RNWRPO", "RNWNLI")

# Generic per-parameter protection (--protect-params). Rows whose 'Parameter'
# is in the protect set are NOT copied from BAU for years >= this boundary, on
# ANY sheet. Same divergence boundary as transmission (2026): years <= 2025 are
# historical and always synced.
PROTECT_PARAMS_FROM_YEAR = TX_DIVERGE_YEAR


def param_path(scenario: str) -> Path:
    return A1_OUTPUTS / f"A1_Outputs_{scenario}" / PARAM_FILENAME


def year_columns(ws, years: set[int]) -> dict[int, int]:
    """Return {year: column_index_1based} for headers in row 1 matching `years`.

    Accepts both integer headers (2023) and string-digit headers ("2023"),
    since some sheets (Capacities, Yearsplit, DaySplit) store year headers as
    strings while others (Secondary Techs, ...) store them as ints.
    """
    found: dict[int, int] = {}
    for col_idx, cell in enumerate(next(ws.iter_rows(min_row=1, max_row=1)), start=1):
        v = cell.value
        y = None
        if isinstance(v, int):
            y = v
        elif isinstance(v, str) and v.strip().isdigit():
            y = int(v.strip())
        if y is not None and y in years:
            found[y] = col_idx
    return found


def header_row(ws) -> list:
    return [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]


def find_param_col(header: list) -> int | None:
    """1-based index of the 'Parameter' column in the header, or None."""
    for i, h in enumerate(header, start=1):
        if isinstance(h, str) and h.strip() == "Parameter":
            return i
    return None


def find_tech_col(header: list) -> int | None:
    """1-based index of the 'Tech' column in the header, or None."""
    for i, h in enumerate(header, start=1):
        if isinstance(h, str) and h.strip() == "Tech":
            return i
    return None


def sync_scenario(src_ws_map, src_wb, tgt_path: Path, years: set[int],
                  apply_changes: bool, protect_params: set[str] | None = None) -> dict:
    """Copy year cells from BAU into one target scenario workbook.

    `src_ws_map` is {sheet_name: source_worksheet}. `protect_params` is a set of
    Parameter names whose rows are NOT copied for years >= PROTECT_PARAMS_FROM_YEAR
    (on any sheet); historical years are still synced. Returns a summary dict.
    Raises ValueError on any structural mismatch (caller aborts the run).
    """
    protect_params = protect_params or set()
    tgt_wb = openpyxl.load_workbook(tgt_path)
    summary = {"path": str(tgt_path), "sheets": [], "total_cells": 0}

    try:
        for sheet in src_wb.sheetnames:
            if sheet not in tgt_wb.sheetnames:
                raise ValueError(
                    f"{tgt_path.name}: sheet '{sheet}' present in BAU but missing here."
                )
            src_ws = src_ws_map[sheet]
            tgt_ws = tgt_wb[sheet]

            src_cols = year_columns(src_ws, years)
            tgt_cols = year_columns(tgt_ws, years)
            if not src_cols:
                summary["sheets"].append({"sheet": sheet, "skipped": "no year columns"})
                continue
            if src_cols != tgt_cols:
                raise ValueError(
                    f"{tgt_path.name}/{sheet}: year-column layout differs from BAU "
                    f"(BAU={src_cols}, target={tgt_cols})."
                )
            if src_ws.max_row != tgt_ws.max_row:
                raise ValueError(
                    f"{tgt_path.name}/{sheet}: row count differs from BAU "
                    f"(BAU={src_ws.max_row}, target={tgt_ws.max_row})."
                )
            if header_row(src_ws) != header_row(tgt_ws):
                raise ValueError(
                    f"{tgt_path.name}/{sheet}: header row differs from BAU."
                )

            param_col = find_param_col(header_row(src_ws))
            tech_col = find_tech_col(header_row(src_ws))
            is_tx_sheet = sheet == TX_SHEET
            cells_changed = 0
            cells_protected = 0
            for row in range(2, src_ws.max_row + 1):
                # Guard against row misalignment: the Parameter column (when
                # present) must match between BAU and target for this row.
                if param_col is not None:
                    sp = src_ws.cell(row=row, column=param_col).value
                    tp = tgt_ws.cell(row=row, column=param_col).value
                    if sp != tp:
                        raise ValueError(
                            f"{tgt_path.name}/{sheet} row {row}: Parameter mismatch "
                            f"(BAU={sp!r}, target={tp!r}); aborting to avoid "
                            f"misaligned copy."
                        )
                # Protect D4's transmission MaxCapInv cap from being snapped back
                # to BAU in 2026+. Only the TX_PARAM rows of TX_PREFIXES techs on
                # the TX_SHEET are shielded, and only for years >= TX_DIVERGE_YEAR.
                protect_tx = False
                if is_tx_sheet and tech_col is not None and param_col is not None:
                    tv = src_ws.cell(row=row, column=tech_col).value
                    pv = src_ws.cell(row=row, column=param_col).value
                    if (isinstance(tv, str) and tv.strip().startswith(TX_PREFIXES)
                            and isinstance(pv, str) and pv.strip() == TX_PARAM):
                        protect_tx = True
                # Generic parameter protection: shield deliberately-diverged
                # parameters (e.g. VGB's LowerLimit / PEGs) from 2026 onward.
                protect_param = False
                if protect_params and param_col is not None:
                    pv2 = src_ws.cell(row=row, column=param_col).value
                    if isinstance(pv2, str) and pv2.strip() in protect_params:
                        protect_param = True
                for y, col in src_cols.items():
                    if protect_tx and y >= TX_DIVERGE_YEAR:
                        cells_protected += 1
                        continue
                    if protect_param and y >= PROTECT_PARAMS_FROM_YEAR:
                        cells_protected += 1
                        continue
                    sv = src_ws.cell(row=row, column=col).value
                    tc = tgt_ws.cell(row=row, column=col)
                    if tc.value != sv:
                        tc.value = sv
                        cells_changed += 1

            summary["sheets"].append({
                "sheet": sheet,
                "years": sorted(src_cols),
                "cells_changed": cells_changed,
                "cells_protected": cells_protected,
            })
            summary["total_cells"] += cells_changed

        if apply_changes:
            ts = datetime.now().strftime("%Y%m%d_%H%M%S")
            backup = tgt_path.with_name(
                f"{tgt_path.stem}.backup_pre_sync_hist_{ts}{tgt_path.suffix}"
            )
            shutil.copy2(tgt_path, backup)
            summary["backup"] = backup.name
            tgt_wb.save(tgt_path)
    finally:
        tgt_wb.close()

    return summary


def main() -> int:
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    grp = ap.add_mutually_exclusive_group()
    grp.add_argument("--apply", action="store_true", help="escribe cambios (crea backup)")
    grp.add_argument("--dry-run", action="store_true", help="preview sin escribir (default)")
    ap.add_argument(
        "--scenarios", default=",".join(DEFAULT_TARGET_SCENARIOS),
        help=f"escenarios destino, separados por coma (default: "
             f"{','.join(DEFAULT_TARGET_SCENARIOS)})",
    )
    ap.add_argument(
        "--years", default=",".join(str(y) for y in DEFAULT_YEARS),
        help=f"anos a copiar, separados por coma (default: "
             f"{','.join(str(y) for y in DEFAULT_YEARS)})",
    )
    ap.add_argument(
        "--protect-params", default="",
        help="parametros (separados por coma) que NO se copian desde BAU para "
             f"anos >= {PROTECT_PARAMS_FROM_YEAR}, en cualquier hoja. Los anos "
             "historicos (<= 2025) siempre se copian.",
    )
    args = ap.parse_args()

    apply_changes = args.apply and not args.dry_run
    mode = "APPLY" if apply_changes else "DRY-RUN"
    targets = [s.strip() for s in args.scenarios.split(",") if s.strip()]
    years = {int(y.strip()) for y in args.years.split(",") if y.strip()}
    protect_params = {p.strip() for p in args.protect_params.split(",") if p.strip()}

    print(f"Mode      : {mode}")
    print(f"Source    : {SOURCE_SCENARIO}")
    print(f"Targets   : {targets}")
    print(f"Years     : {sorted(years)}")
    if protect_params:
        print(f"Protect   : {sorted(protect_params)} (not synced for years "
              f">= {PROTECT_PARAMS_FROM_YEAR})")

    src_path = param_path(SOURCE_SCENARIO)
    if not src_path.exists():
        sys.exit(f"ERROR: no existe source {src_path}")
    src_wb = openpyxl.load_workbook(src_path, data_only=False)
    src_ws_map = {sh: src_wb[sh] for sh in src_wb.sheetnames}

    try:
        grand_total = 0
        for scen in targets:
            tgt_path = param_path(scen)
            if not tgt_path.exists():
                sys.exit(f"ERROR: no existe target {tgt_path}")
            print(f"\n=== {scen} <- {SOURCE_SCENARIO} ===")
            summary = sync_scenario(src_ws_map, src_wb, tgt_path, years,
                                     apply_changes, protect_params)
            for s in summary["sheets"]:
                if "skipped" in s:
                    print(f"  [skip] {s['sheet']:28s} ({s['skipped']})")
                else:
                    prot = s.get("cells_protected", 0)
                    prot_str = f" protected={prot}" if prot else ""
                    print(f"  {s['sheet']:28s} years={s['years']} "
                          f"cells_changed={s['cells_changed']}{prot_str}")
            print(f"  TOTAL cells changed: {summary['total_cells']}")
            if apply_changes and "backup" in summary:
                print(f"  Backup: {summary['backup']}")
            grand_total += summary["total_cells"]

        print(f"\nGran total de celdas {'escritas' if apply_changes else 'a escribir'}: "
              f"{grand_total}")
        if not apply_changes:
            print("[DRY-RUN] no se guardaron cambios. Usa --apply para escribir.")
    finally:
        src_wb.close()

    return 0


if __name__ == "__main__":
    sys.exit(main())
