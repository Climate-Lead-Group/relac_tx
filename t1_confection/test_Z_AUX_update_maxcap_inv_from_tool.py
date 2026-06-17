"""
Test de Z_AUX_update_maxcap_inv_from_tool.py (fuentes A LOGISTIC + B Updated).

Tres comprobaciones, todas con escaneo INDEPENDIENTE de las fuentes (no usan los
lectores del script para construir lo "esperado", asi un bug del lector no puede
auto-enmascararse):

  (R) READERS: lo que read_input_table()/read_complementary_table() del script
      extraen coincide con el escaneo independiente de cada fuente.

  (1) ARCHIVO REAL: el A-O_Parametrization.xlsx real de BAU (tras --apply)
      contiene, en las 186 techs del merge (12 de A + 174 de B), exactamente los
      valores de la fuente correspondiente y Projection.Mode copiado.

  (2) NO-INTERFERENCIA (sandbox, con perturbacion): sobre una COPIA del BAU real
      se "ensucian" las celdas objetivo con un centinela; al correr el writer del
      script debe (a) reescribirlas con los valores del merge y (b) NO tocar
      ninguna otra celda de la hoja. Robusto aunque el real ya este aplicado.

Merge esperado: techs de A con valores de A, MAS techs de B que no estan en A con
valores de B (A tiene precedencia en las coincidentes).

Usage:
    python t1_confection/test_Z_AUX_update_maxcap_inv_from_tool.py
"""
import shutil
import sys
import tempfile
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import Z_AUX_update_maxcap_inv_from_tool as mod  # noqa: E402

TOOL_A = HERE / 'LAC_maxcap_tool.xlsx'
TOOL_B = HERE / 'LAC_maxcap_tool_complementary.xlsx'
REAL_BAU = mod.A1_OUTPUTS / 'A1_Outputs_BAU' / 'A-O_Parametrization.xlsx'
PARAM = mod.PARAM
TOL = 1e-9
SENTINEL_NUM = -987654.0
SENTINEL_PM = '__SENTINEL__'


# --------------------------------------------------------------------------- #
# Escaneo independiente de las fuentes                                         #
# --------------------------------------------------------------------------- #
def _scan(path, sheet, find_logistic):
    """Escaneo propio -> {tech: {proj_mode, years}} para Parameter == PARAM."""
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        ws = wb[sheet]
        if find_logistic:
            title = None
            for r in range(1, ws.max_row + 1):
                v = ws.cell(r, 1).value
                if isinstance(v, str):
                    u = v.upper()
                    if 'LOGISTIC' in u and 'HEADROOM' in u and 'SATURATING CAP' in u \
                            and 'GEOMETRIC' not in u and 'COMPOUNDING' not in u:
                        title = r
                        break
            assert title is not None, 'no se hallo el titulo LOGISTIC'
            header_row = next(r for r in range(title + 1, title + 6)
                              if str(ws.cell(r, 1).value).strip() == 'Tech.ID')
            stop_at_blank = True
        else:
            header_row = next(r for r in range(1, 11)
                              if str(ws.cell(r, 1).value).strip() == 'Tech.ID')
            stop_at_blank = False

        hdr = {str(ws.cell(header_row, c).value).strip(): c
               for c in range(1, ws.max_column + 1)
               if ws.cell(header_row, c).value is not None}
        tcol, pcol, gcol = hdr['Tech'], hdr['Parameter'], hdr['Projection.Mode']
        ycols = {ws.cell(header_row, c).value: c
                 for c in range(1, ws.max_column + 1)
                 if isinstance(ws.cell(header_row, c).value, int)
                 and 2000 <= ws.cell(header_row, c).value <= 2100}

        table = {}
        for r in range(header_row + 1, ws.max_row + 1):
            t = ws.cell(r, tcol).value
            if t is None or str(t).strip() == '':
                if stop_at_blank:
                    break
                continue
            if str(ws.cell(r, pcol).value).strip() != PARAM:
                continue
            years = {}
            for y, c in ycols.items():
                val = ws.cell(r, c).value
                if val is None or (isinstance(val, str) and val.strip() == ''):
                    continue
                assert not isinstance(val, str), f'celda no cacheada {t} {y} = {val!r}'
                years[y] = float(val)
            table[str(t).strip()] = {'proj_mode': ws.cell(r, gcol).value, 'years': years}
        return table
    finally:
        wb.close()


def build_expected():
    tA = _scan(TOOL_A, 'Export', find_logistic=True)
    tB = _scan(TOOL_B, mod.COMPLEMENTARY_SHEET, find_logistic=False)
    merged = dict(tA)  # A precede
    for tech, payload in tB.items():
        if tech not in merged:
            merged[tech] = payload
    return tA, tB, merged


# --------------------------------------------------------------------------- #
# Lectura de la hoja Secondary Techs                                           #
# --------------------------------------------------------------------------- #
def load_sheet(path):
    """Devuelve (snap{(r,c):val}, ycols{año:col}, pmcol, rowidx{tech:fila})."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb['Secondary Techs']
        snap = {}
        header = None
        for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if r == 1:
                header = row
            for c, val in enumerate(row, start=1):
                snap[(r, c)] = val
    finally:
        wb.close()
    ycols, pmcol = {}, None
    for c, v in enumerate(header, start=1):
        if isinstance(v, int) and 2000 <= v <= 2100:
            ycols[v] = c
        elif v is not None and str(v).strip() == 'Projection.Mode':
            pmcol = c
    maxr = max(r for (r, _c) in snap)
    rowidx = {}
    for r in range(2, maxr + 1):
        t, p = snap.get((r, 2)), snap.get((r, 5))
        if isinstance(t, str) and isinstance(p, str) and p.strip() == PARAM:
            rowidx[t.strip()] = r
    return snap, ycols, pmcol, rowidx


def is_empty(v):
    return v is None or (isinstance(v, str) and v.strip() == '')


def vals_equal(a, b):
    if is_empty(a) and is_empty(b):
        return True
    if is_empty(a) != is_empty(b):
        return False
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) <= TOL * max(1.0, abs(a), abs(b))
    return a == b


# --------------------------------------------------------------------------- #
# Tests                                                                        #
# --------------------------------------------------------------------------- #
def test_readers(merged_expected, tA_exp, tB_exp, fails):
    tA_mod, _, _, _ = mod.read_input_table(TOOL_A)
    tB_mod, _, _ = mod.read_complementary_table(TOOL_B, mod.COMPLEMENTARY_SHEET)
    for name, got, exp in (('A', tA_mod, tA_exp), ('B', tB_mod, tB_exp)):
        if set(got) != set(exp):
            fails.append(f'(R) reader {name}: techs difieren ({len(got)} vs {len(exp)})')
        for tech in set(got) & set(exp):
            if mod._norm(got[tech]['proj_mode']) != mod._norm(exp[tech]['proj_mode']):
                fails.append(f'(R) reader {name} proj_mode {tech}')
            if set(got[tech]['years']) != set(exp[tech]['years']):
                fails.append(f'(R) reader {name} años {tech}')
            for y, v in exp[tech]['years'].items():
                gv = got[tech]['years'].get(y)
                if gv is None or abs(gv - v) > TOL * max(1.0, abs(v)):
                    fails.append(f'(R) reader {name} valor {tech} {y}: {gv!r} != {v}')


def test_real_file(merged_expected, fails):
    snap, ycols, pmcol, rowidx = load_sheet(REAL_BAU)
    for tech, payload in merged_expected.items():
        r = rowidx.get(tech)
        if r is None:
            fails.append(f'(1) tech sin fila en BAU real: {tech}')
            continue
        if mod._norm(snap[(r, pmcol)]) != mod._norm(payload['proj_mode']):
            fails.append(f'(1) proj_mode {tech}: {snap[(r, pmcol)]!r} != {payload["proj_mode"]!r}')
        for y, v in payload['years'].items():
            got = snap[(r, ycols[y])]
            if not isinstance(got, (int, float)) or abs(got - v) > TOL * max(1.0, abs(v)):
                fails.append(f'(1) valor {tech} {y}: {got!r} != {v}')


def test_non_interference(merged_expected, fails):
    tmp = Path(tempfile.mkdtemp(prefix='zaux_maxcap_'))
    try:
        sandbox = tmp / 'A1_Outputs' / 'A1_Outputs_BAU' / 'A-O_Parametrization.xlsx'
        sandbox.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(REAL_BAU, sandbox)

        _, ycols, pmcol, rowidx = load_sheet(sandbox)

        # celdas objetivo + valor esperado
        allowed_expected = {}
        for tech, payload in merged_expected.items():
            r = rowidx.get(tech)
            if r is None:
                fails.append(f'(2) cobertura: tech sin fila {tech}')
                continue
            allowed_expected[(r, pmcol)] = payload['proj_mode']
            for y, v in payload['years'].items():
                allowed_expected[(r, ycols[y])] = v

        # perturbar SOLO las celdas objetivo con un centinela
        wb = openpyxl.load_workbook(sandbox)
        ws = wb['Secondary Techs']
        for (r, c) in allowed_expected:
            ws.cell(r, c).value = SENTINEL_PM if c == pmcol else SENTINEL_NUM
        wb.save(sandbox)
        wb.close()

        before = load_sheet(sandbox)[0]

        # correr el writer del script (con la tabla que produce el propio script)
        mod.A1_OUTPUTS = tmp / 'A1_Outputs'
        tA_mod, _, _, _ = mod.read_input_table(TOOL_A)
        tB_mod, _, _ = mod.read_complementary_table(TOOL_B, mod.COMPLEMENTARY_SHEET)
        merged_mod = dict(tA_mod)
        for tech, payload in tB_mod.items():
            if tech not in merged_mod:
                merged_mod[tech] = payload
        mod.write_scenario('BAU', merged_mod, apply_changes=True, do_backup=False)

        after = load_sheet(sandbox)[0]

        # (a) celdas objetivo reescritas con el valor del merge
        for key, want in allowed_expected.items():
            got = after[key]
            if key[1] == pmcol:
                if mod._norm(got) != mod._norm(want):
                    fails.append(f'(2a) proj_mode celda {key}: {got!r} != {want!r}')
            else:
                if not isinstance(got, (int, float)) or abs(got - want) > TOL * max(1.0, abs(want)):
                    fails.append(f'(2a) valor celda {key}: {got!r} != {want}')

        # (b) ninguna otra celda cambio respecto al sandbox perturbado
        outside = 0
        for key in set(before) | set(after):
            if key in allowed_expected:
                continue
            if not vals_equal(before.get(key), after.get(key)):
                outside += 1
                if outside <= 15:
                    fails.append(f'(2b) celda {key} cambio: {before.get(key)!r} -> {after.get(key)!r}')
        if outside:
            fails.append(f'(2b) TOTAL celdas cambiadas fuera del set permitido: {outside}')
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def run_tests():
    for p in (TOOL_A, TOOL_B, REAL_BAU):
        if not p.exists():
            print(f'SKIP: no existe {p}')
            return 0

    tA_exp, tB_exp, merged_expected = build_expected()
    print(f'[info] esperado: A={len(tA_exp)} techs, B={len(tB_exp)} techs, '
          f'merge={len(merged_expected)} (B-only={len(merged_expected) - len(tA_exp)}).')

    fails = []
    test_readers(merged_expected, tA_exp, tB_exp, fails)
    test_real_file(merged_expected, fails)
    test_non_interference(merged_expected, fails)

    print('\n' + '=' * 60)
    if fails:
        print(f'TESTS FAILED: {len(fails)} problema(s):')
        for f in fails:
            print(f'  - {f}')
        return 1
    print('ALL TESTS PASSED')
    return 0


if __name__ == '__main__':
    sys.exit(run_tests())
