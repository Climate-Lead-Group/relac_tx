"""
Verification tests for D2_update_secondary_techs.py fixes.

Reads A-O_Parametrization.xlsx output files and verifies that each fix
produces the expected results. Run after D2_update_secondary_techs.py.

Usage:
    python t1_confection/test_D2_fixes.py
"""
import sys
import openpyxl
from pathlib import Path
from collections import defaultdict


# Expected LowerLimit sums (PJ) for year 2025, from diagnostic report (MAIN_BAU reference)
EXPECTED_LOWER_LIMITS_2025 = {
    'NGS': 1370,
    'HYD': 2600,
    'BIO': 191,
    'OIL': 412,
    'URN': 112,
}

TOLERANCE_PCT = 0.20  # 20% tolerance for informational comparison


def find_scenarios(base_path):
    """Find all A1_Outputs_{scenario} folders."""
    outputs_dir = base_path / "A1_Outputs"
    scenarios = []
    if outputs_dir.exists():
        for item in sorted(outputs_dir.iterdir()):
            if item.is_dir() and item.name.startswith("A1_Outputs_"):
                suffix = item.name.split("A1_Outputs_", 1)[1]
                if suffix:
                    scenarios.append(suffix)
    return scenarios


def read_secondary_techs(param_path):
    """Read Secondary Techs sheet and return structured data."""
    wb = openpyxl.load_workbook(param_path, data_only=True)
    if 'Secondary Techs' not in wb.sheetnames:
        wb.close()
        return None, None, None

    ws = wb['Secondary Techs']

    # Build year column map
    year_col_map = {}
    projection_mode_col = None
    headers = [cell.value for cell in ws[1]]
    for col_idx, header in enumerate(headers, 1):
        if header:
            if str(header).isdigit():
                try:
                    year = int(header)
                    if 2000 <= year <= 2100:
                        year_col_map[year] = col_idx
                except (ValueError, TypeError):
                    pass
            elif str(header).strip() == "Projection.Mode":
                projection_mode_col = col_idx

    # Read all rows
    rows = []
    for row_idx in range(2, ws.max_row + 1):
        tech = ws.cell(row_idx, 2).value
        param = ws.cell(row_idx, 5).value
        proj_mode = ws.cell(row_idx, projection_mode_col).value if projection_mode_col else None

        if not tech or not param:
            continue

        year_values = {}
        for year, col_idx in year_col_map.items():
            year_values[year] = ws.cell(row_idx, col_idx).value

        rows.append({
            'row': row_idx,
            'tech': str(tech).strip(),
            'param': str(param).strip(),
            'proj_mode': str(proj_mode).strip() if proj_mode else '',
            'year_values': year_values,
        })

    wb.close()
    return rows, year_col_map, projection_mode_col


def test_fix1_investment_projection_mode(rows):
    """TEST 1: Investment parameters have correct Projection.Mode."""
    failures = []

    for r in rows:
        tech = r['tech']
        param = r['param']
        mode = r['proj_mode']

        # Rule 1: TotalAnnualMaxCapacityInvestment for PWR -> EMPTY
        if param == 'TotalAnnualMaxCapacityInvestment' and tech.startswith('PWR'):
            if mode != 'EMPTY':
                failures.append(f"Row {r['row']}: {tech} {param} has Projection.Mode='{mode}', expected 'EMPTY'")

        # Rule 2: TotalAnnualMinCapacityInvestment with all empty years -> EMPTY
        if param == 'TotalAnnualMinCapacityInvestment':
            has_value = any(
                v is not None and v != '' and v != 0 and v != 0.0
                for v in r['year_values'].values()
            )
            if not has_value and mode != 'EMPTY':
                failures.append(f"Row {r['row']}: {tech} {param} has all empty years but Projection.Mode='{mode}', expected 'EMPTY'")
            elif has_value and mode != 'User defined':
                failures.append(f"Row {r['row']}: {tech} {param} has values but Projection.Mode='{mode}', expected 'User defined'")

    return failures


def test_fix3_was_residual_capacity(rows):
    """TEST 2: WAS ResidualCapacity preserved.

    Only checks countries that have WAS data in Old_Inputs (workflow 1).
    Countries without WAS capacity in the source data are not flagged.
    """
    # Countries with WAS data in Old_Inputs
    countries_with_was = {'BRA', 'MEX', 'NIC', 'COL', 'PAN', 'PER', 'ECU', 'SLV', 'ARG', 'GTM', 'HND'}

    failures = []
    was_total = 0.0
    bra_value = None

    for r in rows:
        tech = r['tech'].upper()
        param = r['param']

        if tech.startswith('PWRWAS') and param == 'ResidualCapacity':
            country = tech[6:9]  # PWRWASBRAXX -> BRA

            # Check that countries WITH source data have positive values
            has_value = any(
                v is not None and isinstance(v, (int, float)) and v > 0
                for v in r['year_values'].values()
            )
            if not has_value and country in countries_with_was:
                failures.append(f"Row {r['row']}: {r['tech']} ResidualCapacity has no positive values (expected from Old_Inputs)")

            # Check Brazil specifically
            if country == 'BRA':
                val_2023 = r['year_values'].get(2023)
                if val_2023 is not None and isinstance(val_2023, (int, float)):
                    bra_value = val_2023
                    if val_2023 < 3.0:
                        failures.append(f"Row {r['row']}: PWRWASBRAXX ResidualCapacity(2023)={val_2023:.3f}, expected >= 3.0 GW")

            # Sum first year values for total
            first_year = min(r['year_values'].keys()) if r['year_values'] else None
            if first_year:
                val = r['year_values'][first_year]
                if val is not None and isinstance(val, (int, float)):
                    was_total += val

    if bra_value is None:
        failures.append("PWRWASBRAXX ResidualCapacity row not found or has no 2023 value")

    return failures, was_total, bra_value


def test_fix4_no_spurious_limits(rows):
    """TEST 3: No spurious UpperLimit/LowerLimit constraints."""
    failures = []

    for r in rows:
        param = r['param']
        tech = r['tech'].upper()

        # Check for spurious UpperLimit values: UpperLimit > 0 where LowerLimit = 0 or None
        # This catches artifacts where UpperLimit was written for zero LowerLimits
        if param == 'TotalTechnologyAnnualActivityUpperLimit':
            # Find corresponding LowerLimit row
            lower_values = None
            for r2 in rows:
                if (r2['tech'].upper() == tech and
                        r2['param'] == 'TotalTechnologyAnnualActivityLowerLimit'):
                    lower_values = r2['year_values']
                    break

            for year, val in r['year_values'].items():
                if val is not None and isinstance(val, (int, float)) and val > 0:
                    # Check if LowerLimit is 0 or None for this year
                    lower_val = lower_values.get(year) if lower_values else None
                    lower_is_zero = (lower_val is None or
                                     (isinstance(lower_val, (int, float)) and lower_val <= 0))
                    if lower_is_zero:
                        failures.append(f"Row {r['row']}: {r['tech']} UpperLimit({year})={val:.4f} — spurious constraint (LowerLimit=0)")

    return failures


def test_fix4_projection_mode_consistency(rows):
    """TEST 4: Projection.Mode consistent with cell values for activity limits.

    Note: rows with values + EMPTY may be intentional (e.g., migrated from A3
    with EMPTY to keep data as reference without constraining the model).
    Only flag rows with no values but 'User defined' (D2 should have set EMPTY).
    """
    failures = []
    limit_params = (
        'TotalTechnologyAnnualActivityLowerLimit',
        'TotalTechnologyAnnualActivityUpperLimit',
    )

    for r in rows:
        if r['param'] not in limit_params:
            continue

        has_nonzero = any(
            v is not None and isinstance(v, (int, float)) and v > 0
            for v in r['year_values'].values()
        )
        mode = r['proj_mode']

        if not has_nonzero and mode == 'User defined':
            failures.append(f"Row {r['row']}: {r['tech']} {r['param']} has no values but Projection.Mode='User defined', expected 'EMPTY'")

    return failures


def test_fix2_lower_limit_diagnostic(rows):
    """TEST 5: LowerLimit diagnostic comparison (informational)."""
    # Sum LowerLimit values for year 2025 by technology type
    tech_sums = defaultdict(float)

    for r in rows:
        if r['param'] != 'TotalTechnologyAnnualActivityLowerLimit':
            continue
        tech = r['tech'].upper()
        if not tech.startswith('PWR') or len(tech) < 6:
            continue

        fuel = tech[3:6]  # PWRHYDARGXX -> HYD
        val = r['year_values'].get(2025)
        if val is not None and isinstance(val, (int, float)):
            tech_sums[fuel] += val

    return dict(tech_sums)


def run_tests():
    script_dir = Path(__file__).resolve().parent
    base_path = script_dir / "A1_Outputs"

    scenarios = find_scenarios(script_dir)
    if not scenarios:
        print("[ERROR] No A1_Outputs_* folders found.")
        return 1

    # Use BAU as the primary test scenario
    test_scenario = 'BAU' if 'BAU' in scenarios else scenarios[0]
    param_path = base_path / f"A1_Outputs_{test_scenario}" / "A-O_Parametrization.xlsx"

    if not param_path.exists():
        print(f"[ERROR] File not found: {param_path}")
        return 1

    print(f"Testing scenario: {test_scenario}")
    print(f"File: {param_path}")
    print("=" * 80)

    rows, year_col_map, proj_mode_col = read_secondary_techs(param_path)
    if rows is None:
        print("[ERROR] Could not read Secondary Techs sheet.")
        return 1

    print(f"Read {len(rows)} rows from Secondary Techs sheet.")
    print()

    total_failures = 0

    # TEST 1: FIX 1 — Investment Projection.Mode
    print("-" * 60)
    print("TEST 1: Investment parameters Projection.Mode (FIX 1)")
    print("-" * 60)
    failures = test_fix1_investment_projection_mode(rows)
    if failures:
        print(f"  FAIL: {len(failures)} issues found")
        for f in failures[:10]:
            print(f"    - {f}")
        if len(failures) > 10:
            print(f"    ... and {len(failures) - 10} more")
        total_failures += len(failures)
    else:
        print("  PASS")
    print()

    # TEST 2: FIX 3 — WAS ResidualCapacity
    print("-" * 60)
    print("TEST 2: WAS ResidualCapacity preserved (FIX 3)")
    print("-" * 60)
    failures, was_total, bra_value = test_fix3_was_residual_capacity(rows)
    print(f"  WAS Total (first year): {was_total:.3f} GW")
    if bra_value is not None:
        print(f"  Brazil WAS (2023): {bra_value:.3f} GW")
    if failures:
        print(f"  FAIL: {len(failures)} issues found")
        for f in failures:
            print(f"    - {f}")
        total_failures += len(failures)
    else:
        print("  PASS")
    print()

    # TEST 3: FIX 4 — No spurious limits
    print("-" * 60)
    print("TEST 3: No spurious UpperLimit constraints (FIX 4)")
    print("-" * 60)
    failures = test_fix4_no_spurious_limits(rows)
    if failures:
        print(f"  FAIL: {len(failures)} spurious constraints found")
        for f in failures[:10]:
            print(f"    - {f}")
        if len(failures) > 10:
            print(f"    ... and {len(failures) - 10} more")
        total_failures += len(failures)
    else:
        print("  PASS")
    print()

    # TEST 4: FIX 4 — Projection.Mode consistency
    print("-" * 60)
    print("TEST 4: Projection.Mode consistency for activity limits (FIX 4)")
    print("-" * 60)
    failures = test_fix4_projection_mode_consistency(rows)
    if failures:
        print(f"  FAIL: {len(failures)} inconsistencies found")
        for f in failures[:10]:
            print(f"    - {f}")
        if len(failures) > 10:
            print(f"    ... and {len(failures) - 10} more")
        total_failures += len(failures)
    else:
        print("  PASS")
    print()

    # TEST 5: FIX 2 — LowerLimit diagnostic (informational)
    print("-" * 60)
    print("TEST 5: LowerLimit diagnostic for year 2025 (FIX 2 — informational)")
    print("-" * 60)
    tech_sums = test_fix2_lower_limit_diagnostic(rows)
    if tech_sums:
        print(f"  {'Tech':<6} {'D2 Sum (PJ)':>12}  {'Expected':>12}  {'Delta %':>10}  Status")
        print(f"  {'-'*6} {'-'*12}  {'-'*12}  {'-'*10}  {'-'*8}")
        for fuel in sorted(tech_sums.keys()):
            actual = tech_sums[fuel]
            expected = EXPECTED_LOWER_LIMITS_2025.get(fuel)
            if expected:
                delta_pct = (actual - expected) / expected * 100
                status = "OK" if abs(delta_pct) <= TOLERANCE_PCT * 100 else "CHECK"
                print(f"  {fuel:<6} {actual:>12.2f}  {expected:>12.0f}  {delta_pct:>+9.1f}%  {status}")
            else:
                print(f"  {fuel:<6} {actual:>12.2f}  {'n/a':>12}")
    else:
        print("  No LowerLimit data found for year 2025.")
    print()

    # Summary
    print("=" * 80)
    if total_failures == 0:
        print(f"ALL TESTS PASSED (Tests 1-4). Test 5 is informational only.")
        return 0
    else:
        print(f"TESTS FAILED: {total_failures} total issues found.")
        return 1


if __name__ == '__main__':
    sys.exit(run_tests())
