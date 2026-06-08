"""
D4_load_dsptrn_max_cap_inv.py

Escribe TotalAnnualMaxCapacityInvestment (GW) para las tecnologias de TRANSMISION
que inyectan en el nodo 02 (no para el sumador DSPTRN) en la hoja 'Demand Techs'
de A-O_Parametrization.xlsx del escenario destino INV (ver --scenario; OPT esta
desactivado a proposito para no alterar sus datos).

Motivacion:
    El sumador DSPTRN{ISO3}XX tiene capacidad residual 9999 y costo cero, asi que
    cualquier tope sobre el NUNCA es vinculante (BAU daba la misma inversion en
    transmision). En su lugar, topamos directamente las tecnologias de transmision
    reales, usando como cuota lo que BAU ya construyo (NewCapacity), escalado por un
    factor por escenario.
      - INV: factor < 1 (restringe). Como el respaldo PWRBCK inyecta en el mismo
        nodo 02, cubre la diferencia: menos transmision, mas respaldo -> el
        contraste INV vs BAU que se busca, manteniendo factibilidad.
    Los factores por escenario viven en SCENARIO_CONFIG. OPT no esta configurado
    (D4 no topa la transmision de OPT).

Regla por tech/ano Y (tres bloques):

    2023..2025  ->  celda VACIA (None): INV igual a BAU (sin tope). Son los anos
                    historicos/observados; el sync final de A3 los fija a BAU igual.
                    Se LIMPIA cualquier valor que corridas anteriores hubieran dejado.

    2026..2030  ->  cap = NewCapacity_BAU[tech, Y] * 1.0     [GW]
                    Tope = exactamente lo que BAU construyo (sin descuento). Es aqui
                    donde la Transmision de INV empieza a divergir de BAU. La
                    generacion y demas siguen = BAU hasta 2030 porque
                    sync_historical_from_bau.py fue modificado para NO pisar estas
                    filas TRN desde 2026.

    2031..2050  ->  cap = NewCapacity_BAU[tech, Y] * factor(Y)     [GW]
                    factor(Y) es lineal entre cap_start_factor y cap_end_factor
                    del escenario (INV: 0.99->0.80 restringe). Ver SCENARIO_CONFIG.

                    EXCEPCION por pais (solo INV, exempt_countries = PER, CRI, PAN, COL):
                    esos paises NO reciben el descuento; usan factor 1.0 (senda
                    BAU completa) + head-room EXEMPT_MARGIN (CRI: +5%). Sus seis
                    techs de transmision nodo-02 activaban el respaldo PWRBCK
                    (VariableCost 1750) al quedar topadas por debajo del pico;
                    restaurar la senda BAU lo elimina. (Antes esto se hacia
                    aguas abajo en fix_inv_transmission_caps.py; ahora vive aqui
                    como unica fuente de verdad del tope de transmision INV.)

En 2026..2050, donde BAU construyo 0 (o no hay dato) se escribe 0 -> prohibe
inversion ese ano; el respaldo PWRBCK (mismo nodo 02) cubre.

Solo se toca TotalAnnualMaxCapacityInvestment; el piso TotalAnnualMinCapacityInvestment
NO se escribe.

NewCapacity de BAU se lee de RELAC_TX_Combined_Inputs_Outputs.csv (formato largo;
columnas Scenario, YEAR, TECHNOLOGY, NewCapacity). Ese CSV contiene BAU/INV/OPT;
aqui filtramos Scenario == 'BAU'.

Usage:
    python t1_confection/D4_load_dsptrn_max_cap_inv.py --dry-run
    python t1_confection/D4_load_dsptrn_max_cap_inv.py --apply
"""
from __future__ import annotations

import argparse
import csv
import shutil
from datetime import datetime
from pathlib import Path

import openpyxl


HERE = Path(__file__).resolve().parent
COMBINED_CSV = HERE / 'RELAC_TX_Combined_Inputs_Outputs.csv'
A1_OUTPUTS = HERE / 'A1_Outputs'

SHEET = 'Demand Techs'
PARAM = 'TotalAnnualMaxCapacityInvestment'


def param_path(scenario: str) -> Path:
    """Ruta al A-O_Parametrization.xlsx del escenario destino (INV, OPT, ...)."""
    return A1_OUTPUTS / f'A1_Outputs_{scenario}' / 'A-O_Parametrization.xlsx'

# --- Configuracion comun (todos los escenarios) ----------------------------
# Bloque "igual a BAU": anos <= este quedan SIN tope (celda vacia). Solo los
# anos historicos/observados (2023-2025); desde 2026 la Transmision se topa.
EQUAL_THROUGH_YEAR = 2025

# Bloque plano: 2026..FLAT_FACTOR_THROUGH_YEAR usan FLAT_FACTOR (sin descuento
# => cap = NewCapacity_BAU exacto). Es la ventana donde TX diverge de BAU
# mientras el resto del escenario sigue = BAU hasta 2030.
FLAT_FACTOR_THROUGH_YEAR, FLAT_FACTOR = 2030, 1.0

# Bloque con factor lineal en 2031-2050 (los factores extremos van por escenario).
CAP_START_YEAR, CAP_END_YEAR = 2031, 2050

# --- Configuracion por escenario -------------------------------------------
# Para cada escenario destino, el factor lineal CAP_START_FACTOR -> CAP_END_FACTOR
# aplicado a NewCapacity(BAU) en 2031-2050, y exenciones por pais (mantienen
# senda BAU + head-room) para no activar el respaldo PWRBCK en su nodo 02.
#
#   INV: tope < BAU (0.99 -> 0.80) para empujar respaldo / contraste vs BAU,
#        SALVO paises exentos (PER/CRI/PAN/COL) que mantienen senda BAU + margen.
#
# OPT NO se configura aqui a proposito: D4 no debe topar la transmision de OPT
# (eso introducia cambios de datos en OPT respecto a la linea base 2affcec). Un
# `--scenario OPT` se rechaza en main() por escenario no configurado. Solo INV.
SCENARIO_CONFIG = {
    'INV': {
        'cap_start_factor': 0.99,
        'cap_end_factor': 0.80,
        'exempt_countries': ('PER', 'CRI', 'PAN', 'COL'),
        'exempt_restore_factor': 1.0,
        'exempt_margin': {'CRI': 0.05},  # CRI satura las 12 timeslices -> 5% extra
    },
}
DEFAULT_SCENARIOS = ('INV',)  # comportamiento por defecto = solo INV (como antes)

# Escenario de referencia dentro del CSV combinado.
BAU_SCENARIO = 'BAU'

# Familias de transmision que inyectan en el nodo 02 (19 paises c/u = 114 techs).
TRN_PREFIXES = ('PWRTRN', 'TRNNLI', 'TRNRPO', 'RNWTRN', 'RNWRPO', 'RNWNLI')

# Se recorren todos los anos del horizonte; 2023-2025 se vacian, 2026-2050 se topan.
YEARS = list(range(2023, 2051))
# ---------------------------------------------------------------------------


def factor_for_year(year: int, cfg: dict, country: str | None = None) -> float:
    """Factor de tope para los anos topados (2026..2050), segun config `cfg`.

    2026..2030  -> FLAT_FACTOR (1.0): cap = NewCapacity_BAU exacto, sin descuento.
    2031..2050  -> lineal cfg['cap_start_factor'] -> cfg['cap_end_factor'],
                   SALVO paises exentos (cfg['exempt_countries']): esos usan
                   cfg['exempt_restore_factor'] (1.0) + head-room
                   cfg['exempt_margin'], o sea mantienen la senda BAU completa.
    """
    if year <= FLAT_FACTOR_THROUGH_YEAR:
        return FLAT_FACTOR
    if country in cfg['exempt_countries']:
        return cfg['exempt_restore_factor'] * (1.0 + cfg['exempt_margin'].get(country, 0.0))
    span = CAP_END_YEAR - CAP_START_YEAR
    frac = (year - CAP_START_YEAR) / span
    return cfg['cap_start_factor'] + (cfg['cap_end_factor'] - cfg['cap_start_factor']) * frac


def country_of(tech: str) -> str:
    """ISO3 embebido en un codigo de transmision length-11 (chars 6:9).

    Layout: PREFIX(6) + COUNTRY(3) + 'XX'. Todos los TRN_PREFIXES miden 6.
    """
    return tech[6:9]


def load_bau_new_capacity(csv_path: Path) -> dict[tuple[str, int], float]:
    """Lee NewCapacity de BAU para las familias TRN del CSV combinado.

    Devuelve {(tech, year) -> new_capacity_GW} solo para valores > 0. Si hay
    filas repetidas para un (tech, year) se queda con el maximo no nulo
    (NewCapacity no se reparte entre modos/timeslices).
    """
    if not csv_path.exists():
        raise FileNotFoundError(f'No existe: {csv_path}')

    bau: dict[tuple[str, int], float] = {}
    with open(csv_path, newline='', encoding='utf-8') as f:
        reader = csv.DictReader(f)
        for row in reader:
            if row.get('Scenario') != BAU_SCENARIO:
                continue
            tech = (row.get('TECHNOLOGY') or '').strip()
            if not tech.startswith(TRN_PREFIXES):
                continue
            raw_nc = (row.get('NewCapacity') or '').strip()
            if not raw_nc:
                continue
            try:
                nc = float(raw_nc)
                year = int(float(row['YEAR']))
            except (ValueError, TypeError, KeyError):
                continue
            if nc <= 0:
                continue
            key = (tech, year)
            prev = bau.get(key)
            if prev is None or nc > prev:
                bau[key] = nc
    return bau


def build_year_col_map(ws) -> dict[int, int]:
    year_cols: dict[int, int] = {}
    for col_idx, cell in enumerate(ws[1], start=1):
        v = cell.value
        if isinstance(v, int) and 2000 <= v <= 2100:
            year_cols[v] = col_idx
        elif isinstance(v, str) and v.isdigit() and 2000 <= int(v) <= 2100:
            year_cols[int(v)] = col_idx
    return year_cols


def build_row_index(ws) -> dict[tuple[str, str], int]:
    index: dict[tuple[str, str], int] = {}
    for r in range(2, ws.max_row + 1):
        tech = ws.cell(r, 2).value
        param = ws.cell(r, 5).value
        if isinstance(tech, str) and isinstance(param, str):
            index[(tech.strip(), param.strip())] = r
    return index


def run_scenario(scenario: str, cfg: dict, bau: dict, apply_changes: bool) -> int:
    """Aplica (o previsualiza) el tope de transmision para un escenario.

    Devuelve el numero de violaciones detectadas en verificacion (0 en dry-run).
    """
    pth = param_path(scenario)
    print(f'\n========== ESCENARIO {scenario} ==========')
    print(f'Target:    {pth}')
    print(f'Regla:     2023-{EQUAL_THROUGH_YEAR} sin tope (=BAU); '
          f'2026-{FLAT_FACTOR_THROUGH_YEAR} cap = NewCapacity_BAU x {FLAT_FACTOR:.2f}; '
          f'{CAP_START_YEAR}-{CAP_END_YEAR} cap = NewCapacity_BAU x factor '
          f"({cfg['cap_start_factor']:.2f} -> {cfg['cap_end_factor']:.2f})")
    if cfg['exempt_countries']:
        print(f"Exentos:   {cfg['exempt_countries']} -> factor "
              f"{cfg['exempt_restore_factor']:.2f} (senda BAU completa) + margen "
              f"{cfg['exempt_margin']} en {CAP_START_YEAR}-{CAP_END_YEAR}")
    else:
        print('Exentos:   (ninguno)')

    if not pth.exists():
        raise FileNotFoundError(f'No existe: {pth}')

    wb = openpyxl.load_workbook(pth)
    if SHEET not in wb.sheetnames:
        raise KeyError(f'Hoja {SHEET!r} no encontrada. Hojas: {wb.sheetnames}')
    ws = wb[SHEET]

    year_cols = build_year_col_map(ws)
    missing_yrs = [y for y in YEARS if y not in year_cols]
    if missing_yrs:
        raise ValueError(f'Anos faltantes en header de {SHEET}: {missing_yrs}')

    row_idx = build_row_index(ws)
    trn_rows = sorted(
        tech for (tech, param) in row_idx
        if param == PARAM and tech.startswith(TRN_PREFIXES)
    )
    if not trn_rows:
        raise ValueError(
            f'No se hallaron filas {PARAM} para prefijos {TRN_PREFIXES} en {SHEET}.'
        )

    # Aviso si BAU tiene techs que no estan en el xlsx (mapeo por string completo).
    techs_in_bau = sorted({t for (t, _y) in bau})
    not_in_xlsx = sorted(set(techs_in_bau) - set(trn_rows))
    if not_in_xlsx:
        print(f'  [WARN] techs con NewCapacity en BAU pero sin fila {PARAM} en xlsx: {not_in_xlsx}')

    # Escritura: 2023-2025 -> vacio (=BAU); 2026-2050 -> cap = bau * factor(y)
    print(f'\n=== WRITES ({PARAM}) ===')
    cells_cleared = 0
    cells_capped = 0
    techs_touched = 0
    sum_bau = 0.0
    sum_cap = 0.0
    for tech in trn_rows:
        row = row_idx[(tech, PARAM)]
        years_with_quota: list[int] = []  # anos 2026-2050 con NewCapacity_BAU > 0
        for y in YEARS:
            if y <= EQUAL_THROUGH_YEAR:
                ws.cell(row, year_cols[y]).value = None  # igual a BAU: sin tope
                cells_cleared += 1
                continue
            nc = bau.get((tech, y))
            cap_gw = nc * factor_for_year(y, cfg, country_of(tech)) if nc is not None else 0.0  # BAU=0 -> 0
            ws.cell(row, year_cols[y]).value = cap_gw
            cells_capped += 1
            if nc is not None:
                years_with_quota.append(y)
                sum_bau += nc
                sum_cap += cap_gw
        ws.cell(row, 7).value = 'User defined'   # Projection.Mode (col G)
        techs_touched += 1
        if years_with_quota:
            first_y, last_y = years_with_quota[0], years_with_quota[-1]
            first_v = bau[(tech, first_y)] * factor_for_year(first_y, cfg, country_of(tech))
            last_v = bau[(tech, last_y)] * factor_for_year(last_y, cfg, country_of(tech))
            print(f'  {tech}: {len(years_with_quota)} anos con cuota en 2026-2050 '
                  f'(resto=0) | {first_v:.4g} ({first_y}) .. {last_v:.4g} ({last_y})')
        else:
            print(f'  {tech}: sin cuota en 2026-2050 -> todos 0')

    print(f'\nTechs tocadas: {techs_touched} | celdas vaciadas (2023-{EQUAL_THROUGH_YEAR}): '
          f'{cells_cleared} | celdas con tope (2026-{CAP_END_YEAR}): {cells_capped}')
    print(f'Sanity (bloque con tope): suma cap = {sum_cap:.4f} GW  vs  '
          f'suma NewCapacity BAU = {sum_bau:.4f} GW '
          f'(ratio {sum_cap / sum_bau if sum_bau else 0:.4f})')

    if not apply_changes:
        print(f'\n[DRY-RUN] {scenario}: no se guardaron cambios.')
        wb.close()
        return 0

    ts = datetime.now().strftime('%Y%m%d-%H%M%S')
    backup = pth.with_suffix(f'.backup-trn-maxinv-{ts}.xlsx')
    shutil.copy2(pth, backup)
    print(f'\nBackup: {backup.name}')
    wb.save(pth)
    wb.close()
    print(f'Guardado: {pth}')

    # Verificacion
    print('\n=== VERIFICATION ===')
    wb2 = openpyxl.load_workbook(pth, data_only=True)
    ws2 = wb2[SHEET]
    yc2 = build_year_col_map(ws2)
    ri2 = build_row_index(ws2)

    violations = 0
    for tech in trn_rows:
        r = ri2[(tech, PARAM)]
        for y in YEARS:
            v = ws2.cell(r, yc2[y]).value
            if y <= EQUAL_THROUGH_YEAR:
                # debe quedar VACIO (=BAU, sin tope)
                if v is not None:
                    print(f'  NOTEMPTY {tech} {y}: got {v!r}, esperado vacio')
                    violations += 1
                continue
            nc = bau.get((tech, y))
            expected = nc * factor_for_year(y, cfg, country_of(tech)) if nc is not None else 0.0
            if not isinstance(v, (int, float)) or abs(v - expected) > 1e-6 * max(1.0, abs(expected)):
                print(f'  VALUE {tech} {y}: got {v!r}, expected {expected:.6f}')
                violations += 1
    wb2.close()
    print(f'Verificadas {len(trn_rows)} techs. Violations: {violations}')
    return violations


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    group = parser.add_mutually_exclusive_group()
    group.add_argument('--apply', action='store_true', help='escribe cambios (crea backup)')
    group.add_argument('--dry-run', action='store_true', help='preview sin escribir (default)')
    parser.add_argument(
        '--scenario', default=','.join(DEFAULT_SCENARIOS),
        help=f'escenario(s) destino, separados por coma. Opciones: '
             f'{",".join(SCENARIO_CONFIG)}. Default: {",".join(DEFAULT_SCENARIOS)}.')
    args = parser.parse_args()

    apply_changes = args.apply and not args.dry_run
    mode = 'APPLY' if apply_changes else 'DRY-RUN'
    scenarios = [s.strip() for s in args.scenario.split(',') if s.strip()]
    unknown = [s for s in scenarios if s not in SCENARIO_CONFIG]
    if unknown:
        raise SystemExit(
            f'ERROR: escenario(s) no configurado(s): {unknown}. '
            f'Disponibles: {list(SCENARIO_CONFIG)}.')

    print(f'Mode: {mode}')
    print(f'BAU CSV:   {COMBINED_CSV}')
    print(f'Escenarios: {scenarios}')

    bau = load_bau_new_capacity(COMBINED_CSV)
    techs_in_bau = sorted({t for (t, _y) in bau})
    print(f'NewCapacity BAU leida: {len(bau)} celdas (tech,ano) > 0 '
          f'sobre {len(techs_in_bau)} techs de transmision.')

    total_violations = 0
    for scen in scenarios:
        total_violations += run_scenario(scen, SCENARIO_CONFIG[scen], bau, apply_changes)

    if apply_changes and total_violations:
        raise SystemExit(f'\nERROR: {total_violations} violaciones en verificacion.')


if __name__ == '__main__':
    main()
