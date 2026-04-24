"""
Load LAC Market-Realistic Capacity Ceilings into A-O_Parametrization.xlsx

Source: t1_confection/A1_Outputs/A1_Outputs_BAU/LAC_Capacity_Ceilings_Consolidated.md
Target: hoja 'Secondary Techs', parametro TotalAnnualMaxCapacityInvestment
        para PWRSPV{ISO3}XX y PWRWON{ISO3}XX (19 paises LAC, 38 filas).

Regla de coherencia con Min:
    Si Max_MD <= Min_existente para un ano dado, entonces Max = Min * 1.01

Usage:
    python t1_confection/D3_load_lac_max_capacity_caps.py
"""
import openpyxl
import shutil
from pathlib import Path
from datetime import datetime


PARAM_PATH = Path(__file__).parent / 'A1_Outputs' / 'A1_Outputs_BAU' / 'A-O_Parametrization.xlsx'
SHEET_NAME = 'Secondary Techs'
PARAM_MAX = 'TotalAnnualMaxCapacityInvestment'
PARAM_MIN = 'TotalAnnualMinCapacityInvestment'

# Caps del MD en GW/yr por periodo
SPV_CAPS = {
    'ARG': {'23-30': 0.7,  '31-40': 0.9,    '41-50': 1.1},
    'BOL': {'23-30': 0.08, '31-40': 0.12,   '41-50': 0.15},
    'BRA': {'23-30': 5.5,  '31-40': 6.5,    '41-50': 7.5},
    'BRB': {'23-30': 0.1,  '31-40': 0.12,   '41-50': 0.14},
    'CHL': {'23-30': 2.2,  '31-40': 2.7,    '41-50': 3.2},
    'COL': {'23-30': 1.2,  '31-40': 1.4,    '41-50': 1.6},
    'CRI': {'23-30': 0.06, '31-40': 0.09,   '41-50': 0.12},
    'DOM': {'23-30': 0.3,  '31-40': 0.375,  '41-50': 0.45},
    'ECU': {'23-30': 0.1,  '31-40': 0.175,  '41-50': 0.25},
    'SLV': {'23-30': 0.12, '31-40': 0.16,   '41-50': 0.2},
    'GTM': {'23-30': 0.15, '31-40': 0.2,    '41-50': 0.25},
    'HND': {'23-30': 0.18, '31-40': 0.23,   '41-50': 0.28},
    'HTI': {'23-30': 0.01, '31-40': 0.025,  '41-50': 0.04},
    'MEX': {'23-30': 1.0,  '31-40': 1.3,    '41-50': 1.6},
    'NIC': {'23-30': 0.06, '31-40': 0.09,   '41-50': 0.12},
    'PAN': {'23-30': 0.2,  '31-40': 0.275,  '41-50': 0.35},
    'PRY': {'23-30': 0.07, '31-40': 0.135,  '41-50': 0.2},
    'PER': {'23-30': 0.5,  '31-40': 0.75,   '41-50': 1.0},
    'URY': {'23-30': 0.08, '31-40': 0.1,    '41-50': 0.12},
}

WON_CAPS = {
    'ARG': {'23-30': 0.3,   '31-40': 0.5,   '41-50': 0.7},
    'BOL': {'23-30': 0.03,  '31-40': 0.05,  '41-50': 0.06},
    'BRA': {'23-30': 3.5,   '31-40': 4.5,   '41-50': 5.5},
    'BRB': {'23-30': 0.06,  '31-40': 0.08,  '41-50': 0.1},
    'CHL': {'23-30': 1.2,   '31-40': 1.7,   '41-50': 2.2},
    'COL': {'23-30': None,  '31-40': 0.5,   '41-50': 0.7},  # 23-30 via override por ano
    'CRI': {'23-30': 0.07,  '31-40': 0.085, '41-50': 0.1},
    'DOM': {'23-30': 0.12,  '31-40': 0.16,  '41-50': 0.2},
    'ECU': {'23-30': 0.05,  '31-40': 0.085, '41-50': 0.12},
    'SLV': {'23-30': 0.02,  '31-40': 0.035, '41-50': 0.05},
    'GTM': {'23-30': 0.08,  '31-40': 0.11,  '41-50': 0.14},
    'HND': {'23-30': 0.06,  '31-40': 0.09,  '41-50': 0.12},
    'HTI': {'23-30': 0.0,   '31-40': 0.008, '41-50': 0.015},
    'MEX': {'23-30': 0.9,   '31-40': 1.05,  '41-50': 1.2},
    'NIC': {'23-30': 0.03,  '31-40': 0.045, '41-50': 0.06},
    'PAN': {'23-30': 0.04,  '31-40': 0.06,  '41-50': 0.08},
    'PRY': {'23-30': 0.0,   '31-40': 0.02,  '41-50': 0.04},
    'PER': {'23-30': 0.25,  '31-40': 0.425, '41-50': 0.6},
    'URY': {'23-30': 0.03,  '31-40': 0.04,  '41-50': 0.05},
}

# Nota del MD: COL WON 0.1 GW/yr 2025-2027, 0.3 GW/yr 2028-2030.
COL_WON_OVERRIDE_BY_YEAR = {
    2025: 0.1, 2026: 0.1, 2027: 0.1,
    2028: 0.3, 2029: 0.3, 2030: 0.3,
}

ISO3_LIST = list(SPV_CAPS.keys())
# 2023 y 2024 quedan fuera por decision del usuario: los techos aplican desde 2025.
YEARS = list(range(2025, 2051))


def period_for_year(year):
    if 2023 <= year <= 2030:
        return '23-30'
    if 2031 <= year <= 2040:
        return '31-40'
    return '41-50'


def cap_for(iso3, tech, year):
    """Get the MD-derived cap for a given iso3+tech+year."""
    if tech == 'WON' and iso3 == 'COL' and year in COL_WON_OVERRIDE_BY_YEAR:
        return COL_WON_OVERRIDE_BY_YEAR[year]
    caps = SPV_CAPS if tech == 'SPV' else WON_CAPS
    return caps[iso3][period_for_year(year)]


def build_year_col_map(ws):
    """Return dict[year:int] -> col_idx:int based on header row."""
    year_cols = {}
    for col_idx, cell in enumerate(ws[1], start=1):
        val = cell.value
        if isinstance(val, int) and 2000 <= val <= 2100:
            year_cols[val] = col_idx
        elif isinstance(val, str) and val.isdigit():
            yr = int(val)
            if 2000 <= yr <= 2100:
                year_cols[yr] = col_idx
    return year_cols


def build_row_index(ws):
    """Return dict[(tech, parameter)] -> row_idx."""
    index = {}
    for row_idx in range(2, ws.max_row + 1):
        tech = ws.cell(row=row_idx, column=2).value  # col B
        param = ws.cell(row=row_idx, column=5).value  # col E
        if tech and param:
            index[(tech, param)] = row_idx
    return index


def main():
    if not PARAM_PATH.exists():
        raise FileNotFoundError(f'No existe: {PARAM_PATH}')

    print(f'Loading workbook: {PARAM_PATH}')
    wb = openpyxl.load_workbook(PARAM_PATH)
    if SHEET_NAME not in wb.sheetnames:
        raise KeyError(f'Hoja {SHEET_NAME!r} no encontrada. Hojas: {wb.sheetnames}')
    ws = wb[SHEET_NAME]

    year_cols = build_year_col_map(ws)
    missing_years = [y for y in YEARS if y not in year_cols]
    if missing_years:
        raise ValueError(f'Anos faltantes en header: {missing_years}')
    print(f'Mapped {len(year_cols)} year columns (2023 -> col {year_cols[2023]}, 2050 -> col {year_cols[2050]})')

    row_idx = build_row_index(ws)
    print(f'Indexed {len(row_idx)} (tech, parameter) rows')

    # Validar que existan las 38 filas Max (si faltan, hay que crearlas)
    missing_max_rows = []
    missing_min_rows = []
    for iso3 in ISO3_LIST:
        for tech in ('SPV', 'WON'):
            tech_code = f'PWR{tech}{iso3}XX'
            if (tech_code, PARAM_MAX) not in row_idx:
                missing_max_rows.append(tech_code)
            if (tech_code, PARAM_MIN) not in row_idx:
                missing_min_rows.append(tech_code)
    if missing_max_rows:
        raise ValueError(f'Faltan filas Max: {missing_max_rows}')
    if missing_min_rows:
        print(f'WARNING: {len(missing_min_rows)} tech(s) sin fila Min (se asume Min=None): {missing_min_rows}')

    # Procesar
    adjustments = []  # lista de (tech_code, year, max_md, min_val, max_final)
    writes = []       # resumen por fila
    for iso3 in ISO3_LIST:
        for tech in ('SPV', 'WON'):
            tech_code = f'PWR{tech}{iso3}XX'
            max_row = row_idx[(tech_code, PARAM_MAX)]
            min_row = row_idx.get((tech_code, PARAM_MIN))

            row_values = []
            row_adjusted = 0
            for year in YEARS:
                col = year_cols[year]
                max_md = cap_for(iso3, tech, year)
                min_val = None
                if min_row is not None:
                    raw = ws.cell(row=min_row, column=col).value
                    if isinstance(raw, (int, float)):
                        min_val = float(raw)

                if min_val is not None and max_md <= min_val:
                    max_final = min_val * 1.01
                    adjustments.append((tech_code, year, max_md, min_val, max_final))
                    row_adjusted += 1
                else:
                    max_final = max_md

                ws.cell(row=max_row, column=col).value = max_final
                row_values.append(max_final)

            # Projection.Mode = 'User defined' en col G
            ws.cell(row=max_row, column=7).value = 'User defined'
            writes.append((tech_code, row_values, row_adjusted))

    # Summary
    print('\n=== WRITE SUMMARY ===')
    for tech_code, vals, adj in writes:
        sample = f'{vals[0]:.4g} .. {vals[7]:.4g} | {vals[8]:.4g} .. {vals[17]:.4g} | {vals[18]:.4g} .. {vals[-1]:.4g}'
        note = f'  [{adj} adj by Min*1.01]' if adj else ''
        print(f'  {tech_code}: {sample}{note}')

    if adjustments:
        print(f'\n=== MIN*1.01 ADJUSTMENTS ({len(adjustments)} cells) ===')
        for tech_code, year, max_md, min_val, max_final in adjustments:
            print(f'  {tech_code} {year}: MD={max_md:.4g}, Min={min_val:.4g} -> Max={max_final:.6g}')

    # Backup
    ts = datetime.now().strftime('%Y%m%d-%H%M%S')
    backup_path = PARAM_PATH.with_suffix(f'.backup-{ts}.xlsx')
    shutil.copy2(PARAM_PATH, backup_path)
    print(f'\nBackup: {backup_path.name}')

    wb.save(PARAM_PATH)
    print(f'Saved: {PARAM_PATH}')

    # Verificacion post-guardado
    print('\n=== VERIFICATION (re-read) ===')
    wb2 = openpyxl.load_workbook(PARAM_PATH, data_only=True)
    ws2 = wb2[SHEET_NAME]
    year_cols2 = build_year_col_map(ws2)
    row_idx2 = build_row_index(ws2)
    violations = 0
    checked = 0
    for iso3 in ISO3_LIST:
        for tech in ('SPV', 'WON'):
            tech_code = f'PWR{tech}{iso3}XX'
            max_r = row_idx2[(tech_code, PARAM_MAX)]
            min_r = row_idx2.get((tech_code, PARAM_MIN))
            for year in YEARS:
                col = year_cols2[year]
                max_v = ws2.cell(row=max_r, column=col).value
                min_v = ws2.cell(row=min_r, column=col).value if min_r else None
                if max_v is None:
                    print(f'  MISSING: {tech_code} {year} Max is None')
                    violations += 1
                elif isinstance(min_v, (int, float)) and max_v < min_v:
                    print(f'  VIOLATION: {tech_code} {year} Max={max_v} < Min={min_v}')
                    violations += 1
                checked += 1
            pm = ws2.cell(row=max_r, column=7).value
            if pm != 'User defined':
                print(f'  MODE: {tech_code} Projection.Mode={pm!r} (esperado User defined)')
                violations += 1
    print(f'Checked {checked} (tech,year) cells + 38 projection modes. Violations: {violations}')


if __name__ == '__main__':
    main()
